"""
Caja USD con método PEPS (primero entrado, primero salido).

Motor puro (sin Streamlit ni base de datos) para la herramienta «Caja USD».

Reglas del estudio:
- Dos bolsillos: stock de dólares físicos y caja de ahorro (bancario). Cada uno con su cola PEPS.
- Ingreso común: TC vendedor BNA de la fecha; si ese día no hay cotización (fin de semana o
  feriado) se toma el día hábil siguiente (misma convención que inversiones_db.obtener_tc).
  El TC se puede corregir a mano en cada movimiento.
- Depósito en efectivo en la cuenta: sale del stock de dólares físicos por PEPS y entra al banco con
  el TC de origen de cada lote físico (no con el TC vendedor del día). Si el stock físico no alcanza,
  el faltante entra al TC vendedor y queda marcado para revisar.
- Extracción en efectivo: sale de la caja por PEPS y entra al stock físico con los mismos TC de origen.
- Transferencia desde/hacia cuenta propia y rescate/aporte de FCI: el ingreso lleva el TC de origen
  (se carga a mano, sale del análisis FCI o de la cuenta de origen); el egreso sale a costo PEPS, sin
  diferencia de cambio.
- Egreso común: sale por PEPS y se valúa al TC vendedor de la fecha; diferencia = valor - costo PEPS.
- Persona humana: sólo PEPS.
- Sociedad: PEPS y, al último día de cada mes, los USD que quedan (de ambos bolsillos) se reexpresan
  al TC vendedor BNA de ese día. La diferencia se reconoce en el mes y el mes siguiente sigue desde
  ese TC.
"""

from __future__ import annotations

import io
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from typing import Any, Iterable

EPS = 0.005

# ---------------------------------------------------------------------------
# Detalle de cada movimiento
# ---------------------------------------------------------------------------

# clave: (tipo, clase, etiqueta). clase: std | prop | fci | efvo
CATEGORIAS: dict[str, dict[str, str]] = {
    "ing_mep": {"tipo": "ing", "clase": "std", "t": "Acreditación por compra de dólar MEP"},
    "ing_deud": {"tipo": "ing", "clase": "std", "t": "Cobro de deudores por venta"},
    "ing_efvo": {"tipo": "ing", "clase": "efvo", "t": "Depósito en efectivo (sale del stock de dólares físicos)"},
    "ing_vta": {"tipo": "ing", "clase": "std", "t": "Venta de títulos / valores"},
    "ing_renta": {"tipo": "ing", "clase": "std", "t": "Rentas, dividendos o intereses"},
    "ing_prop": {"tipo": "ing", "clase": "prop", "t": "Transferencia de cuenta propia"},
    "ing_fci": {"tipo": "ing", "clase": "fci", "t": "Rescate de FCI"},
    "ing_ident": {"tipo": "ing", "clase": "std", "t": "A identificar"},
    "egr_tit": {"tipo": "egr", "clase": "std", "t": "Compra de títulos / valores"},
    "egr_prov": {"tipo": "egr", "clase": "std", "t": "Pago a proveedores"},
    "egr_gasto": {"tipo": "egr", "clase": "std", "t": "Gastos e impuestos bancarios"},
    "egr_ret": {"tipo": "egr", "clase": "std", "t": "Retiro / venta de USD"},
    "egr_efvo": {"tipo": "egr", "clase": "efvo", "t": "Extracción en efectivo (pasa al stock de dólares físicos)"},
    "egr_prop": {"tipo": "egr", "clase": "prop", "t": "Transferencia a cuenta propia"},
    "egr_fci": {"tipo": "egr", "clase": "fci", "t": "Aporte a FCI"},
    "egr_ident": {"tipo": "egr", "clase": "std", "t": "A identificar"},
}
ETIQUETA_A_CLAVE = {v["t"]: k for k, v in CATEGORIAS.items()}


def _norm(s: Any) -> str:
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode()
    return s.lower()


def sugerir_categoria(desc: str, tipo: str) -> str:
    """Sugerencia por el texto del concepto. Sin coincidencia clara: «A identificar»."""
    d = _norm(desc)
    if tipo == "ing":
        if re.search(r"rescate|fima|\bfci\b|fondo comun|cuotapart", d):
            return "ing_fci"
        if re.search(r"ctas?\.? propias|cuentas? propias|misma titular|entre cuentas|propia", d):
            return "ing_prop"
        if re.search(r"deposito en efectivo|dep\.? efectivo|efectivo", d):
            return "ing_efvo"
        if re.search(r"dolar mep|\bmep\b|compra.*dolar|dolar.*compra", d):
            return "ing_mep"
        if re.search(r"cobro|deudor|cliente|factura", d):
            return "ing_deud"
        if re.search(r"venta.*(titulo|bono|val|accion)|vta.*(titulo|bono|val)|liquidacion.*(titulo|bono)", d):
            return "ing_vta"
        if re.search(r"renta|dividendo|interes|cupon|amortiz", d):
            return "ing_renta"
        return "ing_ident"
    if re.search(r"ctas?\.? propias|cuentas? propias|misma titular|entre cuentas|propia", d):
        return "egr_prop"
    if re.search(r"suscripcion|aporte.*fci|\bfima\b|fondo comun", d):
        return "egr_fci"
    if re.search(r"extraccion.*efectivo|extraccion|retiro en efectivo", d):
        return "egr_efvo"
    if re.search(r"comp\.?.*(titulos|val)|compra.*(titulo|bono|val|accion)|\bal30\b|\bgd30\b", d):
        return "egr_tit"
    if re.search(r"comision|\biva\b|impuesto|sellos|ley 25\.?413|imp\.? ?(deb|cred)|percep|mantenimiento|gasto", d):
        return "egr_gasto"
    if re.search(r"proveedor|pago a|factura", d):
        return "egr_prov"
    if re.search(r"retiro|venta de (usd|dolar)", d):
        return "egr_ret"
    return "egr_ident"


# ---------------------------------------------------------------------------
# Serie de TC BNA
# ---------------------------------------------------------------------------

def _a_fecha(x: Any) -> date | None:
    if x is None:
        return None
    if isinstance(x, datetime):
        return x.date()
    if isinstance(x, date):
        return x
    s = str(x).strip()
    m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})", s)
    if m:
        return date(int(m[1]), int(m[2]), int(m[3]))
    m = re.match(r"^(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{2,4})", s)
    if m:
        y = int(m[3])
        y += 2000 if y < 100 else 0
        return date(y, int(m[2]), int(m[1]))
    return None


@dataclass
class TCResuelto:
    valor: float
    fuente: date
    etiqueta: str
    lejano: bool = False


class SerieTC:
    """Serie BNA (compra/venta) con las reglas de búsqueda del estudio."""

    def __init__(self, filas: Iterable[tuple[Any, float | None, float | None]]):
        self._v: dict[date, float] = {}
        self._c: dict[date, float] = {}
        for f, compra, venta in filas:
            d = _a_fecha(f)
            if d is None:
                continue
            if venta is not None and float(venta) > 0:
                self._v[d] = float(venta)
            if compra is not None and float(compra) > 0:
                self._c[d] = float(compra)
        self._fechas = sorted(self._v)

    def __len__(self) -> int:
        return len(self._fechas)

    @property
    def ultima(self) -> date | None:
        return self._fechas[-1] if self._fechas else None

    def venta_op(self, f: date) -> TCResuelto | None:
        """Vendedor del día; si ese día no hay cotización, el día hábil siguiente.
        Sólo si la fecha supera toda la serie se usa el último día anterior (queda marcado)."""
        if not self._fechas:
            return None
        if f in self._v:
            return TCResuelto(self._v[f], f, "vend. del día")
        sig = next((d for d in self._fechas if d > f), None)
        if sig is not None:
            return TCResuelto(self._v[sig], sig, f"vend. sig. háb. {sig:%d/%m/%Y}", lejano=(sig - f).days > 5)
        ant = self._fechas[-1]
        return TCResuelto(self._v[ant], ant, f"vend. del {ant:%d/%m/%Y} (falta la serie posterior)", lejano=True)

    def venta_cierre(self, f: date) -> TCResuelto | None:
        """Vendedor del último día del mes; si no cotizó, el último día hábil anterior."""
        if not self._fechas:
            return None
        if f in self._v:
            return TCResuelto(self._v[f], f, "vend. último día")
        ant = [d for d in self._fechas if d <= f]
        if ant:
            d = ant[-1]
            return TCResuelto(self._v[d], d, f"vend. últ. háb. {d:%d/%m/%Y}", lejano=(f - d).days > 5)
        d = self._fechas[0]
        return TCResuelto(self._v[d], d, f"vend. del {d:%d/%m/%Y}", lejano=True)

    def a_filas(self) -> list[tuple[date, float | None, float]]:
        return [(d, self._c.get(d), self._v[d]) for d in self._fechas]


# ---------------------------------------------------------------------------
# Motor
# ---------------------------------------------------------------------------

@dataclass
class Lote:
    fecha: date
    usd: float
    tc: float
    origen: str = ""


@dataclass
class ResultadoPeps:
    caja: list[dict] = field(default_factory=list)
    aplicaciones: list[dict] = field(default_factory=list)
    lotes: list[dict] = field(default_factory=list)
    cierres: list[dict] = field(default_factory=list)
    avisos: list[str] = field(default_factory=list)
    resumen: dict[str, Any] = field(default_factory=dict)


def _consumir(lotes: list[Lote], usd: float) -> tuple[list[Lote], float]:
    """Saca `usd` de la cola PEPS. Devuelve (partes consumidas, faltante)."""
    out: list[Lote] = []
    pend = usd
    for l in lotes:
        if pend <= EPS:
            break
        if l.usd <= EPS:
            continue
        t = min(l.usd, pend)
        l.usd -= t
        pend -= t
        out.append(Lote(l.fecha, t, l.tc, l.origen))
    lotes[:] = [l for l in lotes if l.usd > EPS]
    return out, (pend if pend > EPS else 0.0)


def _fin_mes(d: date) -> date:
    nxt = date(d.year + (d.month == 12), (d.month % 12) + 1, 1)
    return nxt - timedelta(days=1)


def _sig_mes(d: date) -> date:
    return date(d.year + (d.month == 12), (d.month % 12) + 1, 1)


def armar_caja_peps(
    movs: list[dict],
    serie: SerieTC,
    *,
    titular: str = "humana",
    lotes_banc_ini: list[dict] | None = None,
    lotes_fis_ini: list[dict] | None = None,
    cierre_hasta: date | None = None,
    cierres_tc: dict[str, float] | None = None,
) -> ResultadoPeps:
    """
    movs: dicts con fecha (date), tipo ('ing'|'egr'), usd, cat (clave de CATEGORIAS), desc,
          tc (opcional: TC manual), bolsillo ('banc' por defecto, 'fis' para dólares físicos).
    lotes_*_ini: dicts con fecha, usd, tc (opcional; vacío = vendedor BNA de esa fecha).
    """
    res = ResultadoPeps()
    soc = titular == "sociedad"
    cierres_tc = cierres_tc or {}
    L: dict[str, list[Lote]] = {"banc": [], "fis": []}
    NOMBRE = {"banc": "Caja de ahorro", "fis": "Dólares físicos"}
    dif_real = 0.0
    dif_cierre = 0.0
    tot_ing = tot_egr = 0.0
    n_id = 0

    def tc_op(f: date) -> TCResuelto | None:
        return serie.venta_op(f)

    # Saldos iniciales
    for b, lotes in (("fis", lotes_fis_ini or []), ("banc", lotes_banc_ini or [])):
        for x in lotes:
            f, usd = _a_fecha(x.get("fecha")), float(x.get("usd") or 0)
            if f is None or usd <= EPS:
                continue
            tc = x.get("tc")
            tc = float(tc) if tc not in (None, "") and float(tc) > 0 else None
            et = "manual"
            if tc is None:
                r = tc_op(f)
                if r is None:
                    res.avisos.append(f"Saldo inicial {NOMBRE[b]} del {f:%d/%m/%Y}: sin TC (cargá el TC o la serie BNA).")
                    continue
                tc, et = r.valor, r.etiqueta
                res.avisos.append(f"Saldo inicial {NOMBRE[b]} del {f:%d/%m/%Y}: sin TC cargado, se usó {et} ({tc:,.2f}).")
            L[b].append(Lote(f, usd, tc, "saldo inicial"))

    def saldo(b: str) -> float:
        return sum(l.usd for l in L[b])

    def costo(b: str) -> float:
        return sum(l.usd * l.tc for l in L[b])

    def fila(fecha, detalle, desc, ing, egr, tc, tc_et, valor, costo_p, dif, revisar="", bolsillo="banc", mid=""):
        res.caja.append({
            "ID": mid, "Fecha": fecha, "Bolsillo": NOMBRE[bolsillo], "Detalle": detalle, "Concepto": desc,
            "Ingresos_USD": ing, "Egresos_USD": egr, "TC": tc, "Origen_TC": tc_et,
            "Valor_ARS": valor, "Costo_PEPS_ARS": costo_p, "Dif_Cambio": dif,
            "Saldo_Caja_USD": round(saldo("banc"), 2), "Saldo_Fisico_USD": round(saldo("fis"), 2),
            "Revisar": revisar,
        })

    def aplic(mid, fecha, desc, bolsillo, partes: list[Lote]):
        for p in partes:
            res.aplicaciones.append({
                "ID": mid, "Fecha_mov": fecha, "Concepto": desc, "Sale_de": NOMBRE[bolsillo],
                "USD": round(p.usd, 2), "Fecha_lote": p.fecha, "TC_lote": p.tc, "Origen_lote": p.origen,
            })

    # Eventos
    ev: list[tuple[date, int, int, Any]] = []
    for i, m in enumerate(movs or []):
        f = _a_fecha(m.get("fecha"))
        usd = float(m.get("usd") or 0)
        if f is None or usd <= EPS:
            continue
        ev.append((f, 1 if m["tipo"] == "ing" else 2, i, m))
    fechas = [e[0] for e in ev]
    if soc and (fechas or L["banc"] or L["fis"]):
        desde = min(fechas + [l.fecha for b in L for l in L[b]])
        hasta = cierre_hasta or (max(fechas) if fechas else desde)
        cur = date(desde.year, desde.month, 1)
        while cur <= date(hasta.year, hasta.month, 1):
            ev.append((_fin_mes(cur), 3, 0, {"cierre": cur.strftime("%Y-%m")}))
            cur = _sig_mes(cur)
    ev.sort(key=lambda e: (e[0], e[1], e[2]))

    for f, _k, _i, m in ev:
        if "cierre" in m:
            ym = m["cierre"]
            if saldo("banc") <= EPS and saldo("fis") <= EPS:
                continue
            man = cierres_tc.get(ym)
            r = None if man else serie.venta_cierre(f)
            if not man and r is None:
                res.avisos.append(f"Cierre {ym}: sin TC vendedor BNA, no se reexpresó.")
                continue
            tcc = float(man) if man else r.valor
            et = "manual" if man else r.etiqueta
            if r is not None and r.lejano:
                res.avisos.append(f"Cierre {ym}: el TC usado es del {r.fuente:%d/%m/%Y}; cargá la serie completa.")
            for b in ("fis", "banc"):
                s = saldo(b)
                if s <= EPS:
                    continue
                c0 = costo(b)
                dif = s * tcc - c0
                dif_cierre += dif
                res.cierres.append({
                    "Mes": ym, "Fecha": f, "Bolsillo": NOMBRE[b], "Saldo_USD": round(s, 2),
                    "Costo_antes_ARS": round(c0, 2), "TC_cierre": tcc, "Origen_TC": et,
                    "Valor_ARS": round(s * tcc, 2), "Dif_Cambio": round(dif, 2),
                })
                L[b] = [Lote(f, s, tcc, f"cierre {ym}")]
                fila(f, "Cierre de mes", f"Reexpresión al TC vendedor BNA del último día ({ym})", None, None, tcc, et,
                     round(s * tcc, 2), round(c0, 2), round(dif, 2), bolsillo=b, mid=f"C{ym}{b}")
            continue

        n_id += 1
        mid = f"M{n_id:04d}"
        usd = float(m["usd"])
        desc = str(m.get("desc") or "")
        b = m.get("bolsillo") or "banc"
        cat = m.get("cat") or sugerir_categoria(desc, m["tipo"])
        info = CATEGORIAS[cat]
        det = info["t"]
        tcm = m.get("tc")
        tcm = float(tcm) if tcm not in (None, "") and float(tcm) > 0 else None
        clase = info["clase"] if b == "banc" else "std"

        if m["tipo"] == "ing":
            tot_ing += usd
            if clase == "efvo" and tcm is None:
                partes, falta = _consumir(L["fis"], usd)
                aplic(mid, f, desc, "fis", partes)
                for p in partes:
                    L["banc"].append(Lote(f, p.usd, p.tc, "efectivo (stock físico)"))
                ok = usd - falta
                c0 = sum(p.usd * p.tc for p in partes)
                rev = ""
                if falta > EPS:
                    r = tc_op(f)
                    if r is None:
                        res.avisos.append(f"{f:%d/%m/%Y}: depósito en efectivo de USD {usd:,.2f} sin stock físico y sin TC.")
                        rev = "sin stock físico y sin TC"
                    else:
                        L["banc"].append(Lote(f, falta, r.valor, "efectivo sin stock físico"))
                        c0 += falta * r.valor
                        rev = f"faltan USD {falta:,.2f} en el stock físico (se usó {r.etiqueta})"
                        res.avisos.append(f"{f:%d/%m/%Y}: {rev}.")
                tcp = c0 / usd if usd else None
                fila(f, det, desc, usd, None, tcp, "stock físico PEPS", round(c0, 2), round(c0, 2), None, rev, mid=mid)
            elif clase in ("prop", "fci"):
                if tcm is None:
                    res.avisos.append(f"{f:%d/%m/%Y}: {det} de USD {usd:,.2f} sin TC de origen; no entra al saldo hasta cargarlo.")
                    fila(f, det, desc, usd, None, None, "falta TC de origen", None, None, None, "falta TC de origen", mid=mid)
                else:
                    L[b].append(Lote(f, usd, tcm, det))
                    fila(f, det, desc, usd, None, tcm, "TC de origen (manual)", round(usd * tcm, 2), None, None, bolsillo=b, mid=mid)
            else:
                if tcm is not None:
                    tc, et, rev = tcm, "manual", ""
                else:
                    r = tc_op(f)
                    if r is None:
                        res.avisos.append(f"{f:%d/%m/%Y}: ingreso de USD {usd:,.2f} sin TC (cargá la serie BNA).")
                        fila(f, det, desc, usd, None, None, "falta TC", None, None, None, "falta TC", bolsillo=b, mid=mid)
                        continue
                    tc, et, rev = r.valor, r.etiqueta, ("TC lejano" if r.lejano else "")
                    if r.lejano:
                        res.avisos.append(f"{f:%d/%m/%Y}: el TC usado es del {r.fuente:%d/%m/%Y}; cargá la serie completa.")
                L[b].append(Lote(f, usd, tc, det))
                fila(f, det, desc, usd, None, tc, et, round(usd * tc, 2), None, None, rev, bolsillo=b, mid=mid)
        else:
            tot_egr += usd
            partes, falta = _consumir(L[b], usd)
            aplic(mid, f, desc, b, partes)
            ok = usd - falta
            c0 = sum(p.usd * p.tc for p in partes)
            rev = ""
            if falta > EPS:
                rev = f"supera el saldo por USD {falta:,.2f}"
                res.avisos.append(f"{f:%d/%m/%Y} [{NOMBRE[b]}]: egreso de USD {usd:,.2f} {rev}.")
            if clase == "efvo" and b == "banc":
                for p in partes:
                    L["fis"].append(Lote(f, p.usd, p.tc, "extracción en efectivo"))
                fila(f, det, desc, None, usd, (c0 / ok if ok > EPS else None), "costo PEPS (pasa al stock físico)", round(c0, 2),
                     round(c0, 2), None, rev, mid=mid)
            elif clase in ("prop", "fci"):
                fila(f, det, desc, None, usd, (c0 / ok if ok > EPS else None), "costo PEPS (traspaso)", round(c0, 2),
                     round(c0, 2), None, rev, bolsillo=b, mid=mid)
            else:
                if tcm is not None:
                    tc, et = tcm, "manual"
                else:
                    r = tc_op(f)
                    if r is None:
                        res.avisos.append(f"{f:%d/%m/%Y}: egreso de USD {usd:,.2f} sin TC (cargá la serie BNA).")
                        fila(f, det, desc, None, usd, None, "falta TC", None, round(c0, 2), None, "falta TC", bolsillo=b, mid=mid)
                        continue
                    tc, et = r.valor, r.etiqueta
                    if r.lejano:
                        rev = (rev + " TC lejano").strip()
                        res.avisos.append(f"{f:%d/%m/%Y}: el TC usado es del {r.fuente:%d/%m/%Y}; cargá la serie completa.")
                valor = ok * tc
                dif = valor - c0
                dif_real += dif
                fila(f, det, desc, None, usd, tc, et, round(valor, 2), round(c0, 2), round(dif, 2), rev, bolsillo=b, mid=mid)

    for b in ("fis", "banc"):
        for l in L[b]:
            res.lotes.append({
                "Bolsillo": NOMBRE[b], "Fecha_lote": l.fecha, "USD": round(l.usd, 2), "TC": l.tc,
                "Costo_ARS": round(l.usd * l.tc, 2), "Origen": l.origen,
            })
    res.resumen = {
        "Saldo_Caja_USD": round(saldo("banc"), 2),
        "Saldo_Fisico_USD": round(saldo("fis"), 2),
        "Costo_Caja_ARS": round(costo("banc"), 2),
        "Costo_Fisico_ARS": round(costo("fis"), 2),
        "Dif_Cambio_Egresos": round(dif_real, 2),
        "Dif_Cambio_Cierres": round(dif_cierre, 2),
        "Dif_Cambio_Total": round(dif_real + dif_cierre, 2),
        "Total_Ingresos_USD": round(tot_ing, 2),
        "Total_Egresos_USD": round(tot_egr, 2),
        "Titular": "Sociedad" if soc else "Persona humana",
    }
    return res


# ---------------------------------------------------------------------------
# Excel (valores de origen tipeados; importes calculados con fórmulas)
# ---------------------------------------------------------------------------

def exportar_caja_peps_excel(res: ResultadoPeps, serie: SerieTC | None = None, titular: str = "humana") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    CONT = '_-* #,##0.00_-;-* #,##0.00_-;_-* "-"??_-;_-@_-'
    FEC = "dd/mm/yyyy"
    HF = PatternFill("solid", fgColor="1F4E79")
    HFONT = Font(bold=True, color="FFFFFF")

    wb = Workbook()

    # --- Aplicaciones (hoja de apoyo del costo PEPS)
    wa = wb.active
    wa.title = "Aplicaciones PEPS"
    ha = ["ID", "Fecha mov.", "Concepto", "Sale de", "USD", "Fecha lote", "TC lote", "Costo ARS"]
    wa.append(ha)
    for i, a in enumerate(res.aplicaciones, start=2):
        wa.append([a["ID"], a["Fecha_mov"], a["Concepto"], a["Sale_de"], a["USD"], a["Fecha_lote"], a["TC_lote"], f"=E{i}*G{i}"])
        wa.cell(i, 2).number_format = FEC
        wa.cell(i, 6).number_format = FEC
        for c in (5, 7, 8):
            wa.cell(i, c).number_format = CONT
    na = max(len(res.aplicaciones) + 1, 2)

    # --- Caja
    wc = wb.create_sheet("Caja", 0)
    hc = ["Fecha", "Bolsillo", "Detalle", "Concepto", "Ingresos USD", "Egresos USD", "TC", "Origen del TC",
          "Valor ARS", "Costo PEPS ARS", "Dif. de cambio", "Saldo caja USD", "Saldo físico USD", "Revisar", "ID"]
    wc.append(hc)
    r0 = 2
    for i, x in enumerate(res.caja, start=r0):
        cierre = x["Detalle"] == "Cierre de mes"
        wc.append([x["Fecha"], x["Bolsillo"], x["Detalle"], x["Concepto"], x["Ingresos_USD"], x["Egresos_USD"], x["TC"], x["Origen_TC"],
                   None, None, None, None, None, x["Revisar"], x["ID"]])
        usd = f"(N(E{i})+N(F{i}))"
        if cierre:
            # cierre: el saldo reexpresado y el costo anterior son datos del motor; la diferencia es fórmula
            wc.cell(i, 9, x["Valor_ARS"])
            wc.cell(i, 10, x["Costo_PEPS_ARS"])
            wc.cell(i, 11, f"=I{i}-J{i}")
        elif x["Egresos_USD"] is not None:
            if "traspaso" in (x["Origen_TC"] or "") or "pasa al stock" in (x["Origen_TC"] or ""):
                wc.cell(i, 10, f"=SUMIFS('Aplicaciones PEPS'!$H$2:$H${na},'Aplicaciones PEPS'!$A$2:$A${na},O{i})")
                wc.cell(i, 9, f"=J{i}")
            else:
                wc.cell(i, 10, f"=SUMIFS('Aplicaciones PEPS'!$H$2:$H${na},'Aplicaciones PEPS'!$A$2:$A${na},O{i})")
                if x["TC"] is not None:
                    wc.cell(i, 9, f"=SUMIFS('Aplicaciones PEPS'!$E$2:$E${na},'Aplicaciones PEPS'!$A$2:$A${na},O{i})*G{i}")
                    wc.cell(i, 11, f"=I{i}-J{i}")
        else:
            if "stock físico" in (x["Origen_TC"] or ""):
                wc.cell(i, 9, x["Valor_ARS"])
            elif x["TC"] is not None:
                wc.cell(i, 9, f"=E{i}*G{i}")
        wc.cell(i, 12, x["Saldo_Caja_USD"])
        wc.cell(i, 13, x["Saldo_Fisico_USD"])
        wc.cell(i, 1).number_format = FEC
        for c in (5, 6, 7, 9, 10, 11, 12, 13):
            wc.cell(i, c).number_format = CONT
    last = r0 + len(res.caja) - 1
    t = last + 2
    wc.cell(t, 4, "Totales")
    wc.cell(t, 4).font = Font(bold=True)
    for col, letter in ((5, "E"), (6, "F"), (11, "K")):
        wc.cell(t, col, f"=SUM({letter}{r0}:{letter}{last})" if last >= r0 else 0)
        wc.cell(t, col).number_format = CONT
        wc.cell(t, col).font = Font(bold=True)

    # --- Lotes vigentes
    wl = wb.create_sheet("Lotes vigentes")
    wl.append(["Bolsillo", "Fecha lote", "USD", "TC", "Costo ARS", "Origen"])
    for i, x in enumerate(res.lotes, start=2):
        wl.append([x["Bolsillo"], x["Fecha_lote"], x["USD"], x["TC"], f"=C{i}*D{i}", x["Origen"]])
        wl.cell(i, 2).number_format = FEC
        for c in (3, 4, 5):
            wl.cell(i, c).number_format = CONT
    if res.lotes:
        n = len(res.lotes) + 1
        wl.cell(n + 2, 1, "Total")
        wl.cell(n + 2, 3, f"=SUM(C2:C{n})")
        wl.cell(n + 2, 5, f"=SUM(E2:E{n})")
        for c in (3, 5):
            wl.cell(n + 2, c).number_format = CONT
            wl.cell(n + 2, c).font = Font(bold=True)

    # --- Cierres mensuales (sociedades)
    if res.cierres:
        wm = wb.create_sheet("Cierres mensuales")
        wm.append(["Mes", "Fecha", "Bolsillo", "Saldo USD", "Costo antes ARS", "TC cierre", "Origen del TC", "Valor ARS", "Dif. de cambio"])
        for i, x in enumerate(res.cierres, start=2):
            wm.append([x["Mes"], x["Fecha"], x["Bolsillo"], x["Saldo_USD"], x["Costo_antes_ARS"], x["TC_cierre"], x["Origen_TC"], f"=D{i}*F{i}", f"=H{i}-E{i}"])
            wm.cell(i, 2).number_format = FEC
            for c in (4, 5, 6, 8, 9):
                wm.cell(i, c).number_format = CONT
        n = len(res.cierres) + 1
        wm.cell(n + 2, 1, "Total")
        wm.cell(n + 2, 9, f"=SUM(I2:I{n})")
        wm.cell(n + 2, 9).number_format = CONT
        wm.cell(n + 2, 9).font = Font(bold=True)

    # --- Cotizaciones usadas
    if serie is not None and len(serie):
        wt = wb.create_sheet("Cotizaciones BNA")
        wt.append(["Fecha", "Comprador", "Vendedor"])
        for i, (d, c, v) in enumerate(serie.a_filas(), start=2):
            wt.append([d, c, v])
            wt.cell(i, 1).number_format = FEC
            wt.cell(i, 2).number_format = CONT
            wt.cell(i, 3).number_format = CONT

    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.fill = HF
            cell.font = HFONT
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.freeze_panes = "A2"
        for col in ws.columns:
            letter = get_column_letter(col[0].column)
            m = max((len(str(c.value or "")) for c in col[:80] if not str(c.value or "").startswith("=")), default=8)
            ws.column_dimensions[letter].width = min(max(m + 2, 11), 48)

    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()
