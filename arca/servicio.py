"""Corridas del emisor para la web: validar, emitir, consultar. Sin reescribir WSFEv1."""
from __future__ import annotations

import argparse
import datetime as dt
import json
import os
import tempfile
import zipfile
from contextlib import contextmanager
from decimal import Decimal
from io import BytesIO
from pathlib import Path

import openpyxl
from openpyxl import Workbook

from arca.codigos import ETIQUETA_CBTE, TIPOS_CBTE
from arca.comprobante import COLUMNAS, solo_digitos
from arca.constancia import consultar_categoria
from arca.pdf import generar_pdf
from arca.persistencia import (
    cargar_emisor,
    emisiones_que_bloquean,
    guardar_emisor,
    insertar_emision,
    listar_emisiones,
    obtener_emision,
)
from arca.secretos import CredencialesError, materializar
from arca.ta_store import obtener_ta_durable
from arca.vista import VistaComprobante
from arca.wsfe import WSFE, WSFEError
from emitir import (
    RES_COLS,
    EmisionAbortada,
    chequear_ambiente,
    correr,
    escribir,
    hoja_resultado,
)

__all__ = ["CredencialesError", "EmisionAbortada", "WSFEError"]


def excel_desde_filas(filas: list[dict]) -> bytes:
    """Arma un xlsx compatible con leer_planilla (una fila de encabezados + datos)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Facturas"
    for i, (_k, encabezado, _ancho) in enumerate(COLUMNAS, 1):
        ws.cell(1, i, encabezado)
    for r, fila in enumerate(filas, 2):
        for i, (clave, _enc, _ancho) in enumerate(COLUMNAS, 1):
            if clave in ("iva", "total"):
                continue
            valor = fila.get(clave)
            if valor not in (None, ""):
                ws.cell(r, i, valor)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _mapa_columnas(ws):
    fila_enc = None
    for r in range(1, 15):
        vals = [str(ws.cell(r, c).value or "") for c in range(1, len(COLUMNAS) + 3)]
        if any(v.startswith("CUIT emisor") for v in vals):
            fila_enc = r
            break
    if fila_enc is None:
        return None, {}
    enc = {str(ws.cell(fila_enc, c).value or "").strip(): c for c in range(1, ws.max_column + 1)}
    mapa = {}
    for clave, encabezado, _ancho in COLUMNAS:
        if encabezado in enc:
            mapa[clave] = enc[encabezado]
    return fila_enc, mapa


def preparar_planilla(path: str) -> None:
    """Completa datos de emisor vacíos desde la base y normaliza el CUIT a 11 dígitos."""
    wb = openpyxl.load_workbook(path)
    if "Facturas" not in wb.sheetnames:
        wb.save(path)
        return
    ws = wb["Facturas"]
    fila_enc, mapa = _mapa_columnas(ws)
    if not fila_enc or "cuit_emisor" not in mapa:
        wb.save(path)
        return
    campos = {
        "razon_emisor": "razon_social",
        "dom_emisor": "domicilio",
        "iibb_emisor": "iibb",
        "inicio_act": "inicio_actividades",
        "cond_iva_emisor": "condicion_iva",
    }
    for r in range(fila_enc + 1, (ws.max_row or fila_enc) + 1):
        crudo = ws.cell(r, mapa["cuit_emisor"]).value
        cuit = solo_digitos(crudo)
        if len(cuit) == 11:
            ws.cell(r, mapa["cuit_emisor"], cuit)
        emisor = cargar_emisor(cuit) if len(cuit) == 11 else None
        if not emisor:
            continue
        for clave, attr in campos.items():
            col = mapa.get(clave)
            if not col:
                continue
            if ws.cell(r, col).value not in (None, ""):
                continue
            valor = emisor.get(attr) or ""
            if not valor:
                continue
            if clave == "inicio_act":
                try:
                    valor = dt.date.fromisoformat(str(valor)[:10])
                except ValueError:
                    pass
            ws.cell(r, col, valor)
    wb.save(path)


def _fila_historial(row: dict) -> list:
    def entero(v):
        if v in (None, ""):
            return None
        try:
            return int(float(v))
        except (TypeError, ValueError):
            return v

    return [
        row.get("fecha_hora"),
        row.get("ambiente"),
        row.get("modo") or "ENVÍO",
        row.get("filas"),
        row.get("id_factura"),
        row.get("cuit_emisor"),
        entero(row.get("pto_vta")),
        row.get("tipo"),
        entero(row.get("numero")),
        row.get("cae"),
        row.get("vto_cae"),
        row.get("estado"),
        row.get("mensajes"),
        row.get("total"),
        row.get("pdf_nombre"),
        row.get("huella"),
        row.get("fecha_cbte"),
        row.get("cbte_asociado"),
    ]


def sembrar_historial(path: str, env: str) -> int:
    """Copia APROBADO/VERIFICAR de la base a la hoja Resultado. Devuelve la última fila sembrada."""
    wb = openpyxl.load_workbook(path)
    ws = hoja_resultado(wb)
    presentes = set()
    for r in range(2, ws.max_row + 1):
        huella = ws.cell(r, 16).value
        if huella:
            presentes.add(str(huella))
    for row in emisiones_que_bloquean(env):
        huella = str(row.get("huella") or "")
        if huella and huella in presentes:
            continue
        escribir(ws, _fila_historial(row))
        if huella:
            presentes.add(huella)
    ultima = ws.max_row
    wb.save(path)
    return ultima


def _snapshot(cb) -> dict | None:
    if cb is None or not getattr(cb, "det", None):
        return None

    def conv(valor):
        if isinstance(valor, dt.datetime):
            return valor.date().isoformat()
        if isinstance(valor, dt.date):
            return valor.isoformat()
        if isinstance(valor, Decimal):
            return float(valor)
        return valor

    f0 = {}
    for clave, valor in (cb.f0 or {}).items():
        if clave == "_fila":
            continue
        f0[clave] = conv(valor)
    items = [{k: conv(v) for k, v in it.items()} for it in (cb.items or [])]
    return {
        "f0": f0,
        "det": cb.det,
        "items": items,
        "clase": cb.clase,
        "pto_vta": cb.pto_vta,
        "cbte_tipo": cb.cbte_tipo,
        "cuit": cb.cuit,
        "fecha": cb.fecha.isoformat() if getattr(cb, "fecha", None) else None,
        "es_fce": bool(cb.es_fce),
        "asoc_desc": cb.asoc_desc,
    }


def _recordar_emisor(cb) -> None:
    if cb is None or not getattr(cb, "cuit", None):
        return
    if cargar_emisor(cb.cuit):
        return
    f0 = cb.f0 or {}
    inicio = f0.get("inicio_act")
    if isinstance(inicio, dt.datetime):
        inicio = inicio.date().isoformat()
    elif isinstance(inicio, dt.date):
        inicio = inicio.isoformat()
    guardar_emisor(
        cb.cuit,
        razon_social=str(f0.get("razon_emisor") or ""),
        domicilio=str(f0.get("dom_emisor") or ""),
        condicion_iva=str(f0.get("cond_iva_emisor") or ""),
        iibb=str(f0.get("iibb_emisor") or ""),
        inicio_actividades=str(inicio or ""),
    )


def _persistir_envio(fila, cb, _pdf_path) -> None:
    if len(fila) < 3 or str(fila[2]) == "DRY-RUN":
        return
    payload = _snapshot(cb)
    insertar_emision(list(fila), json.dumps(payload, ensure_ascii=False) if payload else None)
    estado = str(fila[11] or "") if len(fila) > 11 else ""
    if estado.startswith("APROBADO"):
        _recordar_emisor(cb)


def _leer_filas_nuevas(path: str, desde: int) -> list[dict]:
    wb = openpyxl.load_workbook(path, data_only=True)
    if "Resultado" not in wb.sheetnames:
        return []
    ws = wb["Resultado"]
    cols = [ws.cell(1, c).value or RES_COLS[c - 1] for c in range(1, len(RES_COLS) + 1)]
    salida = []
    for r in range(desde + 1, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, len(RES_COLS) + 1)]
        if all(v in (None, "") for v in vals):
            continue
        salida.append(dict(zip(cols, vals)))
    return salida


def _pdfs_de_filas(filas: list[dict]) -> list[tuple[str, bytes]]:
    vistos = []
    for fila in filas:
        ruta = fila.get("PDF")
        if not ruta or not os.path.isfile(str(ruta)):
            continue
        nombre = os.path.basename(str(ruta))
        vistos.append((nombre, Path(ruta).read_bytes()))
    return vistos


def zip_pdfs(pdfs: list[tuple[str, bytes]]) -> bytes:
    buf = BytesIO()
    with zipfile.ZipFile(buf, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        usados = set()
        for nombre, data in pdfs:
            base = nombre or "comprobante.pdf"
            candidato = base
            n = 2
            while candidato in usados:
                candidato = f"{n}_{base}"
                n += 1
            usados.add(candidato)
            zf.writestr(candidato, data)
    return buf.getvalue()


def _correr_bytes(excel_bytes: bytes, env: str, *, dry_run: bool, confirmar: bool,
                  incluir_ejemplos: bool) -> dict:
    opciones = argparse.Namespace(
        excel="lote.xlsx",
        env=env,
        dry_run=dry_run,
        confirmar_produccion=bool(confirmar),
        incluir_ejemplos=bool(incluir_ejemplos) and not (env == "prod" and not dry_run),
        pdf_borrador=True,
        solo_filas=None,
        cert=None,
        key=None,
        cache=None,
        pdf_dir="pdf",
        dummy=False,
        parametros=None,
    )
    chequear_ambiente(opciones)
    tmp = tempfile.TemporaryDirectory(prefix="arca_fe_")
    try:
        raiz = tmp.name
        path = os.path.join(raiz, "lote.xlsx")
        Path(path).write_bytes(excel_bytes)
        preparar_planilla(path)
        ultima_semilla = sembrar_historial(path, env)
        pdf_dir = os.path.join(raiz, "pdf")
        os.makedirs(pdf_dir, exist_ok=True)
        cert = key = None
        cache = os.path.join(raiz, "cache")
        if not dry_run:
            cert, key = materializar(env, os.path.join(raiz, "cert"))
        opciones.excel = path
        opciones.cert = cert
        opciones.key = key
        opciones.cache = cache
        opciones.pdf_dir = pdf_dir
        corrida = correr(
            opciones,
            on_resultado=_persistir_envio if not dry_run else None,
            base_dir=raiz,
        )
        filas = _leer_filas_nuevas(path, ultima_semilla)
        pdfs = _pdfs_de_filas(filas)
        return {
            "resumen": corrida.get("resumen") or {},
            "filas": filas,
            "pdfs": pdfs,
            "excel": Path(path).read_bytes(),
            "ambiente": env,
            "modo": "DRY-RUN" if dry_run else "ENVÍO",
        }
    finally:
        tmp.cleanup()


def validar_excel(excel_bytes: bytes, env: str, incluir_ejemplos: bool = False) -> dict:
    """Valida y arma PDF borrador. No llama a ARCA."""
    return _correr_bytes(excel_bytes, env, dry_run=True, confirmar=False, incluir_ejemplos=incluir_ejemplos)


def emitir_excel(excel_bytes: bytes, env: str, *, confirmar_produccion: bool = False,
                 incluir_ejemplos: bool = False) -> dict:
    """Emite de verdad. Producción exige confirmar_produccion. Persiste cada resultado."""
    return _correr_bytes(
        excel_bytes, env, dry_run=False, confirmar=confirmar_produccion, incluir_ejemplos=incluir_ejemplos,
    )


def pdf_desde_payload(payload: dict, nro, cae, vto, marca: str | None = None) -> bytes:
    vista = VistaComprobante.from_payload(payload)
    fd, ruta = tempfile.mkstemp(suffix=".pdf")
    os.close(fd)
    try:
        generar_pdf(ruta, vista, int(nro or 0), cae or "0", vto, marca_agua=marca)
        return Path(ruta).read_bytes()
    finally:
        os.remove(ruta)


def pdf_de_emision(emision_id: int) -> bytes | None:
    row = obtener_emision(emision_id)
    if not row or not row.get("payload_json"):
        return None
    payload = json.loads(row["payload_json"])
    marca = "HOMOLOGACIÓN - SIN VALIDEZ" if row.get("ambiente") == "homo" else None
    if str(row.get("estado") or "").startswith("OK-DRYRUN"):
        marca = "BORRADOR - SIN CAE"
    return pdf_desde_payload(payload, row.get("numero") or 0, row.get("cae"), row.get("vto_cae"), marca)


@contextmanager
def _sesion_wsfe(env: str, cuit: str):
    tmp = tempfile.TemporaryDirectory(prefix="arca_ws_")
    try:
        cert, key = materializar(env, tmp.name)
        ta = obtener_ta_durable(env, cert, key, "wsfe", os.path.join(tmp.name, "cache"))
        auth = {"token": ta["token"], "sign": ta["sign"], "cuit": solo_digitos(cuit)}
        yield {"wsfe": WSFE(env), "auth": auth, "ta": ta}
    finally:
        tmp.cleanup()


def listar_puntos_venta(cuit: str, env: str) -> dict:
    with _sesion_wsfe(env, cuit) as ses:
        items, msgs = ses["wsfe"].param(ses["auth"], "FEParamGetPtosVenta", "")
    errores = [f"{c}: {m}" for tipo, c, m in msgs if tipo == "Errors"]
    if errores and not items:
        raise WSFEError("; ".join(errores))
    puntos = []
    omitidos = 0
    for it in items:
        baja = str(it.get("FchBaja") or "").strip()
        bloq = str(it.get("Bloqueado") or "").strip().upper()
        if baja or bloq == "S":
            omitidos += 1
            continue
        puntos.append({
            "Punto de venta": int(float(it.get("Nro") or 0)),
            "Emisión": it.get("EmisionTipo") or "",
            "Bloqueado": bloq or "N",
        })
    puntos.sort(key=lambda p: p["Punto de venta"])
    return {"puntos": puntos, "omitidos": omitidos, "avisos": [f"{t} {c}: {m}" for t, c, m in msgs]}


def ultimos_numeros(cuit: str, env: str, pto_vta: int, incluir_fce: bool = False) -> list[dict]:
    pares = [(etiq, cod) for etiq, (cod, _clase, es_fce) in TIPOS_CBTE.items() if incluir_fce or not es_fce]
    salida = []
    with _sesion_wsfe(env, cuit) as ses:
        for etiq, cod in pares:
            try:
                nro = ses["wsfe"].ultimo_autorizado(ses["auth"], int(pto_vta), cod)
                salida.append({"Tipo": etiq, "Código": cod, "Último número": nro, "Detalle": ""})
            except Exception as ex:
                salida.append({"Tipo": etiq, "Código": cod, "Último número": "", "Detalle": str(ex)})
    return salida


def _payload_guardado(cuit: str, env: str, pto: int, tipo: str, nro: int) -> dict | None:
    for row in listar_emisiones(solo_digitos(cuit), env, limite=500):
        if str(row.get("tipo") or "") != tipo:
            continue
        try:
            if int(float(row.get("pto_vta") or 0)) != int(pto):
                continue
            if int(float(row.get("numero") or 0)) != int(nro):
                continue
        except (TypeError, ValueError):
            continue
        if row.get("payload_json"):
            return json.loads(row["payload_json"])
    return None


def consultar_y_pdf(cuit: str, env: str, pto_vta: int, tipo: str, nro: int) -> dict:
    if tipo not in TIPOS_CBTE:
        raise ValueError(f"Tipo de comprobante inválido: {tipo}")
    cod, _clase, _fce = TIPOS_CBTE[tipo]
    with _sesion_wsfe(env, cuit) as ses:
        cons, msgs = ses["wsfe"].consultar_completo(ses["auth"], int(pto_vta), cod, int(nro))
    mensajes = [f"{t} {c}: {m}" for t, c, m in msgs]
    if not cons:
        return {"ok": False, "mensajes": mensajes or ["ARCA no devolvió el comprobante."], "pdf": None, "resumen": {}}
    cuit_digits = solo_digitos(cuit)
    payload = _payload_guardado(cuit_digits, env, int(pto_vta), tipo, int(nro))
    marca = "HOMOLOGACIÓN - SIN VALIDEZ" if env == "homo" else None
    cae = cons.get("CodAutorizacion")
    vto = cons.get("FchVto")
    if payload:
        pdf = pdf_desde_payload(payload, nro, cae, vto, marca)
        origen = "datos guardados de la emisión"
    else:
        vista = VistaComprobante.from_consulta(cons, cargar_emisor(cuit_digits), cuit_digits, int(pto_vta), cod)
        fd, ruta = tempfile.mkstemp(suffix=".pdf")
        os.close(fd)
        try:
            generar_pdf(ruta, vista, int(nro), cae or "0", vto, marca_agua=marca)
            pdf = Path(ruta).read_bytes()
        finally:
            os.remove(ruta)
        origen = "consulta ARCA (sin el detalle de ítems; cargá los datos del emisor para el encabezado)"
    resumen = {
        "Tipo": f"{tipo} ({cod})",
        "Punto de venta": int(pto_vta),
        "Número": int(nro),
        "Fecha": cons.get("CbteFch") or "",
        "CAE": cae or "",
        "Vencimiento CAE": vto or "",
        "Total": cons.get("ImpTotal") or "",
        "Doc. receptor": f"{cons.get('DocTipo') or ''} {cons.get('DocNro') or ''}".strip(),
        "Resultado": cons.get("Resultado") or "",
    }
    return {"ok": True, "mensajes": mensajes, "pdf": pdf, "resumen": resumen, "origen_pdf": origen,
            "etiqueta": ETIQUETA_CBTE.get(cod, tipo)}


def categoria_monotributo(env: str, cuit: str) -> dict:
    return consultar_categoria(env, solo_digitos(cuit))
