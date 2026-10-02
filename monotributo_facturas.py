"""Papel de trabajo de facturas para la recategorización de monotributo.

La extracción sigue el criterio de referencia (un PDF = un comprobante;
ORIGINAL / DUPLICADO / TRIPLICADO se cuentan una sola vez si coinciden).
Lo que el PDF no trae queda como FALTA DATO: no se completa con otro campo.
"""

from __future__ import annotations

import re
import unicodedata
from datetime import date, datetime
from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FALTA_DATO = "FALTA DATO"

COLUMNAS_PAPEL = [
    "Fecha",
    "Período Desde",
    "Período Hasta",
    "Monto",
    "N° Factura",
    "Tipo de cambio",
    "Moneda",
    "Denominación del comprador",
    "Archivo",
    "CUIT emisor",
    "Emisor",
]

_PREFIJOS = {
    "FACTURA": "FC",
    "NOTA DE CRÉDITO": "NC",
    "NOTA DE CREDITO": "NC",
    "NOTA DE DÉBITO": "ND",
    "NOTA DE DEBITO": "ND",
    "RECIBO": "RECIBO",
}

_ETIQUETA = {
    "FC": "Factura",
    "NC": "Nota de Crédito",
    "ND": "Nota de Débito",
    "RECIBO": "Recibo",
}

_RE_TIPO = r"(FACTURA|NOTA DE CR[ÉE]DITO|NOTA DE D[ÉE]BITO|RECIBO)"
_RE_NF = re.compile(r"^(FC|NC|ND|RECIBO)\s+([ABCEM])\s+(\d+)-(\d+)$")
# En el texto del PDF la razón social y la etiqueta siguiente suelen quedar en la misma línea.
_RE_CORTE_NOMBRE = re.compile(
    r"\s+(?:"
    r"Fecha de Emisi[oó]n"
    r"|Fecha de Inicio(?: de Actividades)?"
    r"|Inicio de Actividades"
    r"|Domicilio(?: Comercial)?"
    r"|CUIT"
    r"|Condici[oó]n(?:\s+frente al IVA|\s+de venta)?"
    r"|Ingresos Brutos"
    r"|Punto de Venta"
    r"|Comp\.?\s*Nro"
    r"|Per[ií]odo(?:\s+Facturado)?"
    r"|Moneda"
    r"|Tipo de Cambio"
    r"|Cotizaci[oó]n"
    r"|C[oó]digo"
    r"|CAE"
    r"|Apellido y Nombre"
    r")\s*:",
    re.IGNORECASE,
)
_MONEDAS_EXTRANJERAS = {"DOL", "USD", "U$S", "US$", "DOLAR", "DOLARES"}

_HEADER_FILL = PatternFill(fill_type="solid", fgColor="305496")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_FALTA_FONT = Font(color="C00000")
_BOLD = Font(bold=True)
_CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)
_CENTRO_DATO = Alignment(horizontal="center")


def numero_impreso(texto: str | None) -> float | None:
    """Miles con punto y decimales con coma, como en el armado de referencia."""
    if texto is None:
        return None
    limpio = str(texto).strip()
    if not limpio or not re.fullmatch(r"[\d.,]+", limpio):
        return None
    try:
        return float(limpio.replace(".", "").replace(",", "."))
    except ValueError:
        return None


def numero_tipo_cambio(texto: str | None) -> float | None:
    """Tipo de cambio impreso. El punto suelto (1.000000) es decimal, no miles."""
    if texto is None:
        return None
    limpio = str(texto).strip()
    if not limpio or not re.fullmatch(r"[\d.,]+", limpio):
        return None
    if "," in limpio:
        return numero_impreso(limpio)
    if re.fullmatch(r"\d{1,3}(?:\.\d{3})+", limpio):
        return numero_impreso(limpio)
    try:
        return float(limpio)
    except ValueError:
        return None


def es_moneda_extranjera(moneda: str | None) -> bool:
    texto = unicodedata.normalize("NFD", str(moneda or ""))
    texto = "".join(c for c in texto if unicodedata.category(c) != "Mn")
    texto = re.sub(r"\s+", "", texto).upper()
    if not texto or texto in {"$", "PES", "ARS", "PESOS"}:
        return False
    return texto in _MONEDAS_EXTRANJERAS or texto.startswith("DOLAR")


def importe_recategorizacion(
    monto: float | None,
    prefijo: str | None,
    moneda: str | None,
    tipo_cambio: str | None,
) -> tuple[float, float, float]:
    """Neto firmado para la categoría: NC resta; moneda extranjera × tipo de cambio.

    Devuelve (importe firmado, dólares, tipo de cambio numérico).
    El papel de trabajo no usa este importe: ahí va el monto impreso, en positivo.
    """
    if monto is None:
        return 0.0, 0.0, 0.0
    tc = numero_tipo_cambio(tipo_cambio) or 0.0
    dolares = 0.0
    base = float(monto)
    if es_moneda_extranjera(moneda) and tc > 0:
        dolares = round(abs(base), 2)
        base = round(dolares * tc, 2)
    signo = -1 if prefijo == "NC" else 1
    return round(abs(base) * signo, 2), dolares, tc


def parsear_texto_recat(texto: str, archivo: str = "") -> dict | None:
    """Una fila por PDF. Las copias se leen para verificar que coincidan."""
    paginas = [p for p in str(texto or "").split("\f") if p.strip()]
    if not paginas:
        return None
    primera = paginas[0]
    tipo = _tipo_comprobante(primera)
    letra = _letra(primera)
    pv = _buscar(primera, r"Punto de Venta:\s*(\d+)")
    nro = _buscar(primera, r"Comp\.?\s*Nro:\s*(\d+)")
    fecha = _buscar(primera, r"Fecha de Emisi[oó]n:\s*(\d\d/\d\d/\d{4})")
    desde = _buscar(primera, r"Desde:\s*(\d\d/\d\d/\d{4})")
    hasta = _buscar(primera, r"Hasta:\s*(\d\d/\d\d/\d{4})")
    simbolo, importe_txt = _importe(primera)
    moneda = _buscar(primera, r"Moneda:\s*(\S+)")
    tc = _buscar(primera, r"(?:Tipo de [Cc]ambio|Cotizaci[oó]n)[^:]*:\s*([\d.,]+)")
    comprador = _comprador(primera)
    emisor = _emisor_nombre(primera)
    cuit_fmt, cuit_digitos = _cuit_emisor(primera)
    if not any([tipo, letra, pv, nro, fecha, importe_txt, comprador, emisor, cuit_digitos]):
        return None

    prefijo = _prefijo(tipo)
    nro_factura = ""
    if tipo and letra and pv and nro and prefijo:
        nro_factura = f"{prefijo} {letra} {pv}-{nro}"

    avisos: list[str] = []
    if len(paginas) > 1 and _copias_difieren(paginas):
        nombre = archivo or "el PDF"
        avisos.append(
            f"{nombre}: las copias no coinciden en Comp. Nro o Importe Total. "
            "Se usó la primera página."
        )

    etiqueta = ""
    if prefijo in _ETIQUETA:
        etiqueta = f"{_ETIQUETA[prefijo]} {letra}".strip() if letra else _ETIQUETA[prefijo]
    elif tipo:
        etiqueta = tipo.title() if tipo.isupper() else tipo

    return {
        "Fecha": fecha or "",
        "Período Desde": desde or "",
        "Período Hasta": hasta or "",
        "Monto": numero_impreso(importe_txt),
        "N° Factura": nro_factura,
        "Tipo de cambio": tc or "",
        "Moneda": (moneda or simbolo or ""),
        "Denominación del comprador": comprador or "",
        "Archivo": archivo,
        "CUIT emisor": cuit_fmt,
        "Emisor": emisor or "",
        "Tipo": etiqueta,
        "_pref": prefijo or "",
        "_moneda_afip": moneda or "",
        "_cuit_digitos": cuit_digitos,
        "_avisos_copias": avisos,
    }


def deduplicar_filas_recat(filas: list[dict]) -> tuple[list[dict], list[dict]]:
    """Una fila por emisor + tipo + punto de venta + número. El primero se conserva."""
    vistos: dict[tuple, dict] = {}
    unicos: list[dict] = []
    duplicados: list[dict] = []
    for fila in filas:
        clave = _clave_fila(fila)
        if clave is not None and clave in vistos:
            conservado = vistos[clave]
            duplicados.append({
                "nro": str(fila.get("N° Factura") or conservado.get("N° Factura") or ""),
                "conservado": str(conservado.get("Archivo") or ""),
                "descartado": str(fila.get("Archivo") or ""),
            })
            continue
        if clave is not None:
            vistos[clave] = fila
        unicos.append(fila)
    return unicos, duplicados


def describir_saltos(filas: list[dict]) -> list[str]:
    """Huecos de numeración por emisor, tipo, letra y punto de venta."""
    grupos: dict[tuple, list[tuple[int, int, str]]] = {}
    cuit_visible: dict[str, str] = {}
    for fila in filas:
        nf = str(fila.get("N° Factura") or "")
        m = _RE_NF.match(nf)
        if not m:
            continue
        pref, letra, pv, nro = m.group(1), m.group(2), m.group(3), m.group(4)
        digitos = _cuit_digitos_de(fila)
        visible = _cuit_visible(fila)
        if digitos and visible:
            cuit_visible.setdefault(digitos, visible)
        grupos.setdefault((digitos, pref, letra, int(pv)), []).append((int(nro), len(nro), pv))

    mensajes: list[str] = []
    for clave, items in sorted(grupos.items()):
        numeros = sorted({n for n, _, _ in items})
        if len(numeros) < 2:
            continue
        presentes = set(numeros)
        faltan = [n for n in range(numeros[0], numeros[-1] + 1) if n not in presentes]
        if not faltan:
            continue
        ancho = max(a for _, a, _ in items)
        pv_txt = max((p for _, _, p in items), key=len)
        muestra = ", ".join(_fmt_nro(n, ancho) for n in faltan[:8])
        extra = f" (+{len(faltan) - 8})" if len(faltan) > 8 else ""
        digitos, pref, letra, _pv = clave
        visible = cuit_visible.get(digitos) or digitos
        quien = f"CUIT {visible} · " if visible else ""
        mensajes.append(
            f"{quien}{pref} {letra} · punto de venta {pv_txt}: falta {muestra}{extra} "
            f"(rango {_fmt_nro(numeros[0], ancho)}–{_fmt_nro(numeros[-1], ancho)})."
        )
    return mensajes


def exportar_papel_facturas(filas, notas: dict | None = None) -> bytes:
    """Una hoja Facturas, con el layout del Excel de referencia."""
    notas = notas or {}
    data = [_fila_excel(fila) for fila in _records(filas)]
    data.sort(key=_clave_orden)

    wb = Workbook()
    ws = wb.active
    ws.title = "Facturas"
    ws.append(COLUMNAS_PAPEL)
    for celda in ws[1]:
        celda.font = _HEADER_FONT
        celda.fill = _HEADER_FILL
        celda.alignment = _CENTRO
    for fila in data:
        ws.append(fila)

    ultima = len(data) + 1
    for row in ws.iter_rows(min_row=2, max_row=max(ultima, 1)):
        if ultima < 2:
            break
        for celda in row[:3]:
            if isinstance(celda.value, datetime):
                celda.number_format = "DD/MM/YYYY"
            celda.alignment = _CENTRO_DATO
        row[3].number_format = "#,##0.00"
        for celda in row[4:7]:
            celda.alignment = _CENTRO_DATO
        for celda in row:
            if celda.value == FALTA_DATO:
                celda.font = _FALTA_FONT

    ultima_col = get_column_letter(len(COLUMNAS_PAPEL))
    ws.auto_filter.ref = f"A1:{ultima_col}{max(ultima, 1)}"
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30

    fila_tot = ultima + 2 if ultima >= 2 else 3
    cursor = _escribir_bloque_totales(ws, fila_tot, ultima)
    anios = _anios_de(data)
    for anio in anios:
        ws.cell(cursor, 3, f"Año {anio}").font = _BOLD
        cursor += 1
        cursor = _escribir_bloque_totales(ws, cursor, ultima, anio=anio)

    emisores = _emisores_distintos(data)
    if len(emisores) > 1:
        for cuit, nombre in emisores:
            titulo = f"Emisor {cuit}"
            if nombre:
                titulo = f"{titulo} — {nombre}"
            celda = ws.cell(cursor, 1, titulo)
            celda.font = _BOLD
            ws.merge_cells(start_row=cursor, start_column=1, end_row=cursor, end_column=len(COLUMNAS_PAPEL))
            cursor += 1
            cursor = _escribir_bloque_totales(ws, cursor, ultima, cuit=cuit)
            for anio in _anios_de(data, cuit):
                ws.cell(cursor, 3, f"Año {anio}").font = _BOLD
                cursor += 1
                cursor = _escribir_bloque_totales(ws, cursor, ultima, anio=anio, cuit=cuit)

    _escribir_notas(ws, cursor, data, notas)
    _ajustar_anchos(ws, data)
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def _escribir_notas(ws, fila: int, data: list[list], notas: dict) -> None:
    duplicados = list(notas.get("duplicados") or [])
    copias = list(notas.get("copias") or [])
    saltos = describir_saltos(_registros_desde_excel(data))
    sin_fecha = any(fila[0] == FALTA_DATO for fila in data)

    lineas = ["Control de duplicados y numeración"]
    if duplicados:
        lineas.append("Duplicados (mismo tipo, punto de venta y número). Se cuenta una sola vez:")
        for dup in duplicados:
            nro = dup.get("nro") or FALTA_DATO
            lineas.append(
                f"- {nro}: también está en {dup.get('descartado') or 'otro archivo'}. "
                f"Se conserva {dup.get('conservado') or 'la primera'}."
            )
    else:
        lineas.append("No hay duplicados de tipo + punto de venta + número.")

    if saltos:
        lineas.append("Saltos de numeración por punto de venta:")
        lineas.extend(f"- {texto}" for texto in saltos)
    else:
        lineas.append("No hay saltos de numeración por punto de venta.")

    if copias:
        lineas.append("Copias del mismo PDF:")
        lineas.extend(f"- {texto}" for texto in copias)
    else:
        lineas.append("Las copias revisadas coinciden (o el PDF trae una sola página).")

    if sin_fecha:
        lineas.append("Hay comprobantes sin fecha: entran en el total general y no en el total por año.")

    for i, texto in enumerate(lineas):
        celda = ws.cell(fila + i, 1, texto)
        if i == 0 or texto.endswith(":"):
            celda.font = _BOLD
        ws.merge_cells(
            start_row=fila + i,
            start_column=1,
            end_row=fila + i,
            end_column=len(COLUMNAS_PAPEL),
        )


def _registros_desde_excel(data: list[list]) -> list[dict]:
    registros = []
    for fila in data:
        nro = fila[4]
        cuit = fila[9] if len(fila) > 9 else ""
        registros.append({
            "N° Factura": "" if nro == FALTA_DATO else nro,
            "CUIT emisor": "" if cuit == FALTA_DATO else cuit,
        })
    return registros


def _escribir_bloque_totales(ws, fila: int, ultima: int, anio: int | None = None, cuit: str | None = None) -> int:
    """Escribe Facturas, NC y Neto. Devuelve la fila siguiente, con un renglón en blanco."""
    _escribir_total(ws, fila, "Total Facturas ($)", _sumifs(ultima, "FC", anio, cuit), ultima)
    _escribir_total(ws, fila + 1, "Total Notas de Crédito ($)", _sumifs(ultima, "NC", anio, cuit), ultima)
    _escribir_total(ws, fila + 2, "Neto (Facturas − NC) ($)", f"=D{fila}-D{fila + 1}", ultima)
    return fila + 4


def _sumifs(ultima: int, prefijo: str, anio: int | None = None, cuit: str | None = None) -> str:
    partes = [f"$D$2:$D${ultima}", f"$E$2:$E${ultima}", f'"{prefijo}*"']
    if cuit is not None:
        partes.extend([f"$J$2:$J${ultima}", _criterio_excel(cuit)])
    if anio is not None:
        partes.extend([
            f"$A$2:$A${ultima}",
            f'">="&DATE({anio},1,1)',
            f"$A$2:$A${ultima}",
            f'"<="&DATE({anio},12,31)',
        ])
    return "=SUMIFS(" + ",".join(partes) + ")"


def _criterio_excel(valor: str) -> str:
    return '"' + str(valor).replace('"', '""') + '"'


def _anios_de(data: list[list], cuit: str | None = None) -> list[int]:
    anios = set()
    for fila in data:
        if cuit is not None and (fila[9] if len(fila) > 9 else "") != cuit:
            continue
        if isinstance(fila[0], datetime):
            anios.add(fila[0].year)
    return sorted(anios)


def _emisores_distintos(data: list[list]) -> list[tuple[str, str]]:
    nombres: dict[str, str] = {}
    for fila in data:
        cuit = fila[9] if len(fila) > 9 else FALTA_DATO
        nombre = fila[10] if len(fila) > 10 else ""
        if nombre == FALTA_DATO:
            nombre = ""
        nombres.setdefault(cuit, "")
        if nombre and not nombres[cuit]:
            nombres[cuit] = nombre
    claves = sorted(nombres, key=lambda c: (c == FALTA_DATO, str(c)))
    return [(c, nombres[c]) for c in claves]


def _escribir_total(ws, fila: int, etiqueta: str, formula: str, ultima: int) -> None:
    if ultima < 2:
        formula = "=0"
    etiqueta_celda = ws.cell(fila, 3, etiqueta)
    etiqueta_celda.font = _BOLD
    etiqueta_celda.alignment = Alignment(horizontal="right")
    valor = ws.cell(fila, 4, formula)
    valor.number_format = "#,##0.00"
    valor.font = _BOLD


def _fila_excel(fila: dict) -> list:
    return [
        _a_fecha(fila.get("Fecha")),
        _a_fecha(fila.get("Período Desde")),
        _a_fecha(fila.get("Período Hasta")),
        _a_monto(fila.get("Monto")),
        _a_texto(fila.get("N° Factura")),
        _a_texto(fila.get("Tipo de cambio")),
        _a_texto(fila.get("Moneda")),
        _a_texto(fila.get("Denominación del comprador")),
        _a_texto(fila.get("Archivo")),
        _a_texto(fila.get("CUIT emisor")),
        _a_texto(fila.get("Emisor")),
    ]


def _clave_orden(fila: list):
    fecha = fila[0] if isinstance(fila[0], datetime) else datetime.max
    nro = fila[4]
    nro_key = nro if nro and nro != FALTA_DATO else "\uffff"
    return (fecha, nro_key)


def _a_fecha(valor):
    if valor is None or valor == "" or _es_nan(valor):
        return FALTA_DATO
    if isinstance(valor, datetime):
        return valor
    if isinstance(valor, date):
        return datetime(valor.year, valor.month, valor.day)
    texto = str(valor).strip()
    if not texto or texto == FALTA_DATO:
        return FALTA_DATO
    try:
        return datetime.strptime(texto, "%d/%m/%Y")
    except ValueError:
        return FALTA_DATO


def _a_monto(valor):
    if valor is None or valor == "" or valor == FALTA_DATO or _es_nan(valor):
        return FALTA_DATO
    if isinstance(valor, bool):
        return FALTA_DATO
    if isinstance(valor, (int, float)):
        return float(valor)
    return FALTA_DATO


def _a_texto(valor) -> str:
    if valor is None or _es_nan(valor):
        return FALTA_DATO
    texto = str(valor).strip()
    return texto or FALTA_DATO


def _es_nan(valor) -> bool:
    try:
        return bool(pd.isna(valor)) and not isinstance(valor, str)
    except (TypeError, ValueError):
        return False


def _records(filas) -> list[dict]:
    if filas is None:
        return []
    if isinstance(filas, pd.DataFrame):
        if filas.empty:
            return []
        return filas.to_dict(orient="records")
    return list(filas)


def _ajustar_anchos(ws, data: list[list]) -> None:
    for i in range(1, len(COLUMNAS_PAPEL) + 1):
        letra = get_column_letter(i)
        largos = [len(COLUMNAS_PAPEL[i - 1])]
        for fila in data:
            largos.append(len(_texto_ancho(fila[i - 1])))
        ancho = max(largos)
        if i == 3:
            ancho = max(ancho, 26)
        ws.column_dimensions[letra].width = ancho + 3


def _texto_ancho(valor) -> str:
    if isinstance(valor, datetime):
        return valor.strftime("%d/%m/%Y")
    if isinstance(valor, float):
        return f"{valor:,.2f}"
    return str(valor)


def _fmt_nro(numero: int, ancho: int) -> str:
    texto = str(numero)
    if ancho > len(texto):
        return texto.zfill(ancho)
    return texto


def partes_numero(nro_factura: str) -> tuple[str, str, str, str] | None:
    """(prefijo, letra, punto de venta, número) tal como figuran en el PDF."""
    m = _RE_NF.match(str(nro_factura or ""))
    if not m:
        return None
    return m.group(1), m.group(2), m.group(3), m.group(4)


def _clave_fila(fila: dict):
    m = _RE_NF.match(str(fila.get("N° Factura") or ""))
    if m:
        return (
            _cuit_digitos_de(fila),
            m.group(1),
            m.group(2),
            int(m.group(3)),
            int(m.group(4)),
        )
    cae = str(fila.get("CAE") or "").strip()
    if cae:
        return ("CAE", _cuit_digitos_de(fila), cae)
    return None


def _cuit_digitos_de(fila: dict) -> str:
    bruto = str(fila.get("_cuit_digitos") or fila.get("CUIT emisor") or fila.get("CUIT Emisor") or "")
    if bruto.strip() == FALTA_DATO:
        return ""
    return re.sub(r"\D", "", bruto)


def _cuit_visible(fila: dict) -> str:
    texto = str(fila.get("CUIT emisor") or "").strip()
    if texto and texto != FALTA_DATO:
        return texto
    digitos = _cuit_digitos_de(fila)
    if len(digitos) == 11:
        return f"{digitos[:2]}-{digitos[2:10]}-{digitos[10:]}"
    return digitos


def _copias_difieren(paginas: list[str]) -> bool:
    claves = set()
    for pagina in paginas:
        nro = re.search(r"Comp\.?\s*Nro:\s*(\d+)", pagina, flags=re.IGNORECASE)
        importe = _importe_token(pagina)
        claves.add((nro.group(1) if nro else None, importe))
    return len(claves) != 1


def _importe_token(pagina: str) -> str | None:
    _simbolo, importe = _importe(pagina)
    return importe


def _importe(pagina: str) -> tuple[str | None, str | None]:
    m = re.search(r"Importe Total:\s*(\S+)\s+([\d.,]+)", pagina, flags=re.IGNORECASE)
    if m and not re.fullmatch(r"[\d.,]+", m.group(1)):
        return m.group(1), m.group(2)
    m = re.search(r"Importe Total:\s*(\$?)\s*([\d.,]+)", pagina, flags=re.IGNORECASE)
    if m:
        simbolo = m.group(1) or None
        return simbolo, m.group(2)
    return None, None


def _recortar_nombre(texto: str) -> str:
    """Corta el nombre donde empieza la etiqueta siguiente de la misma línea."""
    m = _RE_CORTE_NOMBRE.search(texto or "")
    if m:
        texto = texto[:m.start()]
    return re.sub(r"\s+", " ", texto).strip()


def _emisor_nombre(pagina: str) -> str | None:
    """Razón social del emisor. No usa la del comprador (Apellido y Nombre / Razón Social)."""
    for m in re.finditer(r"Raz[oó]n Social:\s*([^\n]*)", pagina, flags=re.IGNORECASE):
        previo = pagina[max(0, m.start() - 40):m.start()]
        if re.search(r"Apellido y Nombre\s*/\s*$", previo, flags=re.IGNORECASE):
            continue
        nombre = _recortar_nombre(m.group(1))
        if nombre:
            return nombre
    return None


def _cuit_emisor(pagina: str) -> tuple[str, str]:
    """Primer CUIT del encabezado: (con guiones, solo dígitos)."""
    m = re.search(
        r"CUIT[:\s]*(\d{2}[-.\s]?\d{8}[-.\s]?\d)",
        pagina,
        flags=re.IGNORECASE,
    )
    if not m:
        return "", ""
    digitos = re.sub(r"\D", "", m.group(1))
    if len(digitos) != 11:
        return digitos, digitos
    return f"{digitos[:2]}-{digitos[2:10]}-{digitos[10:]}", digitos


def _comprador(pagina: str) -> str | None:
    m = re.search(
        r"Apellido y Nombre / Raz[oó]n Social:[ \t]*([^\n]*)",
        pagina,
        flags=re.IGNORECASE,
    )
    if not m:
        return None
    nombre = _recortar_nombre(m.group(1))
    return nombre or None


def _tipo_comprobante(pagina: str) -> str | None:
    m = re.search(rf"\s{_RE_TIPO}\s*\n", pagina, flags=re.IGNORECASE)
    if not m:
        m = re.search(rf"\b{_RE_TIPO}\b", pagina, flags=re.IGNORECASE)
    if not m:
        return None
    return _sin_acentos(m.group(1)).upper()


def _prefijo(tipo: str | None) -> str | None:
    if not tipo:
        return None
    canon = _sin_acentos(tipo).upper()
    canon = re.sub(r"\s+", " ", canon).strip()
    if canon.startswith("NOTA DE CREDITO"):
        return "NC"
    if canon.startswith("NOTA DE DEBITO"):
        return "ND"
    if canon.startswith("FACTURA"):
        return "FC"
    if canon.startswith("RECIBO"):
        return "RECIBO"
    return _PREFIJOS.get(tipo, tipo)


def _letra(pagina: str) -> str | None:
    m = re.search(r"(?:^|\n)\s+.+?\s{3,}([ABCEM])\s{3,}(?:FACTURA|NOTA)", pagina)
    if m:
        return m.group(1)
    # pdfplumber a veces deja un solo espacio: "1 C FACTURA"
    m = re.search(r"(?:^|\s)([ABCEM])\s+(?:FACTURA|NOTA)\b", pagina)
    if m:
        return m.group(1)
    m = re.search(rf"\b{_RE_TIPO}\s+([ABCEM])\b", pagina, flags=re.IGNORECASE)
    if m:
        return m.group(2).upper()
    lineas = pagina.splitlines()
    for i, linea in enumerate(lineas):
        if not re.search(r"FACTURA|NOTA DE|RECIBO", linea, flags=re.IGNORECASE):
            continue
        desde = max(0, i - 4)
        hasta = min(len(lineas), i + 3)
        for vecina in lineas[desde:hasta]:
            sola = re.fullmatch(r"\s*([ABCEM])\s*", vecina)
            if sola:
                return sola.group(1)
            al_final = re.search(r"\s{3,}([ABCEM])\s*$", vecina)
            if al_final:
                return al_final.group(1)
    return None


def _buscar(pagina: str, patron: str) -> str | None:
    m = re.search(patron, pagina, flags=re.IGNORECASE)
    if not m:
        return None
    return m.group(1).strip()


def _sin_acentos(texto: str) -> str:
    normal = unicodedata.normalize("NFD", texto)
    return "".join(c for c in normal if unicodedata.category(c) != "Mn")
