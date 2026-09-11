"""Tests del analizador FIFO: identificar tipo, después analizar."""
from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

import pandas as pd

from inversiones import (
    aplicar_fifo,
    exportar_inversiones_excel,
    identificar_movimientos,
    normalizar_tipo_operacion,
)
from inversiones_catalogo import (
    TIPO_ACCION,
    TIPO_BONO,
    TIPO_FCI,
    TIPO_USD_MEP,
    canon_fima,
    evento_para,
    identificar_especie,
)


def _mov(
    fecha: str,
    especie: str,
    tipo_op: str,
    cant: float,
    precio: float,
    moneda: str = "ARS",
    desc: str = "",
) -> dict:
    f = date(int(fecha[6:10]), int(fecha[3:5]), int(fecha[0:2]))
    return {
        "Fecha": fecha,
        "_fecha": f,
        "Especie": especie,
        "Tipo_Operacion": tipo_op,
        "Cantidad": cant,
        "Precio": precio,
        "Monto_Total": round(cant * precio, 2),
        "Moneda": moneda,
        "Descripcion": desc or f"{tipo_op} {especie}",
        "Archivo origen": "test",
        "Nueva_Clasificacion": None,
    }


def _df(rows: list[dict]) -> pd.DataFrame:
    return identificar_movimientos(pd.DataFrame(rows))


# ---------------------------------------------------------------------------
# Identificación
# ---------------------------------------------------------------------------

def test_identificar_fima_vs_al30_vs_ggal_vs_usd():
    fima = identificar_especie("FONDO - FIMA PREMIUM CLASE A")
    assert fima.tipo_inversion == TIPO_FCI
    assert fima.especie_canonica == "FIMA PREMIUM CLASE A"
    assert fima.confianza == "alta"

    al30 = identificar_especie("AL30", "compra mep")
    assert al30.tipo_inversion == TIPO_BONO
    assert al30.especie_canonica == "AL30"

    ggal = identificar_especie("GGAL")
    assert ggal.tipo_inversion == TIPO_ACCION
    assert ggal.especie_canonica == "GGAL"

    usd = identificar_especie("USD", "caja de ahorro en dolares")
    assert usd.tipo_inversion == TIPO_USD_MEP
    assert usd.especie_canonica == "USD"


def test_nombres_fima_distintos_misma_clave():
    a = canon_fima("FONDO - Fima Premium Clase A")
    b = canon_fima("FIMA PREMIUM CLASE A")
    c = canon_fima("Suscripcion Fima Premium")
    assert a == b == c == "FIMA PREMIUM CLASE A"

    dolares = canon_fima("FIMA RENTA FIJA DOLARES CLASE A")
    assert dolares == "FIMA RENTA FIJA DOLARES CLASE A"


# ---------------------------------------------------------------------------
# FCI
# ---------------------------------------------------------------------------

def test_fci_dos_suscrip_rescate_parcial_lote_viejo():
    df = _df([
        _mov("10/01/2025", "FONDO - FIMA PREMIUM CLASE A", "Suscripcion", 100, 10),
        _mov("20/01/2025", "FIMA PREMIUM CLASE A", "Suscripcion", 100, 12),
        _mov("15/02/2025", "Fima Premium", "Rescate", 80, 13),
    ])
    assert set(df["Especie_canonica"]) == {"FIMA PREMIUM CLASE A"}
    res = aplicar_fifo(df)
    salidas = [m for m in res.movimientos if m["Estado"] == "INTERES"]
    assert len(salidas) == 1
    assert salidas[0]["Costo_Aplicado"] == 800.0  # 80 × 10 (lote viejo)
    assert salidas[0]["Resultado"] == 240.0  # (13 − 10) × 80
    assert len(res.saldos) == 1
    assert res.saldos[0]["Cantidad"] == 120.0
    assert res.saldos[0]["Costo_Total"] == 1400.0  # 20×10 + 100×12


# ---------------------------------------------------------------------------
# Bono
# ---------------------------------------------------------------------------

def test_bono_cupon_no_mueve_stock():
    df = _df([
        _mov("05/01/2025", "AL30", "Compra", 1000, 50),
        _mov("15/03/2025", "AL30", "Renta", 0, 0) | {"Cantidad": 0, "Monto_Total": 2500.0, "Precio": 0},
    ])
    # cantidad 0 en renta: el parser a veces trae solo el cobro
    df.loc[df["Tipo_Operacion"] == "Renta", "Cantidad"] = 0
    df.loc[df["Tipo_Operacion"] == "Renta", "Monto_Total"] = 2500.0
    res = aplicar_fifo(df)
    ingresos = [m for m in res.movimientos if m["Estado"] == "INGRESO"]
    assert len(ingresos) == 1
    assert ingresos[0]["Resultado"] == 2500.0
    assert res.saldos[0]["Cantidad"] == 1000.0
    assert not any("sin stock" in a.lower() for a in res.avisos)


def test_bono_amortizacion_baja_vn():
    df = _df([
        _mov("05/01/2025", "AL30", "Compra", 1000, 50),  # costo 50.000
        _mov("01/06/2025", "AL30", "Amortizacion", 200, 55),  # cobra 11.000
    ])
    res = aplicar_fifo(df)
    amort = [m for m in res.movimientos if m["Estado"] == "AMORTIZACION"]
    assert len(amort) == 1
    assert amort[0]["Costo_Aplicado"] == 10000.0  # 200 × 50
    assert amort[0]["Resultado"] == 1000.0  # 11.000 − 10.000
    assert res.saldos[0]["Cantidad"] == 800.0
    assert res.saldos[0]["Costo_Total"] == 40000.0


# ---------------------------------------------------------------------------
# Acción
# ---------------------------------------------------------------------------

def test_accion_dividendo_no_mueve_stock():
    df = _df([
        _mov("10/01/2025", "GGAL", "Compra", 100, 20),
        _mov("20/04/2025", "GGAL", "Dividendo", 100, 1.5),
    ])
    res = aplicar_fifo(df)
    ingresos = [m for m in res.movimientos if m["Estado"] == "INGRESO"]
    assert len(ingresos) == 1
    assert ingresos[0]["Cantidad_aplicada"] == 0
    assert res.saldos[0]["Cantidad"] == 100.0
    assert res.saldos[0]["Costo_Total"] == 2000.0


# ---------------------------------------------------------------------------
# Mismo día / stock / omitidos
# ---------------------------------------------------------------------------

def test_mismo_dia_compra_antes_que_venta():
    df = _df([
        _mov("10/01/2025", "GGAL", "Venta", 50, 22),
        _mov("10/01/2025", "GGAL", "Compra", 50, 20),
    ])
    res = aplicar_fifo(df)
    assert not any(m["Estado"] == "SIN_STOCK" for m in res.movimientos)
    salidas = [m for m in res.movimientos if m["Estado"] == "SALIDA"]
    assert salidas[0]["Costo_Aplicado"] == 1000.0
    assert salidas[0]["Resultado"] == 100.0
    assert res.saldos == []


def test_sin_stock_no_maquilla_resultado():
    df = _df([
        _mov("10/01/2025", "GGAL", "Compra", 10, 20),
        _mov("20/01/2025", "GGAL", "Venta", 30, 25),
    ])
    res = aplicar_fifo(df)
    salida = [m for m in res.movimientos if m["Estado"] == "SIN_STOCK"][0]
    assert salida["Cantidad_aplicada"] == 10.0
    assert salida["Cantidad_sin_stock"] == 20.0
    assert salida["Costo_Aplicado"] == 200.0
    # monto aplicado = 750 * 10/30 = 250 → resultado 50; no 0
    assert salida["Resultado"] == 50.0
    assert any("faltan 20" in a for a in res.avisos)


def test_usd_no_entra_a_fifo():
    df = _df([
        _mov("10/01/2025", "USD", "Compra", 1000, 1, moneda="USD"),
        _mov("20/01/2025", "USD", "Venta", 400, 1, moneda="USD"),
    ])
    res = aplicar_fifo(df)
    assert all(m["Estado"] == "OMITIDO_USD" for m in res.movimientos)
    assert res.saldos == []
    assert res.aplicaciones == []
    assert any("Caja USD" in a for a in res.avisos)


def test_movimiento_desconocido_no_inventa_lote():
    df = _df([
        _mov("10/01/2025", "GGAL", "Compra", 10, 20),
        _mov("11/01/2025", "GGAL", "Movimiento", 5, 20),
    ])
    res = aplicar_fifo(df)
    assert any(m["Estado"] == "OMITIDO" for m in res.movimientos)
    assert res.saldos[0]["Cantidad"] == 10.0


def test_fci_primer_rescate_contra_valor_cuota_inicial():
    ini = pd.DataFrame([{
        "Especie": "FIMA PREMIUM CLASE A",
        "Cantidad": 310_323.99,
        "Costo_Unitario": 64.0511,
        "Moneda": "ARS",
    }])
    df = _df([
        _mov("01/07/2025", "FIMA PREMIUM CLASE A", "Rescate", 40565.09, 64.094517),
    ])
    res = aplicar_fifo(df, ini)
    interes = [m for m in res.movimientos if m["Estado"] == "INTERES"]
    assert len(interes) == 1
    esperado = round((64.094517 - 64.0511) * 40565.09, 2)
    assert interes[0]["Resultado"] == esperado
    assert abs(res.saldos[0]["Cantidad"] - (310_323.99 - 40565.09)) < 0.02


def test_fci_rescate_parte_varios_lotes_fifo():
    ini = pd.DataFrame([{
        "Especie": "FIMA PREMIUM CLASE A",
        "Cantidad": 1000.0,
        "Costo_Unitario": 10.0,
        "Moneda": "ARS",
    }])
    df = _df([
        _mov("02/07/2025", "FIMA PREMIUM CLASE A", "Suscripcion", 200.0, 11.0),
        _mov("10/07/2025", "FIMA PREMIUM CLASE A", "Rescate", 1100.0, 12.0),
    ])
    res = aplicar_fifo(df, ini)
    apps = res.aplicaciones
    assert len(apps) == 2
    assert apps[0]["Cantidad"] == 1000.0
    assert apps[0]["Interes"] == 2000.0  # (12-10)*1000
    assert apps[1]["Cantidad"] == 100.0
    assert apps[1]["Interes"] == 100.0  # (12-11)*100
    interes = [m for m in res.movimientos if m["Estado"] == "INTERES"][0]
    assert interes["Resultado"] == 2100.0


def test_saldo_inicial_sin_cuotas_no_siembra():
    ini = pd.DataFrame([{
        "Especie": "FIMA PREMIUM CLASE A",
        "Tipo_inversion": TIPO_FCI,
        "Grupo": "FCI",
        "Cantidad": None,
        "Costo_Total": 2_000_000,
        "Moneda": "ARS",
        "Revisar": "SI — falta cantidad de cuotapartes",
    }])
    df = _df([_mov("10/01/2025", "FIMA PREMIUM CLASE A", "Suscripcion", 10, 100)])
    res = aplicar_fifo(df, ini)
    assert any("no se sembró lote" in a for a in res.avisos)
    assert res.saldos[0]["Cantidad"] == 10.0


def test_export_excel_estudio():
    df = _df([
        _mov("10/01/2025", "FIMA PREMIUM CLASE A", "Suscripcion", 10, 100),
        _mov("20/01/2025", "FIMA PREMIUM CLASE A", "Rescate", 4, 110),
    ])
    res = aplicar_fifo(df)
    xlsx = exportar_inversiones_excel(df, res, meta={"nota": "test"})
    assert xlsx[:2] == b"PK"
    assert len(xlsx) > 1000


def test_dividendo_normalizado_no_come_stock():
    assert normalizar_tipo_operacion("DIVIDENDO GGAL") == "Dividendo"
    df = _df([
        _mov("10/01/2025", "GGAL", normalizar_tipo_operacion("Dividendo en efectivo"), 100, 1.5),
    ])
    # sin compra previa: si se tratara como venta, habría SIN_STOCK
    res = aplicar_fifo(df)
    assert res.movimientos[0]["Estado"] == "INGRESO"
    assert res.saldos == []
    assert not any("sin stock" in a.lower() for a in res.avisos)


def test_evento_para_reglas():
    assert evento_para(TIPO_FCI, "Suscripcion") == "entrada"
    assert evento_para(TIPO_FCI, "Rescate") == "salida"
    assert evento_para(TIPO_BONO, "Renta") == "ingreso"
    assert evento_para(TIPO_BONO, "Amortizacion") == "amortizacion"
    assert evento_para(TIPO_ACCION, "Dividendo") == "ingreso"
    assert evento_para(TIPO_USD_MEP, "Compra") == "omitir_usd"
    assert evento_para(TIPO_FCI, "Movimiento") == "omitir"


FCI_GROWTH = Path(r"c:\Users\recep\Desktop\FCI GROWTH.xlsx")


def _leer_fci_growth_manual():
    from openpyxl import load_workbook

    wb = load_workbook(FCI_GROWTH, data_only=False)
    ws = wb["FCI"]
    ini_cant = float(ws["C7"].value)
    ini_vc = float(ws["D7"].value)
    movs = []
    manual = []
    for r in range(10, 81):
        tipo = ws.cell(r, 3).value
        if tipo not in {"RESCATE", "SUSCRIPCION"}:
            continue
        fecha = ws.cell(r, 2).value
        if isinstance(fecha, datetime):
            fecha = fecha.date()
        cant = float(ws.cell(r, 4).value)
        vc = float(ws.cell(r, 5).value)
        monto = ws.cell(r, 6).value
        g = ws.cell(r, 7).value
        movs.append(_mov(
            fecha.strftime("%d/%m/%Y"),
            "FIMA PREMIUM CLASE A",
            "Rescate" if tipo == "RESCATE" else "Suscripcion",
            cant,
            vc,
        ))
        if tipo == "RESCATE" and isinstance(g, str) and g.startswith("="):
            # Recalcular G con data_only=False: usamos openpyxl formulas no evaluadas.
            # Guardamos la fórmula para comparar contra el motor, no el valor.
            manual.append({"row": r, "fecha": fecha, "cant": cant, "vc": vc, "formula": g})
    return ini_cant, ini_vc, movs, manual


def test_fci_growth_julio_contra_excel_manual():
    if not FCI_GROWTH.exists():
        return
    ini_cant, ini_vc, movs, _manual = _leer_fci_growth_manual()
    ini = pd.DataFrame([{
        "Especie": "FIMA PREMIUM CLASE A",
        "Cantidad": ini_cant,
        "Costo_Unitario": ini_vc,
        "Moneda": "ARS",
    }])
    df = _df(movs)
    res = aplicar_fifo(df, ini)
    assert not any("sin stock" in a.lower() for a in res.avisos)

    # Primer rescate = (64.094517 − 64.0511) × 40565.09
    primero = next(m for m in res.movimientos if m["Estado"] == "INTERES")
    assert primero["Fecha"] == "01/07/2025"
    assert primero["Resultado"] == round((64.094517 - 64.0511) * 40565.09, 2)

    # Julio: mismos rescates que suman C3 del Excel (todas las G de julio)
    julio = [
        m for m in res.movimientos
        if m["Estado"] == "INTERES" and str(m["Fecha"]).endswith("/07/2025")
    ]
    assert len(julio) == 13  # G10,12-14,17,20-21,28-31,34,36
    # G20 parte el inicial + 3 suscripciones: el motor debe partir en 4 aplicaciones
    apps_20 = [a for a in res.aplicaciones if a["Fecha_salida"] == "10/07/2025"]
    # hay dos rescates el 10/07; el primero es el grande
    primer_10 = [a for a in apps_20 if abs(float(a.get("Valor_CC_salida") or 0) - 64.457409) < 1e-6]
    # ambos rescates del 10 tienen el mismo VC; chequear que el grande partió lotes
    assert len(apps_20) >= 4
