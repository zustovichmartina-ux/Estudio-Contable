#!/usr/bin/env python3
"""Emisor de facturas electrónicas ARCA (WSFEv1) desde Excel - Estudio Zona Güemes.

Ejemplos:
  python emitir.py --excel Facturas.xlsx --dry-run                     # valida, no envía nada
  python emitir.py --excel Facturas.xlsx --dry-run --incluir-ejemplos --pdf-borrador
  python emitir.py --dummy --env homo                                   # prueba conectividad
  python emitir.py --excel Facturas.xlsx --env homo                     # emite en HOMOLOGACIÓN
  python emitir.py --excel Facturas.xlsx --env prod --confirmar-produccion
"""
import argparse
import datetime as dt
import json
import os
import shutil
import sys
from decimal import Decimal

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from arca.codigos import NOMBRE_CBTE, TIPOS_CBTE
from arca.comprobante import Comprobante, agrupar, leer_planilla
from arca.pdf import generar_pdf
from arca.ta_store import obtener_ta_durable
from arca.wsfe import WSFE, xml_fecaesolicitar

BASE = os.path.dirname(os.path.abspath(__file__))

CUIT_ESTUDIO = "23422843434"
RES_COLS = ["Fecha/hora proceso", "Ambiente", "Modo", "Filas Facturas", "ID factura", "CUIT emisor",
            "Pto. vta", "Tipo", "Número", "CAE", "Vto. CAE", "Estado", "Errores / observaciones",
            "Total", "PDF", "Huella", "Fecha cbte", "Cbte asociado"]
COLOR = {"APROBADO": "C6EFCE", "OK-DRYRUN": "DDEBF7", "RECHAZADO": "FFC7CE", "ERROR": "FFC7CE",
         "VERIFICAR": "FFEB9C", "YA EMITIDO": "EDEDED", "OMITIDO": "EDEDED"}


class EmisionAbortada(Exception):
    """El lote no se envía (producción sin confirmación, o filas EJEMPLO hacia producción)."""


def hoja_resultado(wb):
    if "Resultado" in wb.sheetnames:
        ws = wb["Resultado"]
    else:
        ws = wb.create_sheet("Resultado")
        ws.append(RES_COLS)
        for i, h in enumerate(RES_COLS, 1):
            cell = ws.cell(1, i)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[cell.column_letter].width = [18, 9, 9, 10, 10, 13, 7, 7, 9, 16, 11, 14, 60, 13, 40, 18, 11, 14][i - 1]
        ws.freeze_panes = "A2"
    for i, h in enumerate(RES_COLS, 1):  # planillas de versiones anteriores: agregar columnas nuevas
        if ws.cell(1, i).value in (None, ""):
            ws.cell(1, i, h).font = Font(bold=True, color="FFFFFF")
            ws.cell(1, i).fill = PatternFill("solid", fgColor="1F4E78")
    ws.sheet_state = "visible"
    return ws


def ya_emitidos(ws, env):
    res = {}
    for r in range(2, ws.max_row + 1):
        h, amb, est = ws.cell(r, 16).value, ws.cell(r, 2).value, str(ws.cell(r, 12).value or "")
        if h and amb == env and (est.startswith("APROBADO") or est.startswith("VERIFICAR")):
            res[h] = r
    return res


def historial(ws, env):
    """Comprobantes aprobados en corridas anteriores (por ID) y NC ya emitidas por comprobante asociado."""
    por_id, acreditado = {}, {}
    for r in range(2, ws.max_row + 1):
        v = [ws.cell(r, c).value for c in range(1, len(RES_COLS) + 1)]
        if v[1] != env or not str(v[11] or "").startswith("APROBADO") or not v[8]:
            continue
        tipo = TIPOS_CBTE.get(str(v[7] or "").strip(), (None,))[0]
        if v[4]:
            fch = None
            if v[16]:
                fch = dt.datetime.strptime(str(v[16]), "%Y%m%d").date()
            por_id[(str(v[5]), str(v[4]).strip())] = {"tipo": tipo, "pto": int(v[6]), "nro": int(v[8]),
                                                      "fecha": fch, "total": v[13]}
        if str(v[7] or "").startswith("NC") and v[17]:
            acreditado[v[17]] = acreditado.get(v[17], 0) + float(v[13] or 0)
    return por_id, acreditado


def escribir(ws, fila):
    ws.append(fila)
    r = ws.max_row
    est = str(fila[11] or "")
    for k, col in COLOR.items():
        if est.startswith(k):
            for c in range(1, len(RES_COLS) + 1):
                ws.cell(r, c).fill = PatternFill("solid", fgColor=col)
    ws.cell(r, 13).alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(r, 14).number_format = '#,##0.00'


def guardar(wb, path):
    tmp = path + ".tmp.xlsx"
    wb.save(tmp)
    os.replace(tmp, path)


def log(env, evento, base=None):
    raiz = base or BASE
    os.makedirs(os.path.join(raiz, "salida", "log"), exist_ok=True)
    with open(os.path.join(raiz, "salida", "log", f"emision_{env}_{dt.date.today():%Y%m}.jsonl"), "a") as f:
        f.write(json.dumps({"ts": dt.datetime.now().isoformat(timespec="seconds"), **evento},
                           ensure_ascii=False, default=str) + "\n")


def registrar(cb, nro, por_id_run, acreditado):
    idf = str(cb.f0.get("id") or "").strip()
    if idf:
        por_id_run[(cb.cuit, idf)] = {"tipo": cb.cbte_tipo, "pto": cb.pto_vta, "nro": nro,
                                      "fecha": cb.fecha, "total": cb.det["ImpTotal"]}
    if cb.clase_doc == "NC" and cb.clave_asoc() and nro is not None:
        acreditado[cb.clave_asoc()] = acreditado.get(cb.clave_asoc(), 0) + cb.det["ImpTotal"]


def resolver_asociado(cb, por_id_run, por_id_prev, acreditado, online):
    """Completa cb.asoc (por ID del Excel o datos cargados), verifica contra ARCA si hay conexión
    y aplica las reglas (misma letra, NC <= disponible). Devuelve True si se puede emitir."""
    a = cb.asoc
    total_orig = None
    if "id" in a:
        ref = por_id_run.get((cb.cuit, a["id"])) or por_id_prev.get((cb.cuit, a["id"]))
        if not ref:
            cb.errores.append(f"Cbte asociado: el ID factura '{a['id']}' no está emitido (ni en esta corrida ni "
                              "APROBADO en la hoja Resultado) para este CUIT y ambiente.")
            return False
        cb.asoc = a = {"tipo": ref["tipo"], "pto": ref["pto"], "nro": ref["nro"], "fecha": ref["fecha"],
                       "cuit": cb.cuit, "id": a["id"]}
        total_orig = ref["total"]
    if online and a.get("nro"):
        wsfe, get_auth = online
        try:
            cons, _ = wsfe.consultar(get_auth(), a["pto"], a["tipo"], a["nro"])
        except Exception as ex:
            cb.errores.append(f"No se pudo consultar el comprobante asociado en ARCA: {ex}")
            return False
        if not cons:
            cb.errores.append(f"El comprobante asociado {a['tipo']}/{a['pto']}/{a['nro']} no existe en ARCA "
                              "para este emisor.")
            return False
        total_orig = float(cons.get("ImpTotal") or 0)
        if cons.get("CbteFch"):
            a["fecha"] = dt.datetime.strptime(cons["CbteFch"], "%Y%m%d").date()
        if str(cons.get("DocNro")) != str(cb.det["DocNro"]) or str(cons.get("DocTipo")) != str(cb.det["DocTipo"]):
            cb.avisos.append(f"El receptor difiere del comprobante asociado (ARCA: {cons.get('DocTipo')}/"
                             f"{cons.get('DocNro')}).")
    elif total_orig is None and not online:
        cb.avisos.append("Comprobante asociado externo: existencia e importe se verifican al enviar (FECompConsultar).")
    clave = f"{a['tipo']}-{a['pto']}-{a['nro']}" if a.get("nro") else None
    previo = Decimal(str(acreditado.get(clave, 0))) if clave else Decimal(0)
    return cb.verificar_asociado(total_orig, previo) and not cb.errores


def nombre_pdf(cb, nro):
    return f"{cb.cuit}_{cb.cbte_tipo:03d}_{cb.pto_vta:05d}-{nro:08d}.pdf"


def chequear_ambiente(a):
    if a.env == "prod" and not a.dry_run and not getattr(a, "confirmar_produccion", False) and not getattr(a, "dummy", False):
        raise EmisionAbortada("Producción requiere --confirmar-produccion (o usar --dry-run).")
    if getattr(a, "incluir_ejemplos", False) and a.env == "prod" and not a.dry_run:
        raise EmisionAbortada("Las filas EJEMPLO nunca se envían a producción.")


def correr(a, on_resultado=None, base_dir=None):
    """Misma corrida que el CLI. `on_resultado(fila, cb, pdf_path)` se llama por cada fila escrita."""
    chequear_ambiente(a)
    raiz = base_dir or BASE
    cert = a.cert or os.path.join(raiz, "cert", f"estudiozg-{a.env}.crt")
    a.cache = a.cache or os.path.join(raiz, "cache", a.env)
    wsfe = WSFE(a.env)

    if a.dummy:
        print(a.env, wsfe.dummy())
        return {"resumen": {}, "dummy": wsfe.dummy()}

    def auth_para(cuit):
        ta = obtener_ta_durable(a.env, cert, a.key, "wsfe", a.cache)
        return {"token": ta["token"], "sign": ta["sign"], "cuit": cuit}, ta

    if a.parametros:
        auth, ta = auth_para(a.parametros)
        print(f"TA ({ta['source']}) vence {ta['expiration']}")
        for met in ("FEParamGetCondicionIvaReceptor", "FEParamGetTiposCbte",
                    "FEParamGetTiposIva", "FEParamGetTiposDoc", "FEParamGetPtosVenta"):
            items, msgs = wsfe.param(auth, met, "")
            print(f"\n== {met}")
            for it in items:
                print("  ", it)
            for m in msgs:
                print("   [", *m, "]")
        return {"resumen": {}, "ta": ta}

    if not a.excel:
        raise EmisionAbortada("--excel es obligatorio")

    def anotar(fila, cb=None, pdf_path=None):
        escribir(ws_res, fila)
        if on_resultado:
            on_resultado(fila, cb, pdf_path)

    wb_vals = openpyxl.load_workbook(a.excel, data_only=True)
    wb = openpyxl.load_workbook(a.excel)
    filas = leer_planilla(wb_vals)
    if a.solo_filas:
        sel = {int(x) for x in a.solo_filas.split(",")}
        filas = [f for f in filas if f["_fila"] in sel]
    ws_res = hoja_resultado(wb)
    previos = ya_emitidos(ws_res, a.env)
    por_id_prev, acreditado = historial(ws_res, a.env)
    por_id_run = {}

    if not a.dry_run:
        os.makedirs(os.path.join(raiz, "salida", "backups"), exist_ok=True)
        shutil.copy2(a.excel, os.path.join(raiz, "salida", "backups",
                                           f"{os.path.basename(a.excel)}.{dt.datetime.now():%Y%m%d_%H%M%S}.bak"))
    os.makedirs(a.pdf_dir, exist_ok=True)
    modo = "DRY-RUN" if a.dry_run else "ENVÍO"
    ahora = dt.datetime.now().strftime("%d/%m/%Y %H:%M:%S")
    resumen = {}

    grupos = agrupar(filas)
    # Primero facturas/comprobantes sin referencia a otra fila; después NC/ND que referencian un ID del Excel
    grupos.sort(key=lambda g: 1 if str(g[0].get("asoc_id") or "").strip() else 0)
    for grupo in grupos:
        cb = Comprobante(grupo)
        base = [ahora, a.env, modo, cb.filas_txt, cb.f0.get("id"), cb.f0.get("cuit_emisor"),
                cb.f0.get("pto_vta"), cb.f0.get("tipo")]
        if cb.es_ejemplo and not a.incluir_ejemplos:
            continue
        ok = cb.validar()
        if ok and cb.asoc:
            ok = resolver_asociado(cb, por_id_run, por_id_prev, acreditado,
                                   None if a.dry_run else (wsfe, lambda: auth_para(cb.cuit)[0]))
        huella = cb.huella()
        total = cb.det["ImpTotal"] if cb.det else None
        msgs = " | ".join(cb.errores + [f"Aviso: {x}" for x in cb.avisos])
        if not ok:
            anotar(base + [None, None, None, "RECHAZADO-VALIDACIÓN", msgs, total, None, huella,
                           None, cb.clave_asoc()], cb, None)
            resumen["RECHAZADO-VALIDACIÓN"] = resumen.get("RECHAZADO-VALIDACIÓN", 0) + 1
            print(f"[filas {cb.filas_txt}] RECHAZADO-VALIDACIÓN: {msgs}")
            continue
        if huella in previos:
            anotar(base + [None, None, None, "YA EMITIDO",
                           f"Mismo contenido ya aprobado (Resultado fila {previos[huella]}). No se reenvía.",
                           total, None, huella], cb, None)
            print(f"[filas {cb.filas_txt}] YA EMITIDO, se omite")
            resumen["YA EMITIDO"] = resumen.get("YA EMITIDO", 0) + 1
            continue

        if a.dry_run:
            os.makedirs(os.path.join(raiz, "salida", "dryrun"), exist_ok=True)
            xmlp = os.path.join(raiz, "salida", "dryrun", f"filas_{cb.filas_txt.replace(',', '-')}.xml")
            with open(xmlp, "w") as f:
                f.write(xml_fecaesolicitar({"token": "TOKEN", "sign": "SIGN", "cuit": cb.cuit},
                                           cb.pto_vta, cb.cbte_tipo, cb.con_numero(0)))
            pdfp = None
            if a.pdf_borrador:
                pdfp = os.path.join(a.pdf_dir, f"BORRADOR_filas{cb.filas_txt.replace(',', '-')}_" + nombre_pdf(cb, 0))
                generar_pdf(pdfp, cb, 0, "00000000000000", None,
                            marca_agua="EJEMPLO - SIN VALIDEZ" if cb.es_ejemplo else "BORRADOR - SIN CAE")
            anotar(base + [None, None, None, "OK-DRYRUN", msgs or "Validación OK (no se envió a ARCA)",
                           total, pdfp, huella, cb.det["CbteFch"], cb.clave_asoc()], cb, pdfp)
            registrar(cb, None, por_id_run, acreditado)
            resumen["OK-DRYRUN"] = resumen.get("OK-DRYRUN", 0) + 1
            print(f"[filas {cb.filas_txt}] OK-DRYRUN {NOMBRE_CBTE[cb.cbte_tipo]} {cb.clase} total {total:.2f} -> {xmlp}")
            continue

        # ---- Envío real (homologación o producción) ----
        nro = None
        auth = None
        try:
            auth, ta = auth_para(cb.cuit)
            nro = wsfe.ultimo_autorizado(auth, cb.pto_vta, cb.cbte_tipo) + 1
            det = cb.con_numero(nro)
            log(a.env, {"op": "FECAESolicitar", "cuit": cb.cuit, "pto": cb.pto_vta, "tipo": cb.cbte_tipo,
                        "nro": nro, "det": det}, base=raiz)
            r = wsfe.solicitar_cae(auth, cb.pto_vta, cb.cbte_tipo, det)
            log(a.env, {"op": "respuesta", "nro": nro, "r": r}, base=raiz)
            txt = " | ".join(f"{t} {c}: {m}" for t, c, m in r["mensajes"])
            if msgs:
                txt = (txt + " | " if txt else "") + msgs
            if r["Resultado"] == "A" and r["CAE"]:
                pdfp = os.path.join(a.pdf_dir, nombre_pdf(cb, nro))
                generar_pdf(pdfp, cb, nro, r["CAE"], r["CAEFchVto"],
                            marca_agua="HOMOLOGACIÓN - SIN VALIDEZ" if a.env == "homo" else None)
                estado = "APROBADO" + (" c/observaciones" if any(t == "Observaciones" for t, _, _ in r["mensajes"]) else "")
                anotar(base + [nro, r["CAE"], r["CAEFchVto"], estado, txt, total, pdfp, huella,
                               det["CbteFch"], cb.clave_asoc()], cb, pdfp)
                registrar(cb, nro, por_id_run, acreditado)
            else:
                estado = "RECHAZADO"
                anotar(base + [None, None, None, estado, txt, total, None, huella,
                               det["CbteFch"], cb.clave_asoc()], cb, None)
        except Exception as ex:  # red / timeout / WSAA: verificar si quedó emitido
            estado, txt, cae, vto, pdfp = "ERROR", f"{type(ex).__name__}: {ex}", None, None, None
            if nro and auth is not None:
                try:
                    if wsfe.ultimo_autorizado(auth, cb.pto_vta, cb.cbte_tipo) == nro:
                        cons, _ = wsfe.consultar(auth, cb.pto_vta, cb.cbte_tipo, nro)
                        cae, vto = (cons or {}).get("CodAutorizacion"), (cons or {}).get("FchVto")
                        estado = "VERIFICAR"
                        txt += f" | El nro {nro} figura autorizado en ARCA (CAE {cae}); revisar que corresponda."
                        if cae:
                            pdfp = os.path.join(a.pdf_dir, nombre_pdf(cb, nro))
                            generar_pdf(pdfp, cb, nro, cae, vto,
                                        marca_agua="HOMOLOGACIÓN - SIN VALIDEZ" if a.env == "homo" else None)
                except Exception as ex2:
                    txt += f" | No se pudo verificar: {ex2}"
            anotar(base + [nro if estado == "VERIFICAR" else None, cae, vto, estado, txt, total,
                           pdfp, huella, cb.det["CbteFch"], cb.clave_asoc()], cb, pdfp)
        resumen[estado] = resumen.get(estado, 0) + 1
        print(f"[filas {cb.filas_txt}] {estado} nro={nro} {txt[:300]}")
        guardar(wb, a.excel)  # se guarda después de cada comprobante enviado

    guardar(wb, a.excel)
    print("Resumen:", resumen or "sin comprobantes a procesar")
    return {"resumen": resumen}


def build_parser():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--excel")
    ap.add_argument("--env", choices=["homo", "prod"], default="homo")
    ap.add_argument("--dry-run", action="store_true", help="Valida y arma los XML sin conectarse a ARCA")
    ap.add_argument("--confirmar-produccion", action="store_true")
    ap.add_argument("--incluir-ejemplos", action="store_true", help="Procesa filas marcadas EJEMPLO (nunca en prod)")
    ap.add_argument("--pdf-borrador", action="store_true", help="En dry-run genera PDF de vista previa (sin CAE)")
    ap.add_argument("--solo-filas", help="Lista de filas Excel a procesar, ej: 5,6,9")
    ap.add_argument("--cert", help="Certificado .crt de ARCA (default cert/estudiozg-<env>.crt)")
    ap.add_argument("--key", default=os.path.join(BASE, "cert", "estudiozg.key"))
    ap.add_argument("--cache", help="Directorio de cache del TA (default cache/<env>)")
    ap.add_argument("--pdf-dir", default=os.path.join(BASE, "salida", "pdf"))
    ap.add_argument("--dummy", action="store_true", help="FEDummy (estado de servidores, sin autenticación)")
    ap.add_argument("--parametros", metavar="CUIT", help="Consulta tablas FEParamGet* para esa CUIT representada")
    return ap


def main():
    ap = build_parser()
    a = ap.parse_args()
    if not a.dummy and not a.parametros and not a.excel:
        ap.error("--excel es obligatorio")
    try:
        correr(a)
    except EmisionAbortada as ex:
        sys.exit(str(ex))


if __name__ == "__main__":
    main()
