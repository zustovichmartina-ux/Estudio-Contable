"""Excel interactivo: contraparte, detalle y saldo encadenado."""

from __future__ import annotations

from datetime import date
from io import BytesIO
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from excel_formato_estudio import (
    BODY_FONT,
    BOLD_FONT,
    COLOR_PRIMARIO,
    DATE_FMT,
    HDR_FILL,
    HDR_FONT,
    MONEY_FMT,
    SECTION_FONT,
    SUB_FONT,
    TITLE_FONT,
    ZEBRA,
)
from gastos_mp import CATEGORIAS, CATEGORIAS_NO_GASTO, COLS_MOV

FILL_KPI = PatternFill("solid", fgColor="E8F0F7")
FILL_AVISO = PatternFill("solid", fgColor="FFF8E7")
FONT_KPI_VAL = Font(name="Calibri", bold=True, size=16, color=COLOR_PRIMARIO)
THIN = Border(
    left=Side(style="thin", color="D0D7DE"),
    right=Side(style="thin", color="D0D7DE"),
    top=Side(style="thin", color="D0D7DE"),
    bottom=Side(style="thin", color="D0D7DE"),
)


def _cats_gasto() -> list[str]:
    return [c for c in CATEGORIAS if c not in CATEGORIAS_NO_GASTO]


def _construir_wb_gastos(
    df: pd.DataFrame,
    *,
    subtitulo: str = "Mercado Pago · gastos personales",
    periodo: str = "",
) -> Workbook:
    data = df.copy()
    if data.empty:
        data = pd.DataFrame(columns=COLS_MOV)
    for c in COLS_MOV:
        if c not in data.columns:
            data[c] = None
    data = data[COLS_MOV]
    n = len(data)
    last = max(2, n + 1)
    n_cols = len(COLS_MOV)

    wb = Workbook()
    ws_cat = wb.active
    ws_cat.title = "Categorias"
    for i, cat in enumerate(CATEGORIAS, 1):
        ws_cat.cell(i, 1, cat).font = BODY_FONT
    ws_cat.sheet_state = "hidden"

    ws_m = wb.create_sheet("Movimientos", 0)
    for c, h in enumerate(COLS_MOV, 1):
        cell = ws_m.cell(1, c, h)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = Alignment(horizontal="center")

    for i, row in enumerate(data.itertuples(index=False), 2):
        valores = {h: getattr(row, h) for h in COLS_MOV}
        for c, h in enumerate(COLS_MOV, 1):
            cell = ws_m.cell(i, c, valores[h])
            cell.font = BODY_FONT
            cell.border = THIN
            if h == "Fecha" and valores[h]:
                cell.number_format = DATE_FMT
            if h in {"Importe", "Saldo"}:
                v = valores[h]
                if v is None or (isinstance(v, float) and pd.isna(v)):
                    cell.value = None
                else:
                    cell.value = float(v)
                cell.number_format = MONEY_FMT
            if i % 2 == 0:
                cell.fill = ZEBRA

    extra = last + 8
    for r in range(last + 1, extra + 1):
        for c in range(1, n_cols + 1):
            ws_m.cell(r, c).border = THIN
            if COLS_MOV[c - 1] in {"Importe", "Saldo"}:
                ws_m.cell(r, c).number_format = MONEY_FMT

    tabla_fin = extra
    tab = Table(displayName="Movimientos", ref=f"A1:{chr(64 + n_cols)}{tabla_fin}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
    ws_m.add_table(tab)
    ws_m.freeze_panes = "A2"

    anchos = {
        "A": 14,
        "B": 12,
        "C": 8,
        "D": 32,
        "E": 42,
        "F": 26,
        "G": 14,
        "H": 14,
        "I": 16,
        "J": 16,
    }
    for col, w in anchos.items():
        ws_m.column_dimensions[col].width = w

    col_cat = COLS_MOV.index("Categoria") + 1
    dv = DataValidation(
        type="list",
        formula1=f"=Categorias!$A$1:$A${len(CATEGORIAS)}",
        allow_blank=True,
        showDropDown=False,
        showErrorMessage=True,
        errorTitle="Categoría",
        error="Elegí una categoría de la lista.",
    )
    letra_cat = chr(64 + col_cat)
    dv.add(f"{letra_cat}2:{letra_cat}{tabla_fin}")
    ws_m.add_data_validation(dv)

    rojo = Font(name="Calibri", color="C0392B")
    verde = Font(name="Calibri", color="1E8449")
    col_imp = COLS_MOV.index("Importe") + 1
    letra_imp = chr(64 + col_imp)
    ws_m.conditional_formatting.add(
        f"{letra_imp}2:{letra_imp}{tabla_fin}",
        CellIsRule(operator="lessThan", formula=["0"], font=rojo),
    )
    ws_m.conditional_formatting.add(
        f"{letra_imp}2:{letra_imp}{tabla_fin}",
        CellIsRule(operator="greaterThan", formula=["0"], font=verde),
    )

    ws = wb.create_sheet("Tablero", 0)
    ws["A1"] = "Gastos personales"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:F1")
    ws["A2"] = subtitulo
    ws["A2"].font = SUB_FONT
    ws["A3"] = periodo or f"Actualizado {date.today().strftime('%d/%m/%Y')}"
    ws["A3"].font = SUB_FONT
    ws.merge_cells("A4:H4")
    ws["A4"] = (
        "Contraparte = a quién le pagaste o de quién cobraste. "
        "Saldo = dinero en Mercado Pago después de ese movimiento. "
        "Cambiá Categoría en Movimientos: el tablero se recalcula."
    )
    ws["A4"].font = Font(name="Calibri", italic=True, size=11, color="1F4E79")
    ws["A4"].fill = FILL_AVISO
    ws["A4"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[4].height = 36

    kpis_def = [
        ("A6", "B6", "Saldo actual", "=IFERROR(INDEX(Movimientos[Saldo],1),0)"),
        ("C6", "D6", "Gastado", '=ABS(SUMIF(Movimientos[Importe],"<0"))-ABS(SUMIFS(Movimientos[Importe],Movimientos[Categoria],"Retiro a banco",Movimientos[Importe],"<0"))'),
        ("E6", "F6", "Ingresos", '=SUMIF(Movimientos[Importe],">0")'),
        ("A8", "B8", "Movimientos", "=COUNTA(Movimientos[Id])"),
        ("C8", "D8", "Sin clasificar", '=COUNTIF(Movimientos[Categoria],"Sin clasificar")'),
        ("E8", "F8", "Neto del período", "=F6-D6"),
    ]
    for lab_cell, val_cell, label, formula in kpis_def:
        ws[lab_cell] = label
        ws[lab_cell].font = BOLD_FONT
        ws[lab_cell].fill = FILL_KPI
        ws[val_cell] = formula
        ws[val_cell].font = FONT_KPI_VAL
        ws[val_cell].fill = FILL_KPI
        if label not in {"Movimientos", "Sin clasificar"}:
            ws[val_cell].number_format = MONEY_FMT
        ws[lab_cell].border = THIN
        ws[val_cell].border = THIN

    ws["A10"] = "Por categoría"
    ws["A10"].font = SECTION_FONT
    ws["A11"] = "Categoria"
    ws["B11"] = "Gastado"
    ws["C11"] = "% del total"
    for c in range(1, 4):
        ws.cell(11, c).fill = HDR_FILL
        ws.cell(11, c).font = HDR_FONT

    cats_g = _cats_gasto()
    first_cat = 12
    last_cat = first_cat + len(cats_g) - 1
    for i, cat in enumerate(cats_g):
        r = first_cat + i
        ws.cell(r, 1, cat).font = BODY_FONT
        ws.cell(
            r,
            2,
            f'=ABS(SUMIFS(Movimientos[Importe],Movimientos[Categoria],A{r},Movimientos[Importe],"<0"))',
        )
        ws.cell(r, 2).number_format = MONEY_FMT
        ws.cell(r, 3, f'=IF($D$6=0,0,B{r}/$D$6)')
        ws.cell(r, 3).number_format = "0.0%"
        if i % 2:
            for c in range(1, 4):
                ws.cell(r, c).fill = ZEBRA
        for c in range(1, 4):
            ws.cell(r, c).border = THIN
            ws.cell(r, c).font = BODY_FONT

    total_r = last_cat + 1
    ws.cell(total_r, 1, "TOTAL").font = BOLD_FONT
    ws.cell(total_r, 2, f"=SUM(B{first_cat}:B{last_cat})").font = BOLD_FONT
    ws.cell(total_r, 2).number_format = MONEY_FMT
    ws.cell(total_r, 3, f"=SUM(C{first_cat}:C{last_cat})").font = BOLD_FONT
    ws.cell(total_r, 3).number_format = "0.0%"

    bar = BarChart()
    bar.type = "bar"
    bar.title = "Gastos por categoría"
    bar.x_axis.title = "Importe (ARS)"
    bar.style = 10
    bar.legend = None
    data_ref = Reference(ws, min_col=2, min_row=11, max_row=last_cat)
    cats_ref = Reference(ws, min_col=1, min_row=first_cat, max_row=last_cat)
    bar.add_data(data_ref, titles_from_data=True)
    bar.set_categories(cats_ref)
    bar.width = 18
    bar.height = 12
    if bar.series:
        try:
            bar.series[0].graphicalProperties.solidFill = COLOR_PRIMARIO
        except Exception:
            pass
    ws.add_chart(bar, "E11")

    pie = PieChart()
    pie.title = "Composición de gastos"
    pie.add_data(data_ref, titles_from_data=True)
    pie.set_categories(cats_ref)
    pie.dataLabels = DataLabelList()
    pie.dataLabels.showPercent = True
    pie.dataLabels.showVal = False
    pie.dataLabels.showCatName = False
    pie.width = 14
    pie.height = 10
    ws.add_chart(pie, "E32")

    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 16
    ws.column_dimensions["E"].width = 16
    ws.column_dimensions["F"].width = 18
    ws.page_setup.orientation = "landscape"
    ws.sheet_view.showGridLines = False

    ws_u = wb.create_sheet("Como usar")
    lineas = [
        ("Cómo usar este Excel", TITLE_FONT),
        ("", BODY_FONT),
        ("Contraparte", SECTION_FONT),
        ("Es a quién le transferiste, el comercio donde pagaste, o de quién cobraste. El Detalle junta CVU/CUIT, concepto y medio.", BODY_FONT),
        ("", BODY_FONT),
        ("Saldo", SECTION_FONT),
        ("Cada fila muestra el dinero en Mercado Pago DESPUÉS de ese movimiento. Arriba (más reciente) es el saldo actual.", BODY_FONT),
        ("", BODY_FONT),
        ("Categoría", SECTION_FONT),
        ("Menú desplegable. No uses gasto/ingreso/traspaso: elegí Super, Delivery, Alquiler, Nafta, etc.", BODY_FONT),
    ]
    for i, (txt, font) in enumerate(lineas, 1):
        ws_u.cell(i, 1, txt).font = font
    ws_u.column_dimensions["A"].width = 110
    return wb


def exportar_excel_gastos(
    df: pd.DataFrame,
    ruta: str | Path,
    *,
    subtitulo: str = "Mercado Pago · gastos personales",
    periodo: str = "",
) -> Path:
    ruta = Path(ruta)
    ruta.parent.mkdir(parents=True, exist_ok=True)
    wb = _construir_wb_gastos(df, subtitulo=subtitulo, periodo=periodo)
    wb.save(ruta)
    return ruta


def excel_gastos_bytes(
    df: pd.DataFrame,
    *,
    subtitulo: str = "Mercado Pago · gastos personales",
    periodo: str = "",
) -> bytes:
    wb = _construir_wb_gastos(df, subtitulo=subtitulo, periodo=periodo)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
