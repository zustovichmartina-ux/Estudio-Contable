"""Cola de rutinas, sin Turso ni red. SQLite temporal."""
from __future__ import annotations

import importlib.util
import json
from datetime import datetime
from io import BytesIO
from pathlib import Path

import openpyxl
import pytest

import database
import rutinas

ROOT = Path(__file__).resolve().parents[1]


def _db(tmp_path, monkeypatch):
    monkeypatch.setattr(database, "DB_PATH", tmp_path / "estudio.db")
    monkeypatch.setattr(database, "_turso_conn_obj", None)
    monkeypatch.setattr(database, "_credenciales_turso", lambda: ("", ""))


def _cli():
    ruta = ROOT / "scripts" / "cola_rutinas.py"
    spec = importlib.util.spec_from_file_location("cola_rutinas_cli", ruta)
    assert spec and spec.loader
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


class _CursorSinIter:
    """Cursor de mentira, como el de libsql: tiene fetchall y no es iterable."""

    def __init__(self):
        self.description = [("cid",), ("name",), ("type",)]
        self._filas = [(0, "id", "INTEGER"), (1, "preview", "TEXT")]

    def fetchall(self):
        return list(self._filas)


def test_cursor_turso_se_puede_recorrer():
    crudo = _CursorSinIter()
    with pytest.raises(TypeError, match="not iterable"):
        list(crudo)
    cur = database._CursorCompatTurso(crudo)
    filas = list(cur)
    assert [fila[1] for fila in filas] == ["id", "preview"]
    assert filas[1]["name"] == "preview"
    assert {fila["name"] for fila in database._CursorCompatTurso(crudo)} == {"id", "preview"}


def test_tablas_idempotentes(tmp_path, monkeypatch):
    _db(tmp_path, monkeypatch)
    rutinas.asegurar_tablas()
    pedido = rutinas.crear_pedido("fcc_monotributistas", "Marti", "2026-09")
    rutinas.asegurar_tablas()
    otra = rutinas.obtener_pedido(pedido["id"])
    assert otra is not None
    assert otra["parametros"] == "2026-09"
    with database.obtener_conexion() as conn:
        idx = conn.execute(
            "SELECT name FROM sqlite_master WHERE type = 'index' AND name = 'idx_rutina_pedidos_abierto'"
        ).fetchone()
        cols = {fila[1] for fila in conn.execute("PRAGMA table_info(rutina_pedidos)").fetchall()}
    assert idx is not None
    assert cols == {
        "id",
        "rutina",
        "parametros",
        "solicitado_por",
        "creado_en",
        "estado",
        "tomado_en",
        "terminado_en",
        "resultado",
        "archivos",
        "preview",
    }


def test_agrega_preview_si_la_tabla_era_vieja(tmp_path, monkeypatch):
    _db(tmp_path, monkeypatch)
    with database.obtener_conexion() as conn:
        conn.execute(
            """
            CREATE TABLE rutina_pedidos (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                rutina TEXT NOT NULL,
                parametros TEXT,
                solicitado_por TEXT NOT NULL,
                creado_en TEXT NOT NULL,
                estado TEXT NOT NULL,
                tomado_en TEXT,
                terminado_en TEXT,
                resultado TEXT,
                archivos TEXT
            )
            """
        )
        conn.execute(
            """
            INSERT INTO rutina_pedidos (rutina, solicitado_por, creado_en, estado)
            VALUES ('fcc_monotributistas', 'Marti', '2026-10-01T15:00:00+00:00', 'OK')
            """
        )
        conn.commit()
    rutinas.asegurar_tablas()
    rutinas.asegurar_tablas()
    pedido = rutinas.obtener_pedido(1)
    assert pedido is not None
    assert pedido["preview"] is None
    assert pedido["estado"] == "OK"


def test_catalogo_requisitos_a_confirmar():
    proy = rutinas.rutina_por_codigo(rutinas.CODIGO_PROYECCION)
    assert proy is not None
    assert proy["requisitos"] == (
        "Listado de imputación contable resumido",
        "Copia del F931",
        "Copia de IIBB",
        "Copia de TISH",
    )
    assert proy["ayuda"] == (
        "Todo en una carpeta MMAAAA (ej. 082026) dentro de la carpeta de Proyecciones del cliente.",
    )
    assert "ejercicio" not in " ".join(proy["requisitos"]).casefold()
    assert "ejercicio siguiente" in rutinas.AVISO_EJERCICIO
    reglas = " ".join(rutinas.REGLAS_CONTROL_PROYECCION)
    assert "NO TIENE EMPLEADOS" in reglas
    assert "TISH no es Sí" in reglas
    assert rutinas.ESTADO_SIN_CARPETA in reglas
    assert rutinas.COLUMNAS_RESULTADO_PROYECCION == (
        "Cliente",
        "Carpeta",
        "Imputación",
        "F931",
        "IIBB",
        "TISH",
        "Estado",
    )
    assert rutinas.periodo_mmaaaa("08-2026") == "082026"
    assert rutinas.periodo_mmaaaa("01-2026") == "012026"
    for item in rutinas.RUTINAS:
        if item["codigo"] == rutinas.CODIGO_PROYECCION:
            continue
        assert item["requisitos"]
        for texto in item["requisitos"]:
            assert str(texto).startswith("A CONFIRMAR")


def test_inicializar_bd_engancha_la_tabla():
    texto = (ROOT / "database.py").read_text(encoding="utf-8")
    app = (ROOT / "app.py").read_text(encoding="utf-8")
    assert "inicializar_tablas_rutinas" in texto
    assert '"Rutinas"' in app
    assert "render_rutinas" in app


def test_no_duplica_pendiente_ni_en_curso(tmp_path, monkeypatch):
    _db(tmp_path, monkeypatch)
    primero = rutinas.crear_pedido("control_fcc_portal_iva", "Sol")
    with pytest.raises(rutinas.PedidoDuplicado):
        rutinas.crear_pedido("control_fcc_portal_iva", "Guadi", "otro")
    rutinas.tomar_pedido(primero["id"])
    with pytest.raises(rutinas.PedidoDuplicado):
        rutinas.crear_pedido("control_fcc_portal_iva", "Sol")
    rutinas.terminar_pedido(primero["id"], "OK", "Revisado")
    segundo = rutinas.crear_pedido("control_fcc_portal_iva", "Sol", "  ")
    assert segundo["id"] != primero["id"]
    assert segundo["parametros"] is None
    assert segundo["estado"] == "PENDIENTE"


def test_dos_rutinas_abiertas_a_la_vez(tmp_path, monkeypatch):
    _db(tmp_path, monkeypatch)
    a = rutinas.crear_pedido("aviso_bazan_bajar_archivos", "Lu")
    b = rutinas.crear_pedido("bazan_detalle_items", "Lu", "marzo")
    assert a["rutina"] != b["rutina"]
    assert rutinas.pedido_abierto(a["rutina"])["id"] == a["id"]


def test_cancelar_solo_pendiente(tmp_path, monkeypatch):
    _db(tmp_path, monkeypatch)
    pedido = rutinas.crear_pedido("seguimiento_balances_urgencia", "Tobi")
    cancelado = rutinas.cancelar_pedido(pedido["id"])
    assert cancelado["estado"] == "CANCELADO"
    assert cancelado["terminado_en"]
    assert "Cancelado" in (cancelado["resultado"] or "")
    with pytest.raises(rutinas.EstadoInvalido):
        rutinas.cancelar_pedido(pedido["id"])
    nuevo = rutinas.crear_pedido("seguimiento_balances_urgencia", "Tobi")
    rutinas.tomar_pedido(nuevo["id"])
    with pytest.raises(rutinas.EstadoInvalido):
        rutinas.cancelar_pedido(nuevo["id"])


def test_tomar_es_una_sola_vez(tmp_path, monkeypatch):
    _db(tmp_path, monkeypatch)
    pedido = rutinas.crear_pedido("bazan_detalle_items", "Cami")
    tomado = rutinas.tomar_pedido(pedido["id"])
    assert tomado["estado"] == "EN_CURSO"
    assert tomado["tomado_en"]
    with pytest.raises(rutinas.EstadoInvalido):
        rutinas.tomar_pedido(pedido["id"])
    with pytest.raises(rutinas.PedidoNoEncontrado):
        rutinas.tomar_pedido(99999)


def test_terminar_ok_y_error(tmp_path, monkeypatch):
    _db(tmp_path, monkeypatch)
    pedido = rutinas.crear_pedido("fcc_monotributistas", "Agus", "2026-08")
    with pytest.raises(rutinas.EstadoInvalido):
        rutinas.terminar_pedido(pedido["id"], "OK", "todavía no")
    rutinas.tomar_pedido(pedido["id"])
    with pytest.raises(rutinas.ErrorRutina):
        rutinas.terminar_pedido(pedido["id"], "CANCELADO", "no")
    with pytest.raises(rutinas.ErrorRutina):
        rutinas.terminar_pedido(pedido["id"], "OK", "   ")
    cerrado = rutinas.terminar_pedido(
        pedido["id"],
        "error",
        "Portal caído",
        archivos=r"\\servidor\FCC\salida.xlsx",
    )
    assert cerrado["estado"] == "ERROR"
    assert cerrado["resultado"] == "Portal caído"
    assert cerrado["archivos"] == r"\\servidor\FCC\salida.xlsx"
    assert cerrado["terminado_en"]
    with pytest.raises(rutinas.EstadoInvalido):
        rutinas.terminar_pedido(pedido["id"], "OK", "de nuevo")


def test_nombre_vacio_y_rutina_desconocida(tmp_path, monkeypatch):
    _db(tmp_path, monkeypatch)
    with pytest.raises(rutinas.ErrorRutina):
        rutinas.crear_pedido("fcc_monotributistas", "  ")
    with pytest.raises(rutinas.ErrorRutina):
        rutinas.crear_pedido("no_existe", "Marti")


def test_fecha_argentina():
    assert rutinas.formatear_fecha_ar("2026-10-01T15:30:00+00:00") == "01/10/2026 12:30"
    assert rutinas.formatear_fecha_ar("") == ""
    assert rutinas.formatear_fecha_ar(None) == ""


def test_listar_por_estado(tmp_path, monkeypatch):
    _db(tmp_path, monkeypatch)
    rutinas.crear_pedido("fcc_monotributistas", "Marti")
    otro = rutinas.crear_pedido("bazan_detalle_items", "Marti")
    rutinas.tomar_pedido(otro["id"])
    pendientes = rutinas.listar_pedidos(estado="PENDIENTE")
    assert len(pendientes) == 1
    assert pendientes[0]["rutina"] == "fcc_monotributistas"
    assert len(rutinas.listar_pedidos(estado="EN_CURSO")) == 1
    with pytest.raises(rutinas.ErrorRutina):
        rutinas.listar_pedidos(estado="RARO")


def test_terminar_guarda_preview(tmp_path, monkeypatch):
    _db(tmp_path, monkeypatch)
    pedido = rutinas.crear_pedido("proyecciones_ganancias_iva", "Sol", "2026-09")
    rutinas.tomar_pedido(pedido["id"])
    with pytest.raises(rutinas.ErrorRutina):
        rutinas.normalizar_preview(["no"])
    ruta = r"\\TANGOSRV\Compartido\CLIENTES\ACME\Ganancias.xlsx"
    cerrado = rutinas.terminar_pedido(
        pedido["id"],
        "OK",
        "Proyección actualizada",
        preview={
            "resumen_md": "Quedó la proyección de **Ganancias**.",
            "tablas": [
                {"titulo": "Cambios", "columnas": ["Cliente", "Celda"], "filas": [["ACME", "C12"]]}
            ],
            "archivos": [ruta],
        },
    )
    assert cerrado["preview"]["resumen_md"].startswith("Quedó")
    assert cerrado["preview"]["tablas"][0]["filas"] == [["ACME", "C12"]]
    assert cerrado["archivos"] == ruta
    assert rutinas.filas_tabla_preview(cerrado["preview"]["tablas"][0]) == [
        {"Cliente": "ACME", "Celda": "C12"}
    ]
    assert rutinas.ultimo_pedido("proyecciones_ganancias_iva")["id"] == pedido["id"]


def test_cli_preview_json(tmp_path, monkeypatch, capsys):
    _db(tmp_path, monkeypatch)
    cli = _cli()
    pedido = rutinas.crear_pedido("bazan_detalle_items", "Lu")
    rutinas.tomar_pedido(pedido["id"])
    path = tmp_path / "preview.json"
    path.write_text(
        json.dumps(
            {
                "resumen_md": "Detalle **armado**.",
                "tablas": [
                    {
                        "titulo": "Ítems",
                        "columnas": ["Archivo", "Fila"],
                        "filas": [["items.xlsx", 4]],
                    }
                ],
                "archivos": [r"D:\Bazan\items.xlsx"],
            }
        ),
        encoding="utf-8",
    )
    assert (
        cli.main(
            [
                "terminar",
                str(pedido["id"]),
                "--estado",
                "OK",
                "--resultado",
                "Detalle listo",
                "--preview-json",
                str(path),
            ]
        )
        == 0
    )
    data = json.loads(capsys.readouterr().out)
    assert data["pedido"]["preview"]["tablas"][0]["titulo"] == "Ítems"
    assert data["pedido"]["archivos"] == r"D:\Bazan\items.xlsx"

    otro = rutinas.crear_pedido("bazan_detalle_items", "Lu")
    rutinas.tomar_pedido(otro["id"])
    malo = tmp_path / "malo.json"
    malo.write_text("{", encoding="utf-8")
    assert (
        cli.main(
            [
                "terminar",
                str(otro["id"]),
                "--estado",
                "OK",
                "--resultado",
                "no",
                "--preview-json",
                str(malo),
            ]
        )
        == 1
    )
    fallo = json.loads(capsys.readouterr().out)
    assert fallo["ok"] is False
    assert rutinas.obtener_pedido(otro["id"])["estado"] == "EN_CURSO"
    assert (
        cli.main(
            [
                "terminar",
                str(otro["id"]),
                "--estado",
                "ERROR",
                "--resultado",
                "sin archivo",
                "--preview-json",
                str(tmp_path / "no-esta.json"),
            ]
        )
        == 1
    )


def test_cli_listar_tomar_terminar(tmp_path, monkeypatch, capsys):
    _db(tmp_path, monkeypatch)
    cli = _cli()
    assert cli.main(["listar", "--estado", "PENDIENTE"]) == 0
    assert json.loads(capsys.readouterr().out) == []

    pedido = rutinas.crear_pedido("aviso_bazan_bajar_archivos", "Hernan", "hoy")
    assert cli.main(["listar", "--estado", "PENDIENTE"]) == 0
    lista = json.loads(capsys.readouterr().out)
    assert lista[0]["id"] == pedido["id"]
    assert lista[0]["parametros"] == "hoy"
    assert lista[0]["nombre"] == "Aviso Bazan bajar archivos"

    assert cli.main(["tomar", str(pedido["id"])]) == 0
    tomado = json.loads(capsys.readouterr().out)
    assert tomado["ok"] is True
    assert tomado["pedido"]["estado"] == "EN_CURSO"

    assert cli.main(["tomar", str(pedido["id"])]) == 1
    fallo = json.loads(capsys.readouterr().out)
    assert fallo["ok"] is False
    assert "PENDIENTE" in fallo["error"]

    assert (
        cli.main(
            [
                "terminar",
                str(pedido["id"]),
                "--estado",
                "OK",
                "--resultado",
                "Archivos bajados",
                "--archivos",
                r"D:\Bazan\octubre",
            ]
        )
        == 0
    )
    fin = json.loads(capsys.readouterr().out)
    assert fin["pedido"]["estado"] == "OK"
    assert fin["pedido"]["resultado"] == "Archivos bajados"
    assert fin["pedido"]["archivos"] == r"D:\Bazan\octubre"

    assert cli.main(["listar", "--estado", "PENDIENTE"]) == 0
    assert json.loads(capsys.readouterr().out) == []


def _xlsx(filas: list[list], encabezados: list[str] | None = None) -> bytes:
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.append(encabezados or [etiqueta for _campo, etiqueta in rutinas.COLUMNAS_FICHA])
    for fila in filas:
        hoja.append(fila)
    buffer = BytesIO()
    libro.save(buffer)
    return buffer.getvalue()


def _fila_acme(**extra) -> list:
    base = {
        "sociedad": "ACME",
        "cuit": 30712345671,
        "activa": "Sí",
        "tish": "Sí",
        "mes_inicio": "01-2020",
        "mes_cierre": datetime(2025, 12, 1),
        "dia_revision": 15,
        "proyeccion_vigente": "Ganancias 2025.xlsx",
        "carpeta_proyecciones": r"C:\ACME\Proyecciones",
        "carpeta_iibb": r"C:\ACME\IIBB",
        "carpeta_f931": r"C:\ACME\F931",
        "ultimo_mes_cargado": "09-2026",
        "notas": "",
    }
    base.update(extra)
    return [base[campo] for campo, _etiqueta in rutinas.COLUMNAS_FICHA]


def test_importar_ficha_reemplaza_y_exporta(tmp_path, monkeypatch):
    _db(tmp_path, monkeypatch)
    primera = _xlsx(
        [
            _fila_acme(),
            _fila_acme(sociedad="BETA", cuit="20111222333", activa="No", tish="No", notas="NO TIENE EMPLEADOS"),
        ]
    )
    assert rutinas.importar_ficha_xlsx(primera) == 2
    assert len(rutinas.listar_ficha()) == 2
    acme = rutinas.listar_ficha()[0]
    assert acme["cuit"] == "30712345671"
    assert acme["mes_cierre"] == "12-2025"
    assert acme["dia_revision"] == "15"
    assert rutinas.cliente_activo(acme["activa"])
    assert not rutinas.cliente_activo("No")
    assert not rutinas.cliente_activo("inactiva")
    assert rutinas.cliente_activo("")

    segunda = _xlsx([_fila_acme(sociedad="GAMMA", cuit="27999888776", activa="Sí", tish="")])
    assert rutinas.importar_ficha_xlsx(segunda) == 1
    sociedades = [fila["sociedad"] for fila in rutinas.listar_ficha()]
    assert sociedades == ["GAMMA"]

    exportada = rutinas.exportar_ficha()
    assert list(exportada[0]) == [etiqueta for _campo, etiqueta in rutinas.COLUMNAS_FICHA]
    assert exportada[0]["Sociedad"] == "GAMMA"
    assert exportada[0]["CUIT"] == "27999888776"


_ENCABEZADOS_OFICINA = [
    "Sociedad",
    "CUIT",
    "Activa (Sí/No)",
    "TISH (Sí/No)",
    "Mes inicio ejercicio (1-12)",
    "Mes cierre ejercicio (1-12)",
    "Día revisión mensual (1-28)",
    "Proyección vigente (nombre archivo o ruta)",
    "Carpeta Proyecciones (si no es estándar)",
    "Carpeta IIBB (si no es estándar)",
    "Carpeta F931 Presentaciones (si no es estándar)",
    "Último mes cargado (MM-YYYY)",
    "Notas",
]


def test_ficha_encabezados_oficina(tmp_path, monkeypatch):
    """La hoja Clientes gana aunque no sea la activa, y los meses llegan como número."""
    _db(tmp_path, monkeypatch)
    libro = openpyxl.Workbook()
    instrucciones = libro.active
    instrucciones.title = "Instrucciones"
    instrucciones.append(["Sociedad", "Notas"])
    instrucciones.append(["TRAMPA", "no leer"])
    leyenda = libro.create_sheet("Leyenda meses")
    leyenda.append(["Mes", "Nombre"])
    leyenda.append([1, "Enero"])
    clientes = libro.create_sheet("Clientes")
    clientes.append(_ENCABEZADOS_OFICINA)
    clientes.append(
        [
            "ACME",
            30712345671,
            "Sí",
            "Sí",
            1,
            12.0,
            15.0,
            "Ganancias 2025.xlsx",
            r"C:\ACME\Proyecciones",
            r"C:\ACME\IIBB",
            r"C:\ACME\F931",
            "09-2026",
            "",
        ]
    )
    clientes.append(
        [
            "BETA",
            "20111222333",
            "No",
            "No",
            6.0,
            6,
            8,
            "",
            r"C:\BETA\Proyecciones",
            r"C:\BETA\IIBB",
            r"C:\BETA\F931",
            "08-2026",
            "NO TIENE EMPLEADOS",
        ]
    )
    buffer = BytesIO()
    libro.save(buffer)
    assert libro.sheetnames[0] == "Instrucciones"
    assert rutinas.importar_ficha_xlsx(buffer.getvalue()) == 2
    filas = rutinas.listar_ficha()
    assert [fila["sociedad"] for fila in filas] == ["ACME", "BETA"]
    acme = filas[0]
    assert acme["cuit"] == "30712345671"
    assert acme["activa"] == "Sí"
    assert acme["tish"] == "Sí"
    assert acme["mes_inicio"] == "1"
    assert acme["mes_cierre"] == "12"
    assert acme["dia_revision"] == "15"
    assert acme["proyeccion_vigente"] == "Ganancias 2025.xlsx"
    assert acme["carpeta_proyecciones"] == r"C:\ACME\Proyecciones"
    assert acme["carpeta_iibb"] == r"C:\ACME\IIBB"
    assert acme["carpeta_f931"] == r"C:\ACME\F931"
    assert acme["ultimo_mes_cargado"] == "09-2026"
    assert filas[1]["mes_inicio"] == "6"
    assert filas[1]["notas"] == "NO TIENE EMPLEADOS"
    assert not rutinas.cliente_activo(filas[1]["activa"])
    etiquetas = [nombre for nombre, _fila in rutinas.etiquetas_clientes()]
    assert etiquetas == ["ACME"]


def test_ficha_sin_sociedad_falla(tmp_path, monkeypatch):
    _db(tmp_path, monkeypatch)
    with pytest.raises(rutinas.ErrorRutina, match="Sociedad"):
        rutinas.importar_ficha_xlsx(_xlsx([["x"]], encabezados=["CUIT", "Activa"]))
    with pytest.raises(rutinas.ErrorRutina, match="vacío"):
        rutinas.leer_ficha_xlsx(b"")
    assert rutinas.exportar_ficha() == []


def test_clientes_todos_y_parametros(tmp_path, monkeypatch):
    _db(tmp_path, monkeypatch)
    rutinas.importar_ficha_xlsx(
        _xlsx(
            [
                _fila_acme(),
                _fila_acme(sociedad="ACME", cuit="20999888776", activa="Sí"),
                _fila_acme(sociedad="BETA", cuit="20111222333", activa="inactiva", notas="NO TIENE EMPLEADOS"),
                _fila_acme(sociedad="GAMMA", cuit="", activa="Sí", tish="No"),
            ]
        )
    )
    etiquetas = [nombre for nombre, _fila in rutinas.etiquetas_clientes()]
    assert "BETA" not in etiquetas
    assert "ACME (30712345671)" in etiquetas
    assert "ACME (20999888776)" in etiquetas
    assert "GAMMA" in etiquetas

    todos, es_todos = rutinas.resolver_clientes(["Todos"])
    assert es_todos
    assert {fila["sociedad"] for fila in todos} == {"ACME", "GAMMA"}
    uno, es_todos = rutinas.resolver_clientes(["GAMMA"])
    assert not es_todos
    assert [fila["sociedad"] for fila in uno] == ["GAMMA"]

    payload = json.loads(rutinas.parametros_proyeccion("10-2026", uno, todos=False))
    assert payload == {
        "periodo": "102026",
        "todos": False,
        "clientes": [{"sociedad": "GAMMA", "cuit": ""}],
    }
    pedido = rutinas.crear_pedido(rutinas.CODIGO_PROYECCION, "Marti", json.dumps(payload, ensure_ascii=False))
    assert json.loads(pedido["parametros"])["periodo"] == "102026"
    assert rutinas.resumen_parametros(pedido["parametros"]) == "102026 · GAMMA"
    amplio = rutinas.parametros_proyeccion("08-2026", todos, todos=True)
    assert json.loads(amplio)["periodo"] == "082026"
    assert rutinas.resumen_parametros(amplio).startswith("082026 · Todos")
    assert rutinas.resumen_parametros("octubre a mano") == "octubre a mano"
    with pytest.raises(rutinas.ErrorRutina):
        rutinas.parametros_proyeccion("2026-10", uno, todos=False)
    with pytest.raises(rutinas.ErrorRutina):
        rutinas.parametros_proyeccion("10-2026", [], todos=False)
    assert not rutinas.periodo_valido("13-2026")
    assert rutinas.periodo_valido("01-2026")


def test_tabla_control_proyeccion_en_preview():
    tabla = rutinas.tabla_control_proyeccion(
        [
            {
                "Cliente": "ACME",
                "Carpeta": "OK",
                "Imputación": "OK",
                "F931": "NO APLICA",
                "IIBB": "OK",
                "TISH": "OK",
                "Estado": "FALTA",
            },
            {
                "Cliente": "BETA",
                "Carpeta": rutinas.ESTADO_SIN_CARPETA,
                "Estado": rutinas.ESTADO_SIN_CARPETA,
            },
        ]
    )
    preview = rutinas.normalizar_preview(
        {"resumen_md": "Control de **octubre**.", "tablas": [tabla], "archivos": []}
    )
    filas = rutinas.filas_tabla_preview(preview["tablas"][0])
    assert list(filas[0]) == list(rutinas.COLUMNAS_RESULTADO_PROYECCION)
    assert filas[0]["Carpeta"] == "OK"
    assert filas[0]["F931"] == "NO APLICA"
    assert filas[0]["IIBB"] == "OK"
    assert filas[0]["TISH"] == "OK"
    assert filas[0]["Estado"] == "FALTA"
    assert filas[1]["Carpeta"] == "SIN CARPETA MMAAAA"
    assert filas[1]["Estado"] == "SIN CARPETA MMAAAA"
    assert filas[1]["Imputación"] == ""
    assert "Papel ejercicio siguiente" not in filas[0]
    assert set(filas[0].values()) <= set(rutinas.ESTADOS_CONTROL) | {"ACME"}


def test_cli_ficha_json(tmp_path, monkeypatch, capsys):
    _db(tmp_path, monkeypatch)
    cli = _cli()
    assert cli.main(["ficha", "--json"]) == 0
    assert json.loads(capsys.readouterr().out) == []
    rutinas.importar_ficha_xlsx(_xlsx([_fila_acme(notas="NO TIENE EMPLEADOS")]))
    assert cli.main(["ficha", "--json"]) == 0
    data = json.loads(capsys.readouterr().out)
    assert data[0]["Sociedad"] == "ACME"
    assert data[0]["Notas"] == "NO TIENE EMPLEADOS"
    assert data[0]["Día revisión"] == "15"
    assert data[0]["TISH"] == "Sí"
