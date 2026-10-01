"""Pruebas sin certificado de ARCA ni conexión. Ejecutar: .venv/bin/python -m pytest -q tests"""
import base64
import datetime as dt
import json
import os
import shutil
import subprocess
import sys

import openpyxl
import pytest

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BASE)
from arca import wsaa  # noqa: E402
from arca.comprobante import Comprobante, agrupar, cuit_valido, leer_planilla  # noqa: E402
from arca.pdf import url_qr  # noqa: E402
from arca.wsfe import ORDEN_DET, xml_fecaesolicitar  # noqa: E402

HOY = dt.date(2026, 10, 1)
MODELO = os.path.join(BASE, "plantillas", "Facturas_a_emitir_MODELO.xlsx")


def fila(**kw):
    base = dict(_fila=99, marca=None, id=None, cuit_emisor="30712345671", pto_vta=5, tipo="A",
                cond_iva_emisor="IVA Responsable Inscripto", concepto="Productos", fecha=HOY,
                doc_tipo="CUIT", doc_nro="30709998885", nombre_rec="X", cond_iva="IVA Responsable Inscripto",
                detalle="d", neto=1000, alicuota="21%", fsd=None, fsh=None, fvto=None, cbu=None, fce_transf=None,
                asoc_id=None, asoc_tipo=None, asoc_pto=None, asoc_nro=None, asoc_fecha=None)
    base.update(kw)
    return base


def val(*filas):
    c = Comprobante(list(filas), hoy=HOY)
    c.validar()
    return c


def test_cuit():
    assert cuit_valido("23422843434")  # CUIT del estudio
    assert cuit_valido("20-33444555-1")
    assert not cuit_valido("23422843435")


def test_a_ok_y_montos():
    c = val(fila(), fila(alicuota="10,5%", neto=200))
    assert not c.errores, c.errores
    assert c.det["ImpNeto"] == 1200 and c.det["ImpIVA"] == 231 and c.det["ImpTotal"] == 1431
    assert c.cbte_tipo == 1 and c.det["CondicionIVAReceptorId"] == 1


def test_iva_27_y_0_y_exento():
    c = val(fila(alicuota="27%", neto=100), fila(alicuota="0%", neto=50), fila(alicuota="Exento", neto=10),
            fila(alicuota="No gravado", neto=5))
    assert not c.errores, c.errores
    d = c.det
    assert (d["ImpNeto"], d["ImpIVA"], d["ImpOpEx"], d["ImpTotConc"], d["ImpTotal"]) == (150, 27, 10, 5, 192)
    assert {a["Id"] for a in d["Iva"]} == {6, 3}


def test_c_sin_iva():
    c = val(fila(tipo="C", cond_iva_emisor="Responsable Monotributo", alicuota="No corresponde (C)",
                 cond_iva="Consumidor Final", doc_tipo="DNI", doc_nro="30111222"))
    assert not c.errores, c.errores
    assert c.cbte_tipo == 11 and c.det["ImpIVA"] == 0 and c.det["Iva"] is None and c.det["ImpNeto"] == 1000


def test_c_con_alicuota_falla():
    assert val(fila(tipo="C", cond_iva_emisor="Responsable Monotributo", alicuota="21%")).errores


def test_cond_iva_por_clase():
    e = val(fila(tipo="A", cond_iva="Consumidor Final")).errores
    assert any("10243" in x for x in e)
    e = val(fila(tipo="B", cond_iva="IVA Responsable Inscripto")).errores
    assert any("10243" in x for x in e)
    assert any("RG 5616" in x for x in val(fila(cond_iva=None)).errores)


def test_a_exige_cuit():
    assert any("CUIT" in x for x in val(fila(doc_tipo="DNI", doc_nro="30111222")).errores)


def test_servicios_exige_fechas():
    assert val(fila(concepto="Servicios")).errores
    c = val(fila(concepto="Servicios", fsd=dt.date(2026, 9, 1), fsh=dt.date(2026, 9, 30), fvto=dt.date(2026, 10, 10)))
    assert not c.errores and c.det["FchServDesde"] == "20260901" and c.det["FchVtoPago"] == "20261010"


def test_fecha_fuera_de_rango():
    assert val(fila(fecha=dt.date(2026, 9, 20))).errores  # productos ±5


def test_cf_sin_identificar_umbral():
    ok = val(fila(tipo="B", cond_iva="Consumidor Final", doc_tipo="Consumidor Final (sin identificar)", doc_nro=0,
                  neto=100000))
    assert not ok.errores and ok.det["DocTipo"] == 99 and ok.det["DocNro"] == 0
    big = val(fila(tipo="B", cond_iva="Consumidor Final", doc_tipo="Consumidor Final (sin identificar)", doc_nro=0,
                   neto=9000000))
    assert any("10.000.000" in x for x in big.errores)


def test_fce():
    c = val(fila(tipo="FCE A", fvto=dt.date(2026, 10, 31), cbu="0110599520000012345678",
                 fce_transf="SCA - Transferencia al Sistema de Circulación Abierta"))
    assert not c.errores, c.errores
    assert c.cbte_tipo == 201 and {o["Id"] for o in c.det["Opcionales"]} == {"2101", "27"}
    assert val(fila(tipo="FCE A", fvto=dt.date(2026, 10, 31))).errores


def test_emisor_incoherente():
    assert val(fila(tipo="C", cond_iva_emisor="IVA Responsable Inscripto", alicuota="No corresponde (C)")).errores


def test_xml_orden_y_cuit_representada():
    c = val(fila())
    x = xml_fecaesolicitar({"token": "T", "sign": "S", "cuit": "30712345671"}, 5, 1, c.con_numero(8))
    assert "<Cuit>30712345671</Cuit>" in x and "<CbteDesde>8</CbteDesde>" in x
    pos = [x.find(f"<{k}>") for k in ORDEN_DET if f"<{k}>" in x]
    assert pos == sorted(pos)


def test_qr_ejemplo_oficial():
    """Reproduce el JSON del ejemplo de 'Especificaciones del QR' de ARCA."""
    txt, datos = url_qr(dt.date(2020, 10, 13), "30000000007", 10, 1, 94, 12100, 80, 20000000001,
                        70417054367476, moneda="DOL", ctz=65)
    assert txt.startswith("https://www.arca.gob.ar/fe/qr/?p=")
    js = json.loads(base64.b64decode(txt.split("?p=")[1]))
    assert js == {"ver": 1, "fecha": "2020-10-13", "cuit": 30000000007, "ptoVta": 10, "tipoCmp": 1, "nroCmp": 94,
                  "importe": 12100, "moneda": "DOL", "ctz": 65, "tipoDocRec": 80, "nroDocRec": 20000000001,
                  "tipoCodAut": "E", "codAut": 70417054367476}


def test_modelo_lectura_y_grupos():
    wb = openpyxl.load_workbook(MODELO, data_only=True)
    assert [s for s in wb.sheetnames if wb[s].sheet_state == "visible"] == ["Facturas"]
    filas = leer_planilla(wb)
    grupos = agrupar(filas)
    assert len(filas) == 9 and len(grupos) == 8
    assert all(Comprobante(g).es_ejemplo for g in grupos)


@pytest.fixture
def cert_prueba(tmp_path):
    """Certificado AUTOFIRMADO descartable (no es de ARCA) sólo para probar la firma CMS."""
    k, c = tmp_path / "t.key", tmp_path / "t.crt"
    subprocess.run(["openssl", "req", "-x509", "-newkey", "rsa:2048", "-nodes", "-keyout", str(k), "-out", str(c),
                    "-days", "1", "-subj", "/C=AR/O=PRUEBA/CN=prueba/serialNumber=CUIT 20000000001"],
                   check=True, capture_output=True)
    return str(c), str(k)


def test_firma_cms(cert_prueba, tmp_path):
    c, k = cert_prueba
    tra = wsaa.crear_tra("wsfe")
    b64 = wsaa.firmar_tra(tra, c, k)
    der = tmp_path / "cms.der"
    der.write_bytes(base64.b64decode(b64))
    out = subprocess.run(["openssl", "cms", "-verify", "-inform", "DER", "-in", str(der), "-noverify"],
                         capture_output=True)
    assert out.returncode == 0, out.stderr
    assert b"<service>wsfe</service>" in out.stdout


def test_cache_ta(cert_prueba, tmp_path, monkeypatch):
    c, k = cert_prueba
    cache = tmp_path / "cache"
    cache.mkdir()
    exp = (dt.datetime.now(wsaa.TZ) + dt.timedelta(hours=11)).isoformat()
    f = wsaa._cache_file(str(cache), "homo", "wsfe", c)
    json.dump({"token": "TK", "sign": "SG", "expiration": exp}, open(f, "w"))
    monkeypatch.setattr(wsaa, "soap_post", lambda *a, **kw: pytest.fail("no debía llamar a WSAA"))
    ta = wsaa.obtener_ta("homo", c, k, "wsfe", str(cache))
    assert ta["source"] == "cache" and ta["token"] == "TK"
    # vencido -> debe pedir uno nuevo
    json.dump({"token": "TK", "sign": "SG", "expiration": dt.datetime.now(wsaa.TZ).isoformat()}, open(f, "w"))
    llamado = {}

    class R:
        status_code = 200
        content = ('<soapenv:Envelope xmlns:soapenv="http://schemas.xmlsoap.org/soap/envelope/"><soapenv:Body>'
                   '<loginCmsResponse xmlns="http://wsaa.view.sua.dvadac.desein.afip.gov"><loginCmsReturn>'
                   '&lt;loginTicketResponse&gt;&lt;header&gt;&lt;expirationTime&gt;2026-10-02T08:00:00.000-03:00'
                   '&lt;/expirationTime&gt;&lt;/header&gt;&lt;credentials&gt;&lt;token&gt;NUEVO&lt;/token&gt;'
                   '&lt;sign&gt;FIRMA&lt;/sign&gt;&lt;/credentials&gt;&lt;/loginTicketResponse&gt;'
                   '</loginCmsReturn></loginCmsResponse></soapenv:Body></soapenv:Envelope>').encode()
        text = ""

    def fake(url, body, action):
        llamado["url"] = url
        return R()
    monkeypatch.setattr(wsaa, "soap_post", fake)
    ta = wsaa.obtener_ta("homo", c, k, "wsfe", str(cache))
    assert ta["token"] == "NUEVO" and "wsaahomo" in llamado["url"]
    assert oct(os.stat(f).st_mode & 0o777) == "0o600"


def test_dry_run_excel(tmp_path):
    x = tmp_path / "p.xlsx"
    shutil.copy(MODELO, x)
    r = subprocess.run([sys.executable, os.path.join(BASE, "emitir.py"), "--excel", str(x), "--dry-run",
                        "--incluir-ejemplos", "--pdf-dir", str(tmp_path / "pdf"), "--pdf-borrador"],
                       capture_output=True, text=True)
    assert r.returncode == 0, r.stderr
    wb = openpyxl.load_workbook(x)
    assert "Resultado" in wb.sheetnames
    res = wb["Resultado"]
    estados = [res.cell(i, 12).value for i in range(2, res.max_row + 1)]
    assert estados == ["OK-DRYRUN"] * 8, estados
    # la hoja de entrada queda intacta (fórmulas, validaciones, hoja oculta)
    f = wb["Facturas"]
    orig = openpyxl.load_workbook(MODELO)["Facturas"]
    for row in range(1, 12):
        for col in range(1, f.max_column + 1):
            assert f.cell(row, col).value == orig.cell(row, col).value
    assert len(f.data_validations.dataValidation) == len(orig.data_validations.dataValidation)
    assert wb["Codigos"].sheet_state == "hidden"
    assert len(list((tmp_path / "pdf").glob("*.pdf"))) == 8


def test_sin_ejemplos_no_procesa(tmp_path):
    x = tmp_path / "p.xlsx"
    shutil.copy(MODELO, x)
    r = subprocess.run([sys.executable, os.path.join(BASE, "emitir.py"), "--excel", str(x), "--dry-run"],
                       capture_output=True, text=True)
    assert "sin comprobantes" in r.stdout


def test_prod_exige_confirmacion(tmp_path):
    r = subprocess.run([sys.executable, os.path.join(BASE, "emitir.py"), "--excel", MODELO, "--env", "prod"],
                       capture_output=True, text=True)
    assert r.returncode != 0 and "confirmar-produccion" in r.stderr


# ---------------- Notas de crédito / débito ----------------
C_MONO = dict(tipo="NC C", cond_iva_emisor="Responsable Monotributo", alicuota="No corresponde (C)",
              cond_iva="Consumidor Final", doc_tipo="Consumidor Final (sin identificar)", doc_nro=0)


def test_nc_sin_asociado_falla():
    assert any("10197" in x for x in val(fila(**C_MONO)).errores)


def test_factura_con_asociado_falla():
    assert val(fila(asoc_tipo="A", asoc_pto=5, asoc_nro=1)).errores


def test_nc_c_externa_ok_y_xml():
    c = val(fila(**C_MONO, asoc_tipo="C", asoc_pto=1, asoc_nro=7, asoc_fecha=dt.date(2026, 9, 30)))
    assert not c.errores, c.errores
    assert c.cbte_tipo == 13 and c.clase_doc == "NC"
    x = xml_fecaesolicitar({"token": "T", "sign": "S", "cuit": c.cuit}, 5, 13, c.con_numero(1))
    assert ("<CbtesAsoc><CbteAsoc><Tipo>11</Tipo><PtoVta>1</PtoVta><Nro>7</Nro><Cuit>30712345671</Cuit>"
            "<CbteFch>20260930</CbteFch></CbteAsoc></CbtesAsoc>") in x
    pos = [x.find(f"<{k}>") for k in ORDEN_DET if f"<{k}>" in x]
    assert pos == sorted(pos)


def test_nd_b_y_nc_a_codigos():
    nd = val(fila(tipo="ND B", cond_iva="Consumidor Final", doc_tipo="Consumidor Final (sin identificar)",
                  doc_nro=0, asoc_tipo="B", asoc_pto=5, asoc_nro=3))
    assert not nd.errores and nd.cbte_tipo == 7
    nc = val(fila(tipo="NC A", asoc_tipo="A", asoc_pto=5, asoc_nro=3))
    assert not nc.errores and nc.cbte_tipo == 3


def test_nc_distinta_letra_falla():
    e = val(fila(**C_MONO, asoc_tipo="A", asoc_pto=1, asoc_nro=7)).errores
    assert any("misma letra" in x for x in e)


def test_nc_supera_original():
    c = val(fila(**C_MONO, asoc_tipo="C", asoc_pto=1, asoc_nro=7))
    assert not c.verificar_asociado(total_original=500)
    c2 = val(fila(**C_MONO, asoc_tipo="C", asoc_pto=1, asoc_nro=7))
    from decimal import Decimal
    assert not c2.verificar_asociado(total_original=1500, ya_acreditado=Decimal(600))  # 1000 > 900 disponible
    c3 = val(fila(**C_MONO, asoc_tipo="C", asoc_pto=1, asoc_nro=7))
    assert c3.verificar_asociado(total_original=1000)


def _excel_nc(tmp_path, filas_extra):
    from arca.comprobante import COLUMNAS
    x = tmp_path / "nc.xlsx"
    wb = openpyxl.load_workbook(MODELO)
    ws = wb["Facturas"]
    col = {k: i for i, (k, _, _) in enumerate(COLUMNAS, 1)}
    for r in range(4, 30):
        for k in col:
            if k not in ("iva", "total"):
                ws.cell(r, col[k]).value = None
    base = dict(cuit_emisor="23422843434", razon_emisor="T", cond_iva_emisor="Responsable Monotributo", pto_vta=1,
                concepto="Productos", fecha=dt.date.today(), doc_tipo="Consumidor Final (sin identificar)",
                doc_nro=0, cond_iva="Consumidor Final", detalle="x", alicuota="No corresponde (C)")
    for j, f in enumerate(filas_extra):
        for k, v in {**base, **f}.items():
            ws.cell(4 + j, col[k]).value = v
    wb.save(x)
    r = subprocess.run([sys.executable, os.path.join(BASE, "emitir.py"), "--excel", str(x), "--dry-run",
                        "--pdf-dir", str(tmp_path / "pdf"), "--pdf-borrador"], capture_output=True, text=True)
    assert r.returncode == 0, r.stderr
    res = openpyxl.load_workbook(x)["Resultado"]
    return {res.cell(i, 4).value: (res.cell(i, 12).value, res.cell(i, 13).value) for i in range(2, res.max_row + 1)}


def test_nc_por_id_mismo_excel_aunque_este_antes(tmp_path):
    out = _excel_nc(tmp_path, [dict(tipo="NC C", neto=400, asoc_id="F1"), dict(id="F1", tipo="C", neto=1000)])
    assert out["4"][0] == "OK-DRYRUN" and out["5"][0] == "OK-DRYRUN", out
    pdfs = list((tmp_path / "pdf").glob("BORRADOR_filas4_*_013_*.pdf"))
    assert pdfs
    txt = subprocess.run(["pdftotext", str(pdfs[0]), "-"], capture_output=True, text=True).stdout
    assert "NOTA DE CRÉDITO" in txt and "Comprobante asociado: Factura C (cód. 11) 00001-" in txt


def test_nc_por_id_supera_y_inexistente(tmp_path):
    out = _excel_nc(tmp_path, [dict(id="F1", tipo="C", neto=1000), dict(tipo="NC C", neto=700, asoc_id="F1"),
                               dict(tipo="NC C", neto=400, asoc_id="F1"), dict(tipo="NC C", neto=10, asoc_id="NOEXISTE")])
    assert out["5"][0] == "OK-DRYRUN"
    # en dry-run las NC no "consumen" saldo (no hay número); la de 400 entra sola, la de 1200 no
    assert out["7"][0] == "RECHAZADO-VALIDACIÓN" and "NOEXISTE" in out["7"][1]
    out2 = _excel_nc(tmp_path, [dict(id="F1", tipo="C", neto=1000), dict(tipo="NC C", neto=1200, asoc_id="F1")])
    assert out2["5"][0] == "RECHAZADO-VALIDACIÓN" and "supera" in out2["5"][1]
