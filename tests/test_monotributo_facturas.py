"""Papel de trabajo de recategorización: extracción de referencia, sin PDFs de clientes."""
from __future__ import annotations

import io
from datetime import date

import openpyxl
import pytest
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

import monotributo_facturas as mf
import procesador as proc
from monotributo_proyeccion import proyectar_monotributo

FALTA = "FALTA DATO"


def _pagina(
    *,
    tipo_linea=None,
    pv="00001",
    nro="00000138",
    fecha="15/03/2025",
    desde="01/03/2025",
    hasta="31/03/2025",
    monto="$ 150.000,50",
    moneda="PES",
    tc="1",
    comprador="CLIENTE SRL",
    marca="ORIGINAL",
    emisor="EMISOR QUE NO VA",
    cuit="20-11111111-1",
):
    if tipo_linea is None:
        tipo_linea = f"          CUIT: {cuit}          C          FACTURA"
    bloque_comprador = (
        f"Apellido y Nombre / Razón Social: {comprador}\n" if comprador is not None else ""
    )
    bloque_moneda = f"Moneda: {moneda}\n" if moneda is not None else ""
    bloque_tc = f"Tipo de Cambio: {tc}\n" if tc is not None else ""
    bloque_desde = f"Período Facturado Desde: {desde} Hasta: {hasta}\n" if desde or hasta else ""
    return f"""
{marca}

{tipo_linea}

Razón Social: {emisor}
Punto de Venta: {pv} Comp. Nro: {nro}
Fecha de Emisión: {fecha}
CUIT: {cuit}

{bloque_comprador}Condición frente al IVA: Responsable Monotributo

{bloque_desde}{bloque_moneda}{bloque_tc}Importe Total: {monto}
CAE: 12345678901234
"""


def _rgb(color) -> str:
    if color is None:
        return ""
    return str(getattr(color, "rgb", "") or "")


def _hoja(xlsx: bytes):
    wb = openpyxl.load_workbook(io.BytesIO(xlsx))
    assert wb.sheetnames == ["Facturas"]
    return wb["Facturas"]


def _filas_datos(ws):
    filas = []
    for row in ws.iter_rows(min_row=2, max_col=9):
        if row[2].value in (
            "Total Facturas ($)",
            "Total Notas de Crédito ($)",
            "Neto (Facturas − NC) ($)",
        ):
            break
        if row[0].value is None and row[4].value is None:
            break
        filas.append(row)
    return filas


class _Upload:
    def __init__(self, name: str, data: bytes):
        self.name = name
        self._data = data

    def getvalue(self):
        return self._data


def test_numero_y_comprador_no_toma_al_emisor():
    fila = mf.parsear_texto_recat(_pagina(), "fc.pdf")
    assert fila["N° Factura"] == "FC C 00001-00000138"
    assert fila["Denominación del comprador"] == "CLIENTE SRL"
    assert "EMISOR" not in fila["Denominación del comprador"]
    assert fila["Emisor"] == "EMISOR QUE NO VA"
    assert fila["CUIT emisor"] == "20-11111111-1"
    assert fila["Monto"] == pytest.approx(150000.50)
    assert fila["Moneda"] == "PES"
    assert fila["Tipo de cambio"] == "1"


def test_nota_de_credito_y_campos_ausentes_no_se_inventan():
    texto = _pagina(
        tipo_linea="          CUIT: 20-11111111-1          C          NOTA DE CRÉDITO",
        nro="00000006",
        fecha="02/01/2024",
        desde="",
        hasta="",
        comprador=None,
        moneda=None,
        tc=None,
        monto="6.500,00",
    )
    fila = mf.parsear_texto_recat(texto, "nc.pdf")
    assert fila["N° Factura"] == "NC C 00001-00000006"
    assert fila["Período Desde"] == ""
    assert fila["Período Hasta"] == ""
    assert fila["Denominación del comprador"] == ""
    assert fila["Moneda"] == ""
    assert fila["Tipo de cambio"] == ""
    assert fila["Fecha"] == "02/01/2024"
    assert fila["Monto"] == pytest.approx(6500)
    firmado, dolares, tc = mf.importe_recategorizacion(fila["Monto"], "NC", "", "")
    assert firmado == pytest.approx(-6500)
    assert dolares == 0
    assert tc == 0


def test_tres_copias_cuentan_una_vez_y_si_difieren_se_avisa():
    una = _pagina()
    texto = "\f".join([una.replace("ORIGINAL", m) for m in ("ORIGINAL", "DUPLICADO", "TRIPLICADO")])
    fila = mf.parsear_texto_recat(texto, "fc.pdf")
    assert fila["N° Factura"] == "FC C 00001-00000138"
    assert fila["_avisos_copias"] == []

    distinta = una.replace("Importe Total: $ 150.000,50", "Importe Total: $ 1,00").replace("ORIGINAL", "TRIPLICADO")
    mala = mf.parsear_texto_recat(una + "\f" + distinta, "fc.pdf")
    assert mala["Monto"] == pytest.approx(150000.50)
    assert any("no coinciden" in a for a in mala["_avisos_copias"])


def test_excel_formato_totales_anio_duplicados_y_saltos():
    fc_2025 = mf.parsear_texto_recat(_pagina(fecha="15/03/2025", nro="00000138"), "b.pdf")
    fc_2024 = mf.parsear_texto_recat(
        _pagina(fecha="20/11/2024", nro="00000140", monto="$ 10.000,00", comprador="OTRO SA"),
        "a.pdf",
    )
    nc = mf.parsear_texto_recat(
        _pagina(
            tipo_linea="          CUIT: 20-11111111-1          C          NOTA DE CRÉDITO",
            fecha="02/02/2025",
            nro="00000006",
            monto="$ 2.000,00",
            comprador=None,
            tc=None,
        ),
        "nc.pdf",
    )
    dup = mf.parsear_texto_recat(_pagina(fecha="15/03/2025", nro="00000138"), "b-copia.pdf")
    filas, duplicados = mf.deduplicar_filas_recat([fc_2025, fc_2024, nc, dup])
    assert len(filas) == 3
    assert duplicados[0]["nro"] == "FC C 00001-00000138"
    assert duplicados[0]["descartado"] == "b-copia.pdf"

    xlsx = mf.exportar_papel_facturas(filas, {"duplicados": duplicados, "copias": []})
    ws = _hoja(xlsx)
    assert [c.value for c in ws[1]] == mf.COLUMNAS_PAPEL
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == "A1:K4"
    assert ws.row_dimensions[1].height == 30
    assert ws["A1"].font.bold and _rgb(ws["A1"].font.color).endswith("FFFFFF")
    assert _rgb(ws["A1"].fill.fgColor).endswith("305496")

    datos = _filas_datos(ws)
    assert [c.value for c in datos[0]][0] == date(2024, 11, 20) or datos[0][0].value.year == 2024
    fechas = [c[0].value for c in datos]
    assert fechas == sorted(fechas)
    assert datos[0][0].number_format == "DD/MM/YYYY"
    assert datos[0][3].number_format == "#,##0.00"
    assert datos[0][3].value == pytest.approx(10000)
    numeros = [c[4].value for c in datos]
    assert numeros == [
        "FC C 00001-00000140",
        "NC C 00001-00000006",
        "FC C 00001-00000138",
    ]
    buyer_nc = next(c for c in datos if c[4].value.startswith("NC"))
    assert buyer_nc[7].value == FALTA
    assert _rgb(buyer_nc[7].font.color).endswith("C00000")
    assert all(c[7].value != "EMISOR QUE NO VA" for c in datos)

    textos = [ws.cell(r, 3).value for r in range(1, ws.max_row + 1)]
    assert "Total Facturas ($)" in textos
    assert "Total Notas de Crédito ($)" in textos
    assert "Neto (Facturas − NC) ($)" in textos
    formulas = [ws.cell(r, 4).value for r in range(1, ws.max_row + 1) if isinstance(ws.cell(r, 4).value, str)]
    assert any(f == '=SUMIFS($D$2:$D$4,$E$2:$E$4,"FC*")' for f in formulas)
    assert any(f == '=SUMIFS($D$2:$D$4,$E$2:$E$4,"NC*")' for f in formulas)
    assert any("DATE(2024,1,1)" in f and "FC*" in f for f in formulas)
    assert any("DATE(2025,1,1)" in f and "NC*" in f for f in formulas)
    assert any(str(ws.cell(r, 3).value).startswith("Año ") for r in range(1, ws.max_row + 1))

    nota = "\n".join(
        str(ws.cell(r, 1).value)
        for r in range(1, ws.max_row + 1)
        if isinstance(ws.cell(r, 1).value, str)
    )
    assert "b-copia.pdf" in nota
    assert "00000139" in nota
    assert "punto de venta 00001" in nota
    assert "Emisor 20-11111111-1" not in nota
    assert ws.column_dimensions["A"].width
    assert ws.column_dimensions["C"].width >= 26


def test_usd_en_el_papel_queda_impreso_y_la_categoria_resta_nc(monkeypatch):
    fc = """
  FACTURA C
  Codigo: 011
  CUIT: 20-11222333-9
  Fecha de Emisión: 05/08/2025
  Punto de Venta: 00001
  Comp. Nro: 00000090
  Concepto: 2 - Servicios
  Período Facturado Desde: 01/08/2025 Hasta: 31/08/2025
  Moneda: DOL
  Tipo de Cambio: 1.250,50
  Importe Total: 100,00
  CAE: 33332222111100
  Apellido y Nombre / Razón Social: CLIENTE USD
"""
    nc = """
  NOTA DE CREDITO C
  Codigo: 013
  CUIT: 20-11222333-9
  Fecha de Emisión: 10/08/2025
  Punto de Venta: 00001
  Comp. Nro: 00000002
  Concepto: 2 - Servicios
  Período Facturado Desde: 01/08/2025 Hasta: 31/08/2025
  Moneda: PES
  Importe Total: 2.000,00
  CAE: 12121212121212
  Apellido y Nombre / Razón Social: CLIENTE USD
"""
    parsed_fc = proc.parsear_factura_afip_texto(fc, "usd.pdf")
    assert parsed_fc["N° Factura"] == "FC C 00001-00000090"
    assert parsed_fc["Monto"] == pytest.approx(100)
    assert parsed_fc["Importe Dólares"] == pytest.approx(100)
    assert parsed_fc["Importe Total"] == pytest.approx(125050)
    assert parsed_fc["Período Desde"] == "01/08/2025"
    assert "EMISOR" not in (parsed_fc["Denominación del comprador"] or "")

    sin_periodo = proc.parsear_factura_afip_texto(
        """
        FACTURA C
        Fecha de Emisión: 01/05/2025
        Punto de Venta: 1
        Comp. Nro: 1
        Importe Total: $ 1.000,00
        """,
        "sin-periodo.pdf",
    )
    assert sin_periodo["Período Desde"] == ""
    assert sin_periodo["Fecha"] == "01/05/2025"

    def _fake(pdf_bytes: bytes) -> str:
        return pdf_bytes.decode("utf-8")

    monkeypatch.setattr(proc, "extraer_texto_factura_afip", _fake)
    df, errores = proc.procesar_facturas_monotributo([
        _Upload("usd.pdf", fc.encode("utf-8")),
        _Upload("nc.pdf", nc.encode("utf-8")),
        # Mismos tipo, punto de venta y número; bytes distintos (reimpresión).
        _Upload("usd-dup.pdf", (fc + "\n(reimpresion)\n").encode("utf-8")),
    ])
    assert len(df) == 2
    assert any("duplicado" in str(e.get("motivo", "")).lower() for e in errores)
    assert round(float(df["Importe Total"].sum()), 2) == pytest.approx(123050)
    proy = proyectar_monotributo(df, "A", hoy=date(2026, 9, 10))
    assert proy["facturado_fijo"] == pytest.approx(123050)

    xlsx = proc.exportar_monotributo_excel(df)
    ws = _hoja(xlsx)
    montos = {row[4].value: row[3].value for row in _filas_datos(ws)}
    assert montos["FC C 00001-00000090"] == pytest.approx(100)
    assert montos["NC C 00001-00000002"] == pytest.approx(2000)
    nota = "\n".join(str(ws.cell(r, 1).value or "") for r in range(1, ws.max_row + 1))
    assert "usd-dup.pdf" in nota


def test_varios_emisores_no_se_rechazan_y_los_totales_se_abren():
    uno = _pagina(
        cuit="20-11111111-1",
        emisor="ALFA SA",
        fecha="15/06/2024",
        nro="00000010",
        monto="$ 1.000,00",
    )
    otro = _pagina(
        cuit="27-22222222-3",
        emisor="BETA SRL",
        fecha="20/03/2025",
        nro="00000010",
        monto="$ 4.000,00",
        comprador="OTRO COMPRADOR",
    )
    filas, duplicados = mf.deduplicar_filas_recat([
        mf.parsear_texto_recat(uno, "alfa.pdf"),
        mf.parsear_texto_recat(otro, "beta.pdf"),
    ])
    assert duplicados == []
    assert {f["CUIT emisor"] for f in filas} == {"20-11111111-1", "27-22222222-3"}
    assert {f["Emisor"] for f in filas} == {"ALFA SA", "BETA SRL"}
    assert {f["Denominación del comprador"] for f in filas} >= {"CLIENTE SRL", "OTRO COMPRADOR"}

    xlsx = mf.exportar_papel_facturas(filas)
    ws = _hoja(xlsx)
    assert [c.value for c in ws[1]][-2:] == ["CUIT emisor", "Emisor"]
    formulas = [
        str(ws.cell(r, 4).value)
        for r in range(1, ws.max_row + 1)
        if isinstance(ws.cell(r, 4).value, str) and str(ws.cell(r, 4).value).startswith("=")
    ]
    assert any('"20-11111111-1"' in f and "DATE(2024,1,1)" in f and "FC*" in f for f in formulas)
    assert any('"27-22222222-3"' in f and "DATE(2025,1,1)" in f and "FC*" in f for f in formulas)
    titulos = [str(ws.cell(r, 1).value or "") for r in range(1, ws.max_row + 1)]
    assert any(t.startswith("Emisor 20-11111111-1") and "ALFA SA" in t for t in titulos)
    assert any(t.startswith("Emisor 27-22222222-3") and "BETA SRL" in t for t in titulos)

    def _fake(pdf_bytes: bytes) -> str:
        return pdf_bytes.decode("utf-8")

    with pytest.MonkeyPatch.context() as mp:
        mp.setattr(proc, "extraer_texto_factura_afip", _fake)
        df, errores = proc.procesar_facturas_monotributo(
            [
                _Upload("alfa.pdf", uno.encode("utf-8")),
                _Upload("beta.pdf", otro.encode("utf-8")),
            ],
            cuit_cliente="20111111111",
        )
    assert len(df) == 2
    assert not any("distinto del cliente" in str(e.get("motivo", "")) for e in errores)
    assert round(float(df["Importe Total"].sum()), 2) == pytest.approx(5000)


def test_pdf_sintetico_original_duplicado_triplicado():
    pagina = _pagina(nro="00000007", monto="$ 8.000,00")
    buf = io.BytesIO()
    pdf = canvas.Canvas(buf, pagesize=A4)
    for marca in ("ORIGINAL", "DUPLICADO", "TRIPLICADO"):
        y = 800
        for linea in pagina.replace("ORIGINAL", marca, 1).splitlines():
            pdf.drawString(36, y, linea)
            y -= 14
        pdf.showPage()
    pdf.save()

    texto = proc.extraer_texto_factura_afip(buf.getvalue())
    assert texto.count("\f") == 2
    fila = mf.parsear_texto_recat(texto, "sintetica.pdf")
    assert fila is not None
    assert fila["N° Factura"] == "FC C 00001-00000007"
    assert fila["Monto"] == pytest.approx(8000)
    assert fila["_avisos_copias"] == []
    assert fila["Denominación del comprador"] == "CLIENTE SRL"
