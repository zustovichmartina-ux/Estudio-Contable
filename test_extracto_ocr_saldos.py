"""OCR 11.xxx leído como 71.xxx no debe fabricar débitos de 60 millones."""
from __future__ import annotations

import unittest

import fitz
import numpy as np

from procesador import (
    _RapidOcrAdapter,
    _corregir_filas_extracto_por_saldos,
    _elegir_monto_y_saldo_extracto,
    _parsear_fecha,
    _procesar_un_pdf_extracto,
    _resolver_saldo_ocr,
    _slug_hint_extracto,
)


class TestOcrSaldo71Vs11(unittest.TestCase):
    def test_resuelve_71_a_11_si_cierra_el_debito(self):
        # D Amico 30/10: movimiento 272.250, saldo OCR 71.201.888,77, previo 11.474.138,77
        saldo = _resolver_saldo_ocr(71_201_888.77, 11_474_138.77, 272_250.0)
        self.assertEqual(saldo, 11_201_888.77)

    def test_elige_272250_y_no_60_millones(self):
        mov, saldo = _elegir_monto_y_saldo_extracto(
            [272_250.0, 71_201_888.77],
            saldo_prev=11_474_138.77,
        )
        self.assertEqual(mov, -272_250.0)
        self.assertEqual(saldo, 11_201_888.77)

    def test_cadena_corrige_debito_60m_pegado(self):
        # Cache Gastro: crédito Lab Dominguez + débito D Amico
        movs = [
            {
                "Tipo fila": "Movimiento",
                "Descripcion": "IVA percepcion",
                "Importe": -1599.0,
                "Debito": 1599.0,
                "Credito": None,
                "Saldo": 10_779_874.30,
            },
            {
                "Tipo fila": "Movimiento",
                "Descripcion": "Pago a proveedores recibido",
                "Importe": 60_694_264.47,
                "Debito": None,
                "Credito": 60_694_264.47,
                "Saldo": 71_474_138.77,
            },
            {
                "Tipo fila": "Movimiento",
                "Descripcion": "TRANSFERENCIA REALIZADA",
                "Importe": -60_272_250.0,
                "Debito": 60_272_250.0,
                "Credito": None,
                "Saldo": 11_201_888.77,
            },
        ]
        out = _corregir_filas_extracto_por_saldos(movs)
        self.assertAlmostEqual(float(out[1]["Credito"]), 694_264.47, places=2)
        self.assertAlmostEqual(float(out[1]["Saldo"]), 11_474_138.77, places=2)
        self.assertAlmostEqual(float(out[2]["Debito"]), 272_250.0, places=2)
        self.assertAlmostEqual(float(out[2]["Saldo"]), 11_201_888.77, places=2)

    def test_alquiler_royo_960000(self):
        mov, saldo = _elegir_monto_y_saldo_extracto(
            [960_000.0, 10_854_201.36],
            saldo_prev=11_814_201.36,
        )
        self.assertEqual(mov, -960_000.0)
        self.assertEqual(saldo, 10_854_201.36)

    def test_cadena_arba_71m_luego_royo(self):
        movs = [
            {
                "Tipo fila": "Movimiento",
                "Descripcion": "Retencion arba",
                "Importe": 59_942_823.90,
                "Debito": None,
                "Credito": 59_942_823.90,
                "Saldo": 71_814_201.36,
            },
            {
                "Tipo fila": "Movimiento",
                "Descripcion": "Transferencia inmediata",
                "Importe": -60_960_000.0,
                "Debito": 60_960_000.0,
                "Credito": None,
                "Saldo": 10_854_201.36,
            },
        ]
        out = _corregir_filas_extracto_por_saldos(movs)
        self.assertAlmostEqual(float(out[0]["Saldo"]), 11_814_201.36, places=2)
        self.assertAlmostEqual(float(out[1]["Debito"]), 960_000.0, places=2)

    def test_usa_importe_chico_de_la_leyenda(self):
        movs = [
            {
                "Tipo fila": "Movimiento",
                "Descripcion": "IVA",
                "Importe": -100.0,
                "Debito": 100.0,
                "Credito": None,
                "Saldo": 17_589_304.15,
            },
            {
                "Tipo fila": "Movimiento",
                "Descripcion": "TRANSFERENCIA REALIZADA",
                "Detalle": "2.853,47 pensar digital srl varios var 30715823663",
                "Importe": -17_589_204.15,
                "Debito": 17_589_204.15,
                "Credito": None,
                "Saldo": 17.44,
            },
        ]
        out = _corregir_filas_extracto_por_saldos(movs)
        self.assertAlmostEqual(float(out[1]["Debito"]), 2853.47, places=2)


class TestHintBancoExtracto(unittest.TestCase):
    def test_slug_desde_selector_web(self):
        self.assertEqual(_slug_hint_extracto("Banco Galicia"), "galicia")
        self.assertEqual(_slug_hint_extracto("Banco Provincia"), "provincia")
        self.assertEqual(_slug_hint_extracto("BBVA"), "frances")
        self.assertEqual(_slug_hint_extracto(""), "")

    def test_pdf_sin_banco_en_el_nombre_usa_el_hint(self):
        doc = fitz.open()
        page = doc.new_page()
        page.insert_text(
            (72, 72),
            "01/01/2026 TRANSFERENCIA DE TERCEROS ACME S.A. 1.000,00 5.000,00",
        )
        data = doc.tobytes()
        doc.close()
        filas, meta, err = _procesar_un_pdf_extracto(
            "01-2026.pdf", data, banco_hint="Banco Galicia"
        )
        self.assertIsNone(err, msg=str(err))
        self.assertTrue(filas)
        self.assertEqual(meta.get("banco_slug"), "galicia")

    def test_fecha_2026_es_valida(self):
        self.assertEqual(_parsear_fecha("01/01/2026").year, 2026)
        self.assertEqual(_parsear_fecha("15/01/26").year, 2026)

    def test_rapidocr_adapter_tuplas(self):
        class _Fake:
            def __call__(self, _img):
                box = [[0, 0], [10, 0], [10, 10], [0, 10]]
                return ([[box, "SANTANDER", 0.99]], 0.01)

        filas = _RapidOcrAdapter(_Fake()).readtext(None)
        self.assertEqual(filas[0][1], "SANTANDER")

    def test_rapidocr_adapter_numpy_output_no_revienta(self):
        box = np.array([[0, 0], [10, 0], [10, 10], [0, 10]], dtype=float)

        class _Out:
            boxes = np.stack([box, box])
            txts = np.array(["SANTANDER", "02/02/26"])
            scores = np.array([0.99, 0.95])

        class _FakeObj:
            def __call__(self, _img):
                return _Out()

        filas = _RapidOcrAdapter(_FakeObj()).readtext(None)
        self.assertEqual([t for _, t, _ in filas], ["SANTANDER", "02/02/26"])

        class _FakeTuple:
            def __call__(self, _img):
                return (np.stack([box]), np.array(["SALDO"]), np.array([0.9]))

        filas2 = _RapidOcrAdapter(_FakeTuple()).readtext(None)
        self.assertEqual(filas2[0][1], "SALDO")

    def test_bbox_poligono_y_rectangulo(self):
        from procesador import _bbox_xy

        x, y = _bbox_xy([[0, 10], [20, 10], [20, 30], [0, 30]])
        self.assertEqual(x, 0.0)
        self.assertEqual(y, 20.0)
        x2, y2 = _bbox_xy([0, 10, 20, 30])
        self.assertEqual(x2, 0.0)
        self.assertEqual(y2, 20.0)
        x3, y3 = _bbox_xy(np.array([[0, 10], [20, 10], [20, 30], [0, 30]], dtype=float))
        self.assertEqual(x3, 0.0)
        self.assertEqual(y3, 20.0)

    def test_santander_une_fecha_e_importe_en_la_misma_linea(self):
        from procesador import _parsear_movimientos_santander_paginas, _score_filas_extracto

        texto = (
            "Fecha Comprobante Movimiento Débito Crédito Saldo en cuenta\n"
            "30/08/25 Saldo Inicial $ 17.288.556,31\n"
            "02/09/25 23886 Pago de haberes por cci $ 750.000,00 $ 16.538.556,31\n"
            "23886 Pago de haberes por cci $ 277.150,00 $ 16.261.406,31\n"
            "02/09/25 29143416 Pago haberes $ 500.000,00 $ 15.761.406,31\n"
        )
        movs, _meta = _parsear_movimientos_santander_paginas([(1, texto)], "09-2025.pdf")
        chain, con_monto, n = _score_filas_extracto(movs)
        self.assertGreaterEqual(n, 3)
        self.assertGreaterEqual(con_monto, 2)
        self.assertGreaterEqual(chain, 2)
        self.assertEqual(movs[0].get("Descripcion"), "Saldo Inicial")
        self.assertAlmostEqual(float(movs[1].get("Debito") or 0), 750000.0)


class TestConvertidorSinSaldos(unittest.TestCase):
    def test_excel_solo_fecha_concepto_debitos_creditos(self):
        from io import BytesIO

        import pandas as pd
        from openpyxl import load_workbook

        from procesador import df_extracto_convertidor_sin_saldos, exportar_extracto_bancario_excel

        df = pd.DataFrame(
            [
                {
                    "Tipo fila": "Saldo inicial",
                    "Fecha": "01/09/25",
                    "Descripcion": "Saldo Inicial",
                    "Debito": None,
                    "Credito": None,
                    "Saldo": 1000,
                },
                {
                    "Tipo fila": "Movimiento",
                    "Fecha": "02/09/25",
                    "Descripcion": "Pago haberes",
                    "Debito": 750000,
                    "Credito": None,
                    "Saldo": 250000,
                },
                {
                    "Tipo fila": "Movimiento",
                    "Fecha": "03/09/25",
                    "Descripcion": "Transferencia recibida",
                    "Debito": None,
                    "Credito": 100,
                    "Saldo": 250100,
                },
                {
                    "Tipo fila": "Saldo final",
                    "Fecha": "30/09/25",
                    "Descripcion": "Saldo Final",
                    "Debito": None,
                    "Credito": None,
                    "Saldo": 250100,
                },
            ]
        )
        out = df_extracto_convertidor_sin_saldos(df)
        self.assertEqual(list(out.columns), ["Fecha", "Concepto", "Débitos", "Créditos"])
        self.assertEqual(len(out), 2)
        self.assertNotIn("Saldo", out.columns)
        self.assertFalse(out["Concepto"].str.contains("Saldo", case=False).any())
        self.assertEqual(len(out), 2)

        df_ant = pd.DataFrame(
            [
                {
                    "Tipo fila": "Movimiento",
                    "Fecha": "30/12/25",
                    "Descripcion": "SALDO ANTERIOR",
                    "Debito": None,
                    "Credito": 1327591.30,
                    "Saldo": 1327591.30,
                },
                {
                    "Tipo fila": "Movimiento",
                    "Fecha": "02/01/26",
                    "Descripcion": "CR.DEBIN",
                    "Debito": None,
                    "Credito": 5941.92,
                    "Saldo": 1333533.22,
                },
            ]
        )
        out_ant = df_extracto_convertidor_sin_saldos(df_ant)
        self.assertEqual(len(out_ant), 1)
        self.assertEqual(out_ant.iloc[0]["Concepto"], "CR.DEBIN")

        xlsx = exportar_extracto_bancario_excel(df, {"banco": "Santander"})
        wb = load_workbook(BytesIO(xlsx))
        self.assertIn("Movimientos", wb.sheetnames)
        ws_m = wb["Movimientos"]
        hdr = None
        for row in ws_m.iter_rows(min_row=1, max_row=12, max_col=6, values_only=True):
            vals = [str(v) for v in row if v]
            if "Fecha" in vals and "Concepto" in vals:
                hdr = list(row)
                break
        self.assertIsNotNone(hdr)
        self.assertIn("Fecha", hdr)
        self.assertIn("Concepto", hdr)
        self.assertIn("Débitos", hdr)
        self.assertIn("Créditos", hdr)
        self.assertNotIn("Saldo", hdr)


class TestFormatosExtractoBancos(unittest.TestCase):
    def test_debin_suelto_no_arranca_movimiento(self):
        from procesador import _parece_inicio_movimiento_extracto

        self.assertFalse(
            _parece_inicio_movimiento_extracto("DEBIN", con_fecha_previa=True)
        )
        self.assertFalse(
            _parece_inicio_movimiento_extracto("IVA", con_fecha_previa=True)
        )
        self.assertTrue(
            _parece_inicio_movimiento_extracto(
                "CR.DEBIN CLINICA DEL SOL 30717847810", con_fecha_previa=True
            )
        )

    def test_santander_no_inventa_sin_descripcion_con_saldo(self):
        from procesador import _parsear_movimientos_santander_paginas

        texto = (
            "Fecha Comprobante Movimiento Débito Crédito Saldo en cuenta\n"
            "01/08/26 $ 34.154.678,27\n"
            "02/08/26 12345 CR.DEBIN CLINICA DEL SOL 30717847810 $ 15.000,00 $ 34.169.678,27\n"
            "DEBIN\n"
            "03/08/26 12346 TRANSFERENCIA DE TERCEROS ACME SA $ 1.000,00 $ 34.170.678,27\n"
        )
        movs, _meta = _parsear_movimientos_santander_paginas([(1, texto)], "08-2026.pdf")
        descs = [str(m.get("Descripcion") or "") for m in movs]
        self.assertFalse(any(d.lower() in {"sin descripción", "sin descripcion", ""} for d in descs))
        self.assertFalse(any(d.strip().upper() == "DEBIN" for d in descs))
        self.assertTrue(any("CR.DEBIN" in d.upper() and "CLINICA" in d.upper() for d in descs))
        self.assertEqual(len(movs), 2)

    def test_santander_conserva_saldo_inicial_etiquetado(self):
        from procesador import _parsear_movimientos_santander_paginas

        texto = (
            "30/08/25 Saldo Inicial $ 17.288.556,31\n"
            "02/09/25 23886 Pago de haberes por cci $ 750.000,00 $ 16.538.556,31\n"
        )
        movs, _meta = _parsear_movimientos_santander_paginas([(1, texto)], "09-2025.pdf")
        self.assertEqual(movs[0].get("Descripcion"), "Saldo Inicial")
        self.assertEqual(movs[0].get("Tipo fila"), "Saldo inicial")

    def test_provincia_debin_completo(self):
        from procesador import _parsear_movimientos_provincia_paginas

        texto = (
            "02/01/26 DB.DEBIN PRESTADORES MDP SA 30712345678 15.000,00 34.139.678,27\n"
            "03/01/26 TRANSFERENCIA DE TERCEROS VENOM SA 51.000,00 34.190.678,27\n"
        )
        movs, _meta = _parsear_movimientos_provincia_paginas([(1, texto)], "provincia.pdf")
        self.assertGreaterEqual(len(movs), 1)
        self.assertIn("DEBIN", (movs[0].get("Descripcion") or "").upper())
        self.assertNotEqual((movs[0].get("Descripcion") or "").strip().upper(), "DEBIN")

    def test_nacion_dh_no_toma_saldo_como_movimiento(self):
        from procesador import _extraer_movimientos_desde_texto

        lineas = [
            "01/08/26 SALDO ANTERIOR 34.154.678,27",
            "02/08/26 TRANSF. CTAS. PROPIAS 1.000,00 D 34.153.678,27",
            "03/08/26 ACREDIT. HABERES 50.000,00 H 34.203.678,27",
        ]
        movs = _extraer_movimientos_desde_texto(lineas, 1, "nacion", "bna.pdf")
        descs = [m.descripcion.upper() for m in movs]
        self.assertFalse(any("SALDO ANTERIOR" in d for d in descs))
        self.assertGreaterEqual(len(movs), 2)

    def test_clasifica_debin_como_deudores_a_identificar(self):
        from procesador import clasificar_movimiento_extracto

        self.assertEqual(
            clasificar_movimiento_extracto("CR.DEBIN CLINICA DEL SOL", "", 15000),
            "Deudores por ventas a identificar",
        )

    def test_score_prefiere_descripciones_largas(self):
        from procesador import _elegir_mejor_parseo_extracto

        cortos = [
            {
                "Descripcion": "DEBIN",
                "Debito": 100.0,
                "Credito": 0.0,
                "Saldo": 900.0,
            }
        ] * 20
        largos = [
            {
                "Descripcion": "CR.DEBIN CLINICA DEL SOL 30717847810",
                "Debito": 0.0,
                "Credito": 100.0,
                "Saldo": 1100.0,
            },
            {
                "Descripcion": "TRANSFERENCIA DE TERCEROS ACME SA",
                "Debito": 50.0,
                "Credito": 0.0,
                "Saldo": 1050.0,
            },
        ]
        _movs, _meta, tag = _elegir_mejor_parseo_extracto(
            (cortos, {}, "explotado"),
            (largos, {}, "layout_y"),
        )
        self.assertEqual(tag, "layout_y")

    def test_partir_debin_deja_contraparte(self):
        from procesador import _partir_concepto_y_detalle

        desc, _det = _partir_concepto_y_detalle(
            "CR.DEBIN CLINICA DEL SOL 30717847810", None
        )
        self.assertIn("CLINICA", desc.upper())
        self.assertIn("30717847810", desc)


if __name__ == "__main__":
    unittest.main()

