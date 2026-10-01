#!/usr/bin/env python3
"""Genera Facturas_a_emitir_MODELO.xlsx (hoja visible 'Facturas' + hoja oculta 'Codigos')."""
import datetime as dt
import os
import sys

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

BASE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, BASE)
from arca.codigos import (ALICUOTAS, COND_IVA_RECEPTOR, CONCEPTOS, DOC_TIPOS, TIPOS_CBTE,  # noqa: E402
                          TRANSFERENCIA_FCE)
from arca.comprobante import COLUMNAS, COND_IVA_EMISOR, COND_VENTA  # noqa: E402

FILA_ENC, PRIMERA, ULTIMA = 3, 4, 500
COL = {k: get_column_letter(i) for i, (k, _, _) in enumerate(COLUMNAS, 1)}
HOY = dt.date(2026, 10, 1)


def crear(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Facturas"
    cod = wb.create_sheet("Codigos")

    # ---------- Hoja oculta de códigos ----------
    tablas = [
        ("Tipo comprobante", "Código ARCA", [(k, v[0]) for k, v in TIPOS_CBTE.items()], "L_Tipo"),
        ("Concepto", "Código", list(CONCEPTOS.items()), "L_Concepto"),
        ("Doc. tipo", "Código", list(DOC_TIPOS.items()), "L_Doc"),
        ("Condición IVA receptor", "Código", [(d, k) for k, (d, _) in COND_IVA_RECEPTOR.items()], "L_CondIVA"),
        ("Alícuota IVA", "Tasa", [(k, v[1]) for k, v in ALICUOTAS.items()], "L_Alic"),
        ("Condición de venta", "", [(x, "") for x in COND_VENTA], "L_CondVenta"),
        ("Condición IVA emisor", "", [(x, "") for x in COND_IVA_EMISOR], "L_CondEmisor"),
        ("Transferencia FCE", "Código", list(TRANSFERENCIA_FCE.items()), "L_Transf"),
        ("Marca", "", [("EJEMPLO", ""), ("NO EMITIR", "")], "L_Marca"),
    ]
    col = 1
    for t1, t2, filas, nombre in tablas:
        cod.cell(1, col, t1).font = Font(bold=True)
        cod.cell(1, col + 1, t2).font = Font(bold=True)
        for i, (a, b) in enumerate(filas, 2):
            cod.cell(i, col, a)
            cod.cell(i, col + 1, b)
        L1, L2 = get_column_letter(col), get_column_letter(col + 1)
        ult = len(filas) + 1
        wb.defined_names[nombre] = DefinedName(nombre, attr_text=f"Codigos!${L1}$2:${L1}${ult}")
        if nombre == "L_Alic":
            wb.defined_names["T_Alic"] = DefinedName("T_Alic", attr_text=f"Codigos!${L1}$2:${L2}${ult}")
        cod.column_dimensions[L1].width = max(len(str(a)) for a, _ in filas) + 2
        col += 3
    # Matriz condición IVA vs clase (referencia, manual WSFEv1 v4.7 Anexo)
    cod.cell(1, col, "Cond. IVA admitida por clase (Anexo manual WSFEv1)").font = Font(bold=True)
    for j, h in enumerate(["Código", "Descripción", "A", "B", "C"]):
        cod.cell(2, col + j, h).font = Font(bold=True)
    for i, (k, (d, cl)) in enumerate(COND_IVA_RECEPTOR.items(), 3):
        cod.cell(i, col, k)
        cod.cell(i, col + 1, d)
        for j, c in enumerate("ABC"):
            cod.cell(i, col + 2 + j, "X" if c in cl else "")
    cod.sheet_state = "hidden"

    # ---------- Hoja Facturas ----------
    azul = PatternFill("solid", fgColor="1F4E78")
    gris = PatternFill("solid", fgColor="F2F2F2")
    amarillo = PatternFill("solid", fgColor="FFF2CC")
    fino = Side(style="thin", color="BFBFBF")
    borde = Border(left=fino, right=fino, top=fino, bottom=fino)
    nc = len(COLUMNAS)
    ws.cell(1, 1, "Facturas a emitir por Web Service ARCA (WSFEv1) — Estudio Zona Güemes").font = Font(bold=True, size=14, color="1F4E78")
    ws.cell(2, 1, ("Una fila por comprobante (o varias filas con el mismo 'ID factura' para un comprobante con varios ítems/alícuotas). "
                   "Filas marcadas EJEMPLO no se emiten. Columnas grises = fórmula, no editar. "
                   "Para C cargar el importe total y alícuota 'No corresponde (C)'. Para A/B cargar el NETO; IVA y Total se calculan. "
                   "Servicios: completar fechas desde/hasta y Vto. pago. El resultado (CAE, nro) se escribe en la hoja 'Resultado'.")).alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=nc)
    ws.row_dimensions[2].height = 45
    ayudas = {
        "cuit_emisor": "CUIT del cliente que factura (11 dígitos, sin guiones). Debe haber delegado 'Facturación Electrónica' (wsfe) a la CUIT 23-42284343-4.",
        "pto_vta": "Punto de venta dado de alta para Web Services (RECE para aplicativo y web services / Factura Electrónica - Monotributo - Web Services). NO usar el de Comprobantes en línea.",
        "tipo": "C = monotributista/exento (cód. 11). B = RI a consumidor final/exento (cód. 6). A = RI a RI/monotributista (cód. 1). FCE = Factura de Crédito MiPyME (201/206/211). NC/ND = notas de crédito/débito (A 3/2, B 8/7, C 13/12): requieren comprobante asociado.",
        "concepto": "Productos (1), Servicios (2) o Productos y Servicios (3). Servicios exige fechas desde/hasta y Vto. pago.",
        "fecha": "Productos: ±5 días de hoy. Servicios: ±10 días. FCE: -5/+1 día.",
        "cond_iva": "Obligatoria (RG 5616). A admite: RI, Monotributo, Monotributista Social, Monotr. Trab. Indep. Promovido. B admite: Exento, Consumidor Final, No Categorizado, Exterior, Liberado 19.640, No Alcanzado. C admite todas.",
        "doc_tipo": "A y FCE exigen CUIT. Consumidor final sin identificar sólo si total < $10.000.000 (RG 5824/2026).",
        "neto": "A/B: importe NETO gravado (sin IVA). C: importe total.",
        "alicuota": "A/B: 21%, 10,5%, 27%, 0%, Exento o No gravado. C: 'No corresponde (C)'.",
        "id": "Opcional. Filas con el mismo ID y mismo CUIT emisor forman UN comprobante (varios ítems). Dejar vacío para 1 fila = 1 factura.",
        "cbu": "Sólo FCE: CBU del emisor (22 dígitos).",
        "asoc_id": "Sólo NC/ND. Opción 1: el 'ID factura' de otra fila de este Excel (se emite primero y se usa el número que devuelve ARCA) o de una corrida anterior (hoja Resultado). Si se usa esto, dejar vacíos tipo/PtoVta/Nro/fecha.",
        "asoc_tipo": "Sólo NC/ND. Opción 2: comprobante ya emitido fuera de este Excel. Tipo (misma letra: NC C asocia C/ND C/NC C).",
        "asoc_pto": "Punto de venta del comprobante asociado.",
        "asoc_nro": "Número del comprobante asociado (sin punto de venta).",
        "asoc_fecha": "Fecha del comprobante asociado (al enviar se toma la de ARCA vía FECompConsultar).",
    }
    for i, (k, h, w) in enumerate(COLUMNAS, 1):
        c = ws.cell(FILA_ENC, i, h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = azul
        c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        c.border = borde
        ws.column_dimensions[get_column_letter(i)].width = w
        if k in ayudas:
            c.comment = Comment(ayudas[k], "Estudio ZG", width=320, height=140)
    ws.row_dimensions[FILA_ENC].height = 48
    ws.freeze_panes = "D4"  # congela encabezado y columnas Marca, ID y CUIT emisor

    # Fórmulas, formatos y estilo de todas las filas de carga
    for r in range(PRIMERA, ULTIMA + 1):
        N, T, I = f"{COL['neto']}{r}", f"{COL['alicuota']}{r}", f"{COL['tipo']}{r}"
        ws[f"{COL['iva']}{r}"] = f'=IF(OR({N}="",{T}=""),"",IF(RIGHT({I},1)="C",0,ROUND({N}*IFERROR(VLOOKUP({T},T_Alic,2,FALSE),0),2)))'
        ws[f"{COL['total']}{r}"] = f'=IF({N}="","",{N}+IF({COL["iva"]}{r}="",0,{COL["iva"]}{r}))'
        for k in ("iva", "total"):
            ws[f"{COL[k]}{r}"].fill = gris
        for k in ("neto", "iva", "total"):
            ws[f"{COL[k]}{r}"].number_format = '#,##0.00'
        for k in ("fecha", "fsd", "fsh", "fvto", "inicio_act", "asoc_fecha"):
            ws[f"{COL[k]}{r}"].number_format = "dd/mm/yyyy"
        for k in ("cuit_emisor", "doc_nro", "cbu"):
            ws[f"{COL[k]}{r}"].number_format = "0"
        for i in range(1, nc + 1):
            ws.cell(r, i).border = borde

    # Desplegables
    rng = lambda k: f"{COL[k]}{PRIMERA}:{COL[k]}{ULTIMA}"  # noqa: E731
    for k, nombre, titulo in [("tipo", "L_Tipo", "Tipo"), ("concepto", "L_Concepto", "Concepto"),
                              ("doc_tipo", "L_Doc", "Doc. tipo"), ("cond_iva", "L_CondIVA", "Condición IVA"),
                              ("alicuota", "L_Alic", "Alícuota"), ("cond_venta", "L_CondVenta", "Cond. venta"),
                              ("cond_iva_emisor", "L_CondEmisor", "Condición IVA emisor"),
                              ("fce_transf", "L_Transf", "Transferencia FCE"), ("marca", "L_Marca", "Marca"),
                              ("asoc_tipo", "L_Tipo", "Cbte asociado tipo")]:
        dv = DataValidation(type="list", formula1=f"={nombre}", allow_blank=True, showErrorMessage=True,
                            errorTitle=titulo, error="Elegir un valor de la lista.")
        dv.add(rng(k))
        ws.add_data_validation(dv)
    for k in ("fecha", "fsd", "fsh", "fvto", "asoc_fecha"):
        dv = DataValidation(type="date", operator="greaterThan", formula1="DATE(2020,1,1)", allow_blank=True,
                            showErrorMessage=True, error="Fecha inválida (dd/mm/aaaa).")
        dv.add(rng(k))
        ws.add_data_validation(dv)
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2="99999", allow_blank=True,
                        showErrorMessage=True, error="Punto de venta 1 a 99999.")
    dv.add(rng("pto_vta"))
    dv.add(rng("asoc_pto"))
    ws.add_data_validation(dv)
    dv = DataValidation(type="decimal", operator="greaterThan", formula1="0", allow_blank=True,
                        showErrorMessage=True, error="Importe mayor a 0.")
    dv.add(rng("neto"))
    ws.add_data_validation(dv)
    dv = DataValidation(type="textLength", operator="equal", formula1="11", allow_blank=True,
                        showErrorMessage=True, error="CUIT: 11 dígitos sin guiones.")
    dv.add(rng("cuit_emisor"))
    ws.add_data_validation(dv)

    # Formato condicional: filas EJEMPLO en amarillo; condición IVA incompatible con la clase en rojo
    M = COL["marca"]
    ws.conditional_formatting.add(f"A{PRIMERA}:{get_column_letter(nc)}{ULTIMA}",
                                  FormulaRule(formula=[f'${M}{PRIMERA}="EJEMPLO"'], fill=amarillo))
    I, CI = COL["tipo"], COL["cond_iva"]
    claseA = '{"IVA Responsable Inscripto","Responsable Monotributo","Monotributista Social","Monotributo Trabajador Independiente Promovido"}'
    f_err = (f'AND(${CI}{PRIMERA}<>"",${I}{PRIMERA}<>"",RIGHT(${I}{PRIMERA},1)<>"C",'
             f'OR(AND(RIGHT(${I}{PRIMERA},1)="A",ISNA(MATCH(${CI}{PRIMERA},{claseA},0))),'
             f'AND(RIGHT(${I}{PRIMERA},1)="B",ISNUMBER(MATCH(${CI}{PRIMERA},{claseA},0)))))')
    ws.conditional_formatting.add(f"{CI}{PRIMERA}:{CI}{ULTIMA}",
                                  FormulaRule(formula=[f_err], fill=PatternFill("solid", fgColor="FFC7CE"),
                                              font=Font(color="9C0006", bold=True)))

    # ---------- Filas de ejemplo (datos ficticios) ----------
    d = lambda y, m, dd: dt.date(y, m, dd)  # noqa: E731
    mono = dict(cuit_emisor="27301234568", razon_emisor="EJEMPLO Pérez Ana (monotributo)",
                dom_emisor="Güemes 1234, Mar del Plata", iibb_emisor="27301234568", inicio_act=d(2019, 3, 1), pto_vta=3,
                cond_iva_emisor="Responsable Monotributo")
    ri = dict(cuit_emisor="30712345671", razon_emisor="EJEMPLO Servicios del Sur SRL",
              dom_emisor="Av. Colón 2500, Mar del Plata", iibb_emisor="901-123456-7", inicio_act=d(2015, 8, 10), pto_vta=5,
              cond_iva_emisor="IVA Responsable Inscripto")
    ejemplos = [
        {**mono, "id": "EJ-C1", "tipo": "C", "concepto": "Servicios", "fecha": HOY, "cond_venta": "Transferencia bancaria",
         "doc_tipo": "DNI", "doc_nro": "30111222", "nombre_rec": "EJEMPLO Gómez Juan", "dom_rec": "Alberti 100, MdP",
         "cond_iva": "Consumidor Final", "detalle": "Honorarios septiembre 2026", "neto": 150000,
         "alicuota": "No corresponde (C)", "fsd": d(2026, 9, 1), "fsh": d(2026, 9, 30), "fvto": d(2026, 10, 10)},
        {**mono, "tipo": "C", "concepto": "Servicios", "fecha": HOY, "cond_venta": "Cuenta corriente",
         "doc_tipo": "CUIT", "doc_nro": "30709998885", "nombre_rec": "EJEMPLO Distribuidora Atlántica SA",
         "dom_rec": "Juan B. Justo 4000, MdP", "cond_iva": "IVA Responsable Inscripto",
         "detalle": "Mantenimiento mensual", "neto": 480000, "alicuota": "No corresponde (C)",
         "fsd": d(2026, 9, 1), "fsh": d(2026, 9, 30), "fvto": d(2026, 10, 15)},
        {**ri, "tipo": "B", "concepto": "Productos", "fecha": HOY, "cond_venta": "Contado",
         "doc_tipo": "Consumidor Final (sin identificar)", "doc_nro": "0", "nombre_rec": "",
         "cond_iva": "Consumidor Final", "detalle": "Mercadería varias", "neto": 82644.63, "alicuota": "21%"},
        {**ri, "id": "EJ-A1", "tipo": "A", "concepto": "Productos y Servicios", "fecha": HOY,
         "cond_venta": "Cuenta corriente", "doc_tipo": "CUIT", "doc_nro": "30709998885",
         "nombre_rec": "EJEMPLO Distribuidora Atlántica SA", "dom_rec": "Juan B. Justo 4000, MdP",
         "cond_iva": "IVA Responsable Inscripto", "detalle": "Servicio técnico", "neto": 200000, "alicuota": "21%",
         "fsd": d(2026, 9, 1), "fsh": d(2026, 9, 30), "fvto": d(2026, 10, 20)},
        {**ri, "id": "EJ-A1", "tipo": "A", "concepto": "Productos y Servicios", "fecha": HOY,
         "cond_venta": "Cuenta corriente", "doc_tipo": "CUIT", "doc_nro": "30709998885",
         "nombre_rec": "EJEMPLO Distribuidora Atlántica SA", "dom_rec": "Juan B. Justo 4000, MdP",
         "cond_iva": "IVA Responsable Inscripto", "detalle": "Repuestos (bienes de capital)", "neto": 50000,
         "alicuota": "10,5%", "fsd": d(2026, 9, 1), "fsh": d(2026, 9, 30), "fvto": d(2026, 10, 20)},
        {**ri, "tipo": "A", "concepto": "Servicios", "fecha": HOY, "cond_venta": "Transferencia bancaria",
         "doc_tipo": "CUIT", "doc_nro": "20334445551", "nombre_rec": "EJEMPLO López Marta (monotributo)",
         "dom_rec": "San Martín 2200, MdP", "cond_iva": "Responsable Monotributo",
         "detalle": "Provisión de gas/telefonía (27%)", "neto": 10000, "alicuota": "27%",
         "fsd": d(2026, 9, 1), "fsh": d(2026, 9, 30), "fvto": d(2026, 10, 10)},
        {**ri, "tipo": "FCE A", "concepto": "Productos", "fecha": HOY, "cond_venta": "Cuenta corriente",
         "doc_tipo": "CUIT", "doc_nro": "30500001115", "nombre_rec": "EJEMPLO Gran Empresa SA",
         "dom_rec": "Av. Corrientes 1000, CABA", "cond_iva": "IVA Responsable Inscripto",
         "detalle": "Venta mayorista", "neto": 5000000, "alicuota": "21%", "fvto": d(2026, 10, 31),
         "cbu": "0110599520000012345678", "fce_transf": "SCA - Transferencia al Sistema de Circulación Abierta"},
        {**mono, "tipo": "NC C", "concepto": "Servicios", "fecha": HOY, "cond_venta": "Transferencia bancaria",
         "doc_tipo": "DNI", "doc_nro": "30111222", "nombre_rec": "EJEMPLO Gómez Juan", "dom_rec": "Alberti 100, MdP",
         "cond_iva": "Consumidor Final", "detalle": "Bonificación honorarios septiembre (NC de la fila EJ-C1)",
         "neto": 50000, "alicuota": "No corresponde (C)", "fsd": d(2026, 9, 1), "fsh": d(2026, 9, 30),
         "fvto": d(2026, 10, 10), "asoc_id": "EJ-C1"},
        {**ri, "tipo": "NC A", "concepto": "Productos", "fecha": HOY, "cond_venta": "Cuenta corriente",
         "doc_tipo": "CUIT", "doc_nro": "30709998885", "nombre_rec": "EJEMPLO Distribuidora Atlántica SA",
         "dom_rec": "Juan B. Justo 4000, MdP", "cond_iva": "IVA Responsable Inscripto",
         "detalle": "Devolución mercadería (NC de factura emitida antes)", "neto": 20000, "alicuota": "21%",
         "asoc_tipo": "A", "asoc_pto": 5, "asoc_nro": 123, "asoc_fecha": d(2026, 9, 25)},
    ]
    for j, ej in enumerate(ejemplos):
        r = PRIMERA + j
        ej["marca"] = "EJEMPLO"
        for k, v in ej.items():
            ws[f"{COL[k]}{r}"] = v
    ws.auto_filter.ref = f"A{FILA_ENC}:{get_column_letter(nc)}{ULTIMA}"
    ws.sheet_view.zoomScale = 90
    wb.active = 0
    wb.save(path)
    print("Modelo creado:", path)


if __name__ == "__main__":
    crear(sys.argv[1] if len(sys.argv) > 1 else os.path.join(BASE, "plantillas", "Facturas_a_emitir_MODELO.xlsx"))
