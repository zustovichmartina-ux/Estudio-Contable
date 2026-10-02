"""Reporte de Impuestos - Ganancias (Personas Humanas), H. Trujillo & Asociados.

Puerto a la app de la lógica de ModuloReporte.bas / _simular_macro.py:
  - Ubica en la hoja "2026" de la Proyección del cliente las filas de Ventas,
    Gastos, Retenciones y Anticipos por su título en la columna C.
  - Toma la escala del Art. 94 de la misma hoja.
  - Cruza con los topes de deducciones de Tango (acumulados por mes) y la
    ficha del cliente (cónyuge, hijos, hijos incapacitados, tipo de
    deducción especial), guardados en Turso.
  - Calcula mes a mes: resultado neto, deducciones, ganancia neta sujeta a
    impuesto, impuesto determinado (escala Art. 94 prorrateada) y saldo a
    ingresar / a favor.
  - Arma el Excel de salida (solapas "Topes 2026" + "Reporte", con fórmulas
    vivas idénticas a las de la macro) a partir de una Plantilla embebida.

Coeficientes de deducción especial (Art. 30 inc. c), por defecto:
  Ap. 1 - Autonomos            = MNI x 3.5
  Ap. 1 - Nuevos profesionales = MNI x 4
  Ap. 2 - Relacion de dependencia = "Deducción Especial" de Tango
(tomados de la solapa Config del Generador; quedan para validar).
"""
from __future__ import annotations

import datetime
from copy import copy
from pathlib import Path
from typing import Optional

import openpyxl
from openpyxl.utils import get_column_letter as L

HOJA_PROY = "2026"
HOJA_TOPES = "Topes 2026"
ANIO = 2026
MESES = ["Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"]
MESES_LARGO = [
    "enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto",
    "septiembre", "octubre", "noviembre", "diciembre",
]

TIPOS_DEDUCCION_ESPECIAL = (
    "Ap. 1 - Autonomos",
    "Ap. 1 - Nuevos profesionales",
    "Ap. 2 - Relacion de dependencia",
    "No corresponde",
)

COEF_AUTONOMOS_DEFAULT = 3.5
COEF_NUEVOS_PROF_DEFAULT = 4.0

CONCEPTO_MNI = "M.N.I."
CONCEPTO_DEESP = "Deducción Especial"
CONCEPTO_CONYUGE = "Cónyuge / Unión Convivencial"
CONCEPTO_HIJOS = "Hijos"
CONCEPTO_HINC = "Hijos incapacitados para el trabajo"

T_PRIMERA, T_ULTIMA = 4, 23
T_CONY, T_HIJOS, T_HINC, T_DESP, T_COEF1, T_COEF2, T_MES = 27, 28, 29, 30, 31, 32, 33
T_RMNI, T_RDE, T_RCF = 36, 37, 38

PLANTILLA_PATH = Path(__file__).resolve().parent / "assets" / "plantilla_reporte_ganancias.xlsx"


class ErrorReporteGanancias(Exception):
    """Error esperable (archivo mal armado, faltan datos) para mostrar en la UI."""


# --------------------------------------------------------------------------- lectura proyección
def _norm(v) -> str:
    return " ".join(str(v or "").split()).upper()


def _fila_titulo(ws, texto: str, desde: int = 1, exacto: bool = False) -> Optional[int]:
    for r in range(desde, ws.max_row + 1):
        v = _norm(ws.cell(r, 3).value)
        if (v == texto) if exacto else (texto in v):
            return r
    return None


def localizar_filas_proyeccion(ws) -> dict:
    rV = _fila_titulo(ws, "RESULTADO BRUTO VENTAS")
    rG = _fila_titulo(ws, "TOTAL DE GASTOS")
    if not rV or not rG:
        raise ErrorReporteGanancias(
            "No encontré en la columna C de la hoja 2026 los títulos "
            "'RESULTADO BRUTO VENTAS' y/o 'TOTAL DE GASTOS'."
        )
    rR = _fila_titulo(ws, "RETENCIONES", rG)
    rA = _fila_titulo(ws, "ANTICIPOS", rG, exacto=True)
    if not rR or not rA:
        raise ErrorReporteGanancias(
            "No encontré en la columna C de la hoja 2026 los títulos "
            "'...RETENCIONES...' y/o 'ANTICIPOS'."
        )
    return {"rV": rV, "rG": rG, "rR": rR, "rA": rA}


def localizar_escala(ws) -> dict:
    hdr = None
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and "Ganancia neta imponible acumulada" in c.value:
                hdr = c
                break
        if hdr:
            break
    if hdr is None:
        raise ErrorReporteGanancias(
            "No encontré la escala del Art. 94 (título 'Ganancia neta imponible "
            "acumulada') en la hoja 2026."
        )
    eC, eR0 = hdr.column, hdr.row + 1
    eR1 = eR0
    while isinstance(ws.cell(eR1 + 1, eC).value, (int, float)):
        eR1 += 1
    tramos = []
    for r in range(eR0, eR1 + 1):
        desde = ws.cell(r, eC).value
        fijo = ws.cell(r, eC + 2).value
        pct = ws.cell(r, eC + 3).value
        excedente = ws.cell(r, eC + 4).value
        if desde is None:
            continue
        tramos.append(
            {
                "desde": float(desde),
                "fijo": float(fijo or 0),
                "pct": float(pct or 0),
                "excedente": float(excedente or 0),
            }
        )
    if not tramos:
        raise ErrorReporteGanancias("La escala del Art. 94 está vacía en la hoja 2026.")
    tramos.sort(key=lambda t: t["desde"])
    return {"col": eC, "r0": eR0, "r1": eR1, "tramos": tramos}


def _fila_valores(ws, fila: int) -> list[float]:
    return [float(ws.cell(fila, 4 + i).value or 0) for i in range(12)]


def _total_p(ws, fila: int) -> float:
    return float(ws.cell(fila, 16).value or 0)


# --------------------------------------------------------------------------- topes de Tango
def parsear_topes_tango(file_like) -> dict[str, list[Optional[float]]]:
    """Lee la exportación de Tango (solapa 'Deducciones': Período, Deducción,
    Importe tope, [Eliminado]) y arma {concepto: [12 valores acumulados]}."""
    wb = openpyxl.load_workbook(file_like, data_only=True)
    if "Deducciones" not in wb.sheetnames:
        raise ErrorReporteGanancias(
            "El archivo no tiene la hoja 'Deducciones' (exportación de Tango)."
        )
    sh = wb["Deducciones"]
    topes: dict[str, list[Optional[float]]] = {}
    for row in sh.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        periodo, concepto, importe = row[0], row[1], row[2]
        eliminado = row[3] if len(row) > 3 else None
        if isinstance(periodo, datetime.datetime):
            mm, yy = periodo.month, periodo.year
        else:
            p = str(periodo).strip()
            if len(p) == 7 and p[2] == "/":
                mm, yy = int(p[:2]), int(p[-4:])
            else:
                continue
        if yy != ANIO or not (1 <= mm <= 12):
            continue
        if str(eliminado or "").strip().upper() in ("SI", "SÍ"):
            continue
        if concepto is None:
            continue
        k = str(concepto).strip()
        arr = topes.setdefault(k, [None] * 12)
        arr[mm - 1] = float(importe) if importe is not None else None
    if not topes:
        raise ErrorReporteGanancias(f"El archivo de topes no tiene datos del año {ANIO}.")
    return topes


def _tope(topes: dict, concepto: str, mes_idx: int) -> float:
    arr = topes.get(concepto)
    if not arr:
        return 0.0
    v = arr[mes_idx]
    return float(v) if v is not None else 0.0


# --------------------------------------------------------------------------- cálculo puro (para pantalla)
def mes_informado(ventas: list[float]) -> int:
    n = sum(1 for v in ventas if v)
    return max(1, n)


def deduccion_especial_mensual(
    mni_mes: float, tipo_de: str, coef1: float, coef2: float, topes: dict, mes_idx: int
) -> float:
    if tipo_de == "Ap. 1 - Autonomos":
        return mni_mes * coef1
    if tipo_de == "Ap. 1 - Nuevos profesionales":
        return mni_mes * coef2
    if tipo_de == "Ap. 2 - Relacion de dependencia":
        return _tope(topes, CONCEPTO_DEESP, mes_idx)
    return 0.0


def _tramo_para(g_anual: float, tramos: list[dict]) -> dict:
    elegido = tramos[0]
    for t in tramos:
        if t["desde"] <= g_anual:
            elegido = t
        else:
            break
    return elegido


def impuesto_determinado(ganancia_acum: float, n: int, tramos: list[dict]) -> float:
    if ganancia_acum is None or ganancia_acum <= 0:
        return 0.0
    g_anual = ganancia_acum * 12 / n
    t = _tramo_para(g_anual, tramos)
    return (n / 12) * t["fijo"] + (ganancia_acum - (n / 12) * t["excedente"]) * t["pct"]


def calcular_reporte(
    ventas: list[float],
    gastos: list[float],
    retenciones: list[float],
    anticipos_mensual: list[float],
    anticipos_total: float,
    tramos: list[dict],
    topes: dict,
    conyuge: int,
    hijos: int,
    hijos_incapacitados: int,
    tipo_deduccion_especial: str,
    coef1: float = COEF_AUTONOMOS_DEFAULT,
    coef2: float = COEF_NUEVOS_PROF_DEFAULT,
) -> dict:
    """Devuelve el cuadro mes a mes + totales, igual que la solapa Reporte."""
    n = mes_informado(ventas)
    anticipos_solo_total = round(sum(anticipos_mensual), 2) != round(anticipos_total, 2)

    meses = []
    resultado_acum = 0.0
    retenciones_acum = 0.0
    anticipos_acum = 0.0
    fila_n = None
    for i in range(12):
        mes_num = i + 1
        v, g = ventas[i], gastos[i]
        resultado = v - g
        resultado_acum += resultado
        retenciones_acum += retenciones[i]
        anticipos_acum += anticipos_mensual[i]
        fila = {"mes": MESES[i], "ventas": v, "gastos": g, "resultado": resultado}
        if mes_num > n:
            fila.update(
                mni=None, de=None, cf=None, ganancia=None, impuesto=None,
                retenciones=None, anticipos=None, saldo=None,
            )
        else:
            mni = _tope(topes, CONCEPTO_MNI, i)
            de = deduccion_especial_mensual(mni, tipo_deduccion_especial, coef1, coef2, topes, i)
            cf = (
                conyuge * _tope(topes, CONCEPTO_CONYUGE, i)
                + hijos * _tope(topes, CONCEPTO_HIJOS, i)
                + hijos_incapacitados * _tope(topes, CONCEPTO_HINC, i)
            )
            ganancia = resultado_acum - mni - de - cf
            impuesto = impuesto_determinado(ganancia, n, tramos)
            anticipos_fila = None if anticipos_solo_total else anticipos_acum
            saldo = None if anticipos_solo_total else (impuesto - retenciones_acum - anticipos_fila)
            fila.update(
                mni=mni, de=de, cf=cf, ganancia=ganancia, impuesto=impuesto,
                retenciones=retenciones_acum, anticipos=anticipos_fila, saldo=saldo,
            )
            fila_n = fila
        meses.append(fila)

    totales = {
        "ventas": sum(ventas),
        "gastos": sum(gastos),
        "resultado": sum(ventas) - sum(gastos),
        "mni": fila_n["mni"] if fila_n else 0.0,
        "de": fila_n["de"] if fila_n else 0.0,
        "cf": fila_n["cf"] if fila_n else 0.0,
        "ganancia": fila_n["ganancia"] if fila_n else 0.0,
        "impuesto": fila_n["impuesto"] if fila_n else 0.0,
        "retenciones": fila_n["retenciones"] if fila_n else 0.0,
        "anticipos": anticipos_total,
        "mes_informado": n,
        "mes_informado_nombre": MESES_LARGO[n - 1],
    }
    totales["saldo"] = (totales["impuesto"] or 0.0) - (totales["retenciones"] or 0.0) - (totales["anticipos"] or 0.0)
    return {"meses": meses, "totales": totales}


# --------------------------------------------------------------------------- avisos
def avisos_reporte(topes: dict, n: int, tipo_de: str) -> list[str]:
    avisos = []
    tiene_datos_mes = any(
        (arr[n - 1] is not None) for arr in topes.values() if len(arr) >= n
    )
    if not tiene_datos_mes:
        avisos.append(
            "El mes informado no tiene topes oficiales en el archivo de Tango "
            "(puede estar proyectado con el último valor)."
        )
    if tipo_de == "No corresponde":
        avisos.append(
            "Falta indicar el tipo de deducción especial en la ficha del cliente."
        )
    return avisos


# --------------------------------------------------------------------------- Excel de salida
def _copiar_estilo(a, b):
    b.font, b.fill, b.border = copy(a.font), copy(a.fill), copy(a.border)
    b.number_format, b.alignment = a.number_format, copy(a.alignment)


def generar_excel(
    archivo_cliente,
    topes: dict,
    conyuge: int,
    hijos: int,
    hijos_incapacitados: int,
    tipo_deduccion_especial: str,
    coef1: float = COEF_AUTONOMOS_DEFAULT,
    coef2: float = COEF_NUEVOS_PROF_DEFAULT,
) -> bytes:
    """Arma el .xlsx de salida: la Proyección del cliente + solapas Topes 2026
    y Reporte, con las mismas fórmulas vivas que escribe la macro de Excel."""
    wb = openpyxl.load_workbook(archivo_cliente)
    if HOJA_PROY not in wb.sheetnames:
        raise ErrorReporteGanancias(f"El archivo no tiene la hoja '{HOJA_PROY}'.")
    s = wb[HOJA_PROY]

    filas = localizar_filas_proyeccion(s)
    rV, rG, rR, rA = filas["rV"], filas["rG"], filas["rR"], filas["rA"]
    escala = localizar_escala(s)
    eC, eR0, eR1 = escala["col"], escala["r0"], escala["r1"]

    def rng(col):
        return f"'{HOJA_PROY}'!${L(col)}${eR0}:${L(col)}${eR1}"

    eU, eW, eX, eY = rng(eC), rng(eC + 2), rng(eC + 3), rng(eC + 4)

    for nm in ("Reporte", "Informe", HOJA_TOPES):
        if nm in wb.sheetnames:
            del wb[nm]

    # --- Topes 2026 ------------------------------------------------------
    tp = wb.create_sheet(HOJA_TOPES)
    tp["B1"] = f"TOPES DE DEDUCCIONES PERSONALES {ANIO} - ACUMULADOS AL MES (fuente: Tango)"
    tp["B1"].font = openpyxl.styles.Font(bold=True, size=12)
    tp["B2"] = f"Actualizado el {datetime.date.today().strftime('%d/%m/%Y')} desde la app"
    tp["B2"].font = openpyxl.styles.Font(italic=True, size=8)
    tp["B3"] = "Deducción"
    for i, m in enumerate(MESES):
        tp.cell(3, 3 + i, m)
    tp["B3:N3"[:2]]  # no-op guard
    r = T_PRIMERA
    for k, arr in topes.items():
        tp.cell(r, 2, k)
        for i, v in enumerate(arr):
            if v is not None:
                tp.cell(r, 3 + i, v)
        r += 1
        if r > T_ULTIMA:
            break
    tp[f"B{T_CONY - 1}"] = "CONFIGURACIÓN DEL CONTRIBUYENTE"
    tp[f"B{T_CONY}"] = "Cónyuge / Unión convivencial (1 = sí, 0 = no)"
    tp[f"B{T_HIJOS}"] = "Hijos (cantidad)"
    tp[f"B{T_HINC}"] = "Hijos incapacitados para el trabajo (cantidad)"
    tp[f"B{T_DESP}"] = "Deducción especial que corresponde"
    tp[f"C{T_CONY}"], tp[f"C{T_HIJOS}"] = conyuge, hijos
    tp[f"C{T_HINC}"], tp[f"C{T_DESP}"] = hijos_incapacitados, tipo_deduccion_especial
    tp[f"C{T_COEF1}"], tp[f"C{T_COEF2}"] = coef1, coef2
    tp[f"C{T_MES}"] = f"=MAX(1,COUNTIF('{HOJA_PROY}'!$D${rV}:$O${rV},\"<>0\"))"

    def busca(c, concepto):
        return (
            f'IFERROR(INDEX({c}${T_PRIMERA}:{c}${T_ULTIMA},'
            f'MATCH("{concepto}",$B${T_PRIMERA}:$B${T_ULTIMA},0)),0)'
        )

    for i in range(12):
        c = L(3 + i)
        tp[f"{c}{T_RMNI}"] = "=" + busca(c, CONCEPTO_MNI)
        tp[f"{c}{T_RDE}"] = (
            f'=IF($C${T_DESP}="Ap. 1 - Autonomos",{c}{T_RMNI}*$C${T_COEF1},'
            f'IF($C${T_DESP}="Ap. 1 - Nuevos profesionales",{c}{T_RMNI}*$C${T_COEF2},'
            f'IF($C${T_DESP}="Ap. 2 - Relacion de dependencia",{busca(c, CONCEPTO_DEESP)},0)))'
        )
        tp[f"{c}{T_RCF}"] = (
            f"=$C${T_CONY}*{busca(c, CONCEPTO_CONYUGE)}+$C${T_HIJOS}*{busca(c, CONCEPTO_HIJOS)}"
            f"+$C${T_HINC}*{busca(c, CONCEPTO_HINC)}"
        )
    tp.column_dimensions["B"].width = 52
    for col in "CDEFGHIJKLMN":
        tp.column_dimensions[col].width = 14

    # --- Reporte (desde la Plantilla embebida) ----------------------------
    pl_wb = openpyxl.load_workbook(PLANTILLA_PATH)
    pl = pl_wb["Plantilla"]
    ws = wb.create_sheet("Reporte", 0)
    for row in pl.iter_rows():
        for c in row:
            x = ws.cell(c.row, c.column)
            if not isinstance(c, openpyxl.cell.cell.MergedCell):
                x.value = c.value
            if c.has_style:
                _copiar_estilo(c, x)
    for m in pl.merged_cells.ranges:
        ws.merge_cells(str(m))
    for k, dim in pl.column_dimensions.items():
        ws.column_dimensions[k].width = dim.width
    for k, dim in pl.row_dimensions.items():
        if dim.height:
            ws.row_dimensions[k].height = dim.height
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = ws.page_setup.fitToHeight = 1
    ws.print_area = "A1:P33"

    H, T = f"'{HOJA_PROY}'!", f"'{HOJA_TOPES}'!"
    MES = T + f"$C${T_MES}"
    ws["B5"] = f"=UPPER({H}C1)"
    ws["B6"] = f'="CUIT "&IF(ISNUMBER({H}D1),TEXT({H}D1,"00-00000000-0"),{H}D1)'
    ws["B7"] = (
        '="Información acumulada al mes de "&INDEX({"enero","febrero","marzo","abril",'
        '"mayo","junio","julio","agosto","septiembre","octubre","noviembre","diciembre"},'
        + MES + f')&" de {ANIO}"'
    )
    solo = f"{H}$P${rA}<>SUM({H}$D${rA}:$O${rA})"
    for i in range(12):
        n, c, dd = i + 1, L(3 + i), L(4 + i)
        hide = f'IF({n}>{MES},"",'
        ws[f"{c}21"] = f"={H}{dd}{rV}"
        ws[f"{c}22"] = f"={H}{dd}{rG}"
        ws[f"{c}23"] = f"={c}21-{c}22"
        ws[f"{c}24"] = f"={hide}{T}{c}{T_RMNI})"
        ws[f"{c}25"] = f"={hide}{T}{c}{T_RDE})"
        ws[f"{c}26"] = f"={hide}{T}{c}{T_RCF})"
        ws[f"{c}27"] = f"={hide}SUM($C23:{c}23)-{c}24-{c}25-{c}26)"
        m = f"MATCH({c}27*12/{n},{eU},1)"
        ws[f"{c}28"] = (
            f"={hide}IF(N({c}27)<=0,0,{n}/12*INDEX({eW},{m})+"
            f"({c}27-{n}/12*INDEX({eY},{m}))*INDEX({eX},{m})))"
        )
        ws[f"{c}29"] = f"={hide}SUM({H}$D{rR}:{dd}{rR}))"
        ws[f"{c}30"] = f'=IF(OR({n}>{MES},{solo}),"",SUM({H}$D{rA}:{dd}{rA}))'
        ws[f"{c}31"] = f'=IF(OR({n}>{MES},{solo}),"",{c}28-{c}29-{c}30)'
        v = s[f"{dd}{rV}"].value
        ws.column_dimensions[c].width = 6.6 if v in (None, 0) else 20
    ws.column_dimensions["O"].width = 21
    for r in (21, 22, 23):
        ws[f"O{r}"] = f"=SUM(C{r}:N{r})"
    for r in range(24, 30):
        ws[f"O{r}"] = f"=INDEX($C{r}:$N{r},{MES})"
    ws["O30"] = f"={H}$P${rA}"
    ws["O31"] = "=O28-O29-O30"
    ws["B14"], ws["C14"], ws["F14"], ws["I14"], ws["L14"] = "=O21", "=O22", "=O27", "=O28", "=O31"
    ws["B15"] = (
        '="Margen neto sobre ventas: "&TEXT(IF(O21=0,0,O23/O21),"0.0%")&'
        '"     ·     Tasa efectiva del impuesto: "&TEXT(IF(N(O27)<=0,0,O28/O27),"0.0%")&'
        '"     ·     Desde las deducciones personales, importes acumulados al mes informado"'
    )
    ws["L33"] = '="Mar del Plata · Emitido el "&TEXT(TODAY(),"dd/mm/yyyy")'

    wb.active = 0
    wb.calculation.fullCalcOnLoad = True
    import io
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
