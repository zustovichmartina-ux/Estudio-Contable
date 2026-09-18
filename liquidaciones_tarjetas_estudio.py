# -*- coding: utf-8 -*-
"""Liquidaciones de tarjetas — formato del estudio (Liquidaciones de Tarjetas 1.xlsx).

Hojas: Resumen + una solapa por medio (Master/Visa crédito-débito, CABAL, FAVA, Naranja).
Cada liquidación es una fila: nro, local, tipo, fechas, ventas, gravado 21%, IVA CF,
percepciones, SIRTAC, neto PDF, neto calculado y diferencia.
"""
from __future__ import annotations

import re
import tempfile
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from procesadores.tarjeta_parser import (
    detectar_entidad_por_texto,
    extraer_con_plantilla,
)

COLOR_TITULO = "1F4E78"
COLOR_HEADER = "1F4E78"
COLOR_TOTAL = "D9E1F2"
COLOR_ZEBRA = "F2F2F2"
COLOR_BORDE = "BFBFBF"
TOL = 0.05
_FONT = "Arial"
_NUM = r'\$#,##0.00;[Red]"-$"#,##0.00'
_THIN = Border(
    left=Side(style="thin", color=COLOR_BORDE),
    right=Side(style="thin", color=COLOR_BORDE),
    top=Side(style="thin", color=COLOR_BORDE),
    bottom=Side(style="thin", color=COLOR_BORDE),
)
_FILL_HDR = PatternFill("solid", fgColor=COLOR_HEADER)
_FILL_ZEBRA = PatternFill("solid", fgColor=COLOR_ZEBRA)
_FILL_TOT = PatternFill("solid", fgColor=COLOR_TOTAL)
_FONT_TITLE = Font(name=_FONT, bold=True, size=16, color=COLOR_TITULO)
_FONT_SUB = Font(name=_FONT, size=10, color="666666")
_FONT_HDR = Font(name=_FONT, bold=True, color="FFFFFF", size=10)
_FONT_BODY = Font(name=_FONT, size=10)
_FONT_BOLD = Font(name=_FONT, bold=True, size=10)
_WRAP = Alignment(wrap_text=True, horizontal="center", vertical="center")

MEDIOS_ORDEN = (
    "Master Crédito",
    "Master Débito",
    "Visa Crédito",
    "Visa Débito",
    "CABAL",
    "FAVA",
    "Naranja",
)
HDR_DETALLE = [
    "Nro. Liquidación",
    "Local",
    "Tipo",
    "F. Presentación",
    "F. Pago (Acreditación)",
    "Ventas Brutas",
    "Neto Gravado 21% (Arancel+Dto.Financ.)",
    "IVA Crédito Fiscal 21%",
    "Percepción IIBB",
    "Retención IIBB SIRTAC",
    "Percepción IVA",
    "Neto según Liquidación",
    "Neto Calculado",
    "Diferencia",
]
NOTA_HOJA = (
    'Notas: "Local" indica de qué comercio proviene cada liquidación '
    "(JBJ=Juan B. Justo, GUE=Güemes), sumados en esta solapa. "
    '"Neto Gravado 21%" e "IVA Crédito Fiscal 21%" agrupan Arancel '
    "(+cuotas/costo financiero) y Dto.Financiación con su IVA. "
    'Filas "Ajuste retenc." son correcciones sin venta asociada; '
    '"Nota de crédito" son devoluciones con importes negativos.'
)

RE_PAGO = re.compile(
    r"IMPORTE\s*NETO\s*DE\s*PAGOS\s*\$\s*([\d.]+,\d{2})(-)?"
    r".{0,250}?"
    r"el\s+d[ií]a\s+(\d{2}/\d{2}/\d{4})"
    r".{0,180}?"
    r"Nro\.?\s*Liq:\s*(\d+)"
    r"(?:F\.?\s*Pres\s*(\d{2}/\d{2}/\d{4}))?",
    re.I | re.S,
)
RE_VENTA_CTDO = re.compile(
    r"Venta\s+ctdo\s+(\d{2}/\d{2}/\d{2})\s+\S+\s+\S+\s+\S+\s+\S+\s+"
    r"[\d,.]+\s+([\d.]+,\d{2})\s+([\d.]+,\d{2})\s+([\d.]+,\d{2})",
    re.I,
)
RE_VTA_CUO = re.compile(
    r"vta\s+\d+\s+cuo\s+(\d{2}/\d{2}/\d{2})\s+\S+\s+\S+\s+\S+\s+\S+\s+"
    r"Cup[^\s]*\s+([\d.]+,\d{2})\s+([\d.]+,\d{2})\s+([\d.]+,\d{2})",
    re.I,
)
RE_TOTAL_PRESENTADO = re.compile(r"Total\s+presentado:\s*([\d.]+,\d{2})", re.I)
RE_CUIT = re.compile(r"CUIT:\s*(\d{2}-\d{8}-\d)")
RE_NRO_COM = re.compile(r"N.?º?\s*Comercio:\s*([0-9/ ]+)", re.I)
RE_TITULO_FD = re.compile(
    r"TARJETA\s+DE\s+(CR[EÉ]DITO|D[EÉ]BITO)\s+PESOS\s+(\w+)\s+(\d{4})",
    re.I,
)


def money(s: str | None, neg: str | None = None) -> float:
    if s is None:
        return 0.0
    t = str(s).strip().replace("$", "").replace(" ", "")
    if not t:
        return 0.0
    signo = -1 if (neg or t.endswith("-") or t.startswith("-") or t.startswith("(")) else 1
    t = t.replace("-", "").replace("(", "").replace(")", "")
    if "," in t:
        t = t.replace(".", "").replace(",", ".")
    try:
        return round(signo * float(t), 2)
    except ValueError:
        return 0.0


def texto_pdf(path: Path) -> str:
    with pdfplumber.open(str(path)) as pdf:
        return "\n".join((pg.extract_text() or "") for pg in pdf.pages)


def clasificar_pdf(nombre: str, texto: str) -> str:
    n = (nombre or "").upper()
    t = (texto or "").upper()
    blob = n + "\n" + t[:5000]
    if "FAVACARD" in blob or re.search(r"\bFAVA\b", n):
        return "FAVA"
    if "NARANJA" in blob:
        return "NARANJA"
    if "LIQUIDACION PARA PAGO A COMERCIOS" in blob or (
        "CABAL" in blob and "ARANCEL DE DESCUENTO" in blob
    ):
        return "CABAL"
    if "IMPORTE NETO DE PAGOS" in blob and (
        "VENTAS C/DESCUENTO" in blob
        or "VENTAS C/DTO" in blob
        or "FIRST DATA" in blob
        or "FISERV" in blob
        or "TARJETA DE CREDITO PESOS" in blob
        or "TARJETA DE DEBITO PESOS" in blob
    ):
        return "FIRST_DATA"
    return "OTRO"


def _fecha(s: str) -> str:
    m = re.search(r"(\d{2})/(\d{2})/(\d{2,4})", str(s or ""))
    if not m:
        return str(s or "")
    d, mo, y = m.groups()
    if len(y) == 2:
        y = "20" + y
    return f"{d}/{mo}/{y}"


def _local(nombre: str, texto: str = "") -> str:
    n = (nombre or "").upper()
    if re.search(r"\bGUE\b", n) or "6378" in n:
        return "GUE"
    if re.search(r"\bJBJ\b", n) or "5031" in n or "1614" in n:
        return "JBJ"
    blob = f"{nombre} {texto[:2000]}".upper()
    if re.search(r"\bGUE\b|GUEMES|GÜEMES", blob):
        return "GUE"
    if re.search(r"\bJBJ\b|JUSTO", blob):
        return "JBJ"
    return ""


def _medio_fd(nombre: str, texto: str) -> str:
    n = (nombre or "").upper()
    t = (texto or "")[:2500].upper()
    deb = bool(
        re.search(r"DEBIT|DEB\b", n)
        or re.search(r"TARJETA DE D[EÉ]BITO", t)
    )
    if "VISA" in n or "VISA" in t:
        return "Visa Débito" if deb else "Visa Crédito"
    if "MASTER" in n or "MASTERCARD" in t:
        return "Master Débito" if deb else "Master Crédito"
    m = RE_TITULO_FD.search(texto or "")
    if m:
        deb = "DEB" in m.group(1).upper()
        return "Master Débito" if deb else "Master Crédito"
    return "Master Crédito"


@dataclass
class VentaCupon:
    fecha: str
    ventas: float
    arancel: float
    dto_fin: float


@dataclass
class LiqFD:
    archivo: str
    nro: str
    fecha: str
    ventas_contado: float = 0.0
    ventas_cuotas: float = 0.0
    arancel: float = 0.0
    arancel_cuotas: float = 0.0
    dto_cont: float = 0.0
    dto_cuota: float = 0.0
    iva_arancel: float = 0.0
    iva_arancel_cuotas: float = 0.0
    iva_dto_cont: float = 0.0
    iva_dto_cuota: float = 0.0
    perc_iva_15: float = 0.0
    perc_iva_30: float = 0.0
    sirtac: float = 0.0
    per_iibb_ba: float = 0.0
    perc_caba: float = 0.0
    perc_iva_qr: float = 0.0
    serv_int: float = 0.0
    iva_serv_int: float = 0.0
    otras: float = 0.0
    neto: float = 0.0
    fecha_pres: str = ""
    local: str = ""
    tipo: str = "Liquidación de ventas"
    medio: str = ""
    presentado: float = 0.0
    cupones: list[VentaCupon] = field(default_factory=list)
    otras_det: list[str] = field(default_factory=list)

    @property
    def ventas(self) -> float:
        return round(self.ventas_contado + self.ventas_cuotas, 2)

    @property
    def neto_gravado(self) -> float:
        return round(
            self.arancel + self.arancel_cuotas + self.dto_cont + self.dto_cuota + self.serv_int,
            2,
        )

    @property
    def iva_cf(self) -> float:
        return round(
            self.iva_arancel
            + self.iva_arancel_cuotas
            + self.iva_dto_cont
            + self.iva_dto_cuota
            + self.iva_serv_int,
            2,
        )

    @property
    def perc_iibb(self) -> float:
        return round(self.per_iibb_ba + self.perc_caba, 2)

    @property
    def perc_iva(self) -> float:
        return round(self.perc_iva_15 + self.perc_iva_30 + self.perc_iva_qr, 2)

    @property
    def deducciones(self) -> float:
        return round(
            self.neto_gravado + self.iva_cf + self.perc_iibb + self.sirtac + self.perc_iva + self.otras,
            2,
        )

    @property
    def neto_calc(self) -> float:
        return round(self.ventas - self.deducciones, 2)


@dataclass
class FilaLiq:
    nro: str
    local: str
    tipo: str
    f_pres: str
    f_pago: str
    ventas: float
    neto_gravado: float
    iva_cf: float
    perc_iibb: float
    sirtac: float
    perc_iva: float
    neto: float
    medio: str
    declarado: float = 0.0
    archivo: str = ""


@dataclass
class ResultadoTarjetasEstudio:
    ok: bool
    excel_bytes: bytes = b""
    nombre_archivo: str = "Liquidaciones_de_Tarjetas.xlsx"
    mensajes: list[str] = field(default_factory=list)
    advertencias: list[str] = field(default_factory=list)
    n_pdf: int = 0
    n_fd: int = 0
    n_ventas: int = 0
    n_control_ok: int = 0
    n_control_diff: int = 0
    comercio: str = ""
    periodo: str = ""
    error: str = ""
    hojas: list[str] = field(default_factory=list)


def _aplicar_deduccion(liq: LiqFD, concepto: str, monto: float) -> None:
    u = re.sub(r"\s+", " ", concepto.upper())
    if "IVA ARANCEL CUOTAS" in u or ("IVA" in u and "ARANCEL CUOTAS" in u):
        liq.iva_arancel_cuotas = round(liq.iva_arancel_cuotas + monto, 2)
    elif "ARANCEL CUOTAS" in u:
        liq.arancel_cuotas = round(liq.arancel_cuotas + monto, 2)
    elif re.search(r"\bARANCEL\b", u) and "IVA" not in u:
        liq.arancel = round(liq.arancel + monto, 2)
    elif "IVA CRED" in u or "S/ARANC" in u:
        liq.iva_arancel = round(liq.iva_arancel + monto, 2)
    elif "IVA S/DTO FIN ADQ CUOTA" in u:
        liq.iva_dto_cuota = round(liq.iva_dto_cuota + monto, 2)
    elif "IVA S/DTO FIN ADQ CONT" in u:
        liq.iva_dto_cont = round(liq.iva_dto_cont + monto, 2)
    elif "DTO S/VENTAS FIN ADQ CUOTA" in u:
        liq.dto_cuota = round(liq.dto_cuota + monto, 2)
    elif "DTO S/VENTAS FIN ADQ CONT" in u:
        liq.dto_cont = round(liq.dto_cont + monto, 2)
    elif "PERCEPCION IVA" in u and "1,50" in u:
        liq.perc_iva_15 = round(liq.perc_iva_15 + monto, 2)
    elif "PERCEPCION IVA" in u and "3,00" in u:
        liq.perc_iva_30 = round(liq.perc_iva_30 + monto, 2)
    elif "PERC IVA 2408" in u or "PERCEPCION IVA 2408" in u:
        liq.perc_iva_qr = round(liq.perc_iva_qr + monto, 2)
    elif "SIRTAC" in u:
        liq.sirtac = round(liq.sirtac + monto, 2)
    elif "PER B.A" in u or "PER BA" in u.replace(".", ""):
        liq.per_iibb_ba = round(liq.per_iibb_ba + monto, 2)
    elif "PERC.IB.CABA" in u or "IB.CABA" in u:
        liq.perc_caba = round(liq.perc_caba + monto, 2)
    elif "SERV.OPER" in u or "SERVICIO OPER" in u:
        if "IVA" in u:
            liq.iva_serv_int = round(liq.iva_serv_int + monto, 2)
        else:
            liq.serv_int = round(liq.serv_int + monto, 2)
    else:
        liq.otras = round(liq.otras + monto, 2)
        liq.otras_det.append(f"{concepto} {monto}")


def parsear_fd(path: Path, texto: str | None = None) -> list[LiqFD]:
    text = texto if texto is not None else texto_pdf(path)
    presentado = 0.0
    m_tp = RE_TOTAL_PRESENTADO.search(text)
    if m_tp:
        presentado = money(m_tp.group(1))
    medio = _medio_fd(path.name, text)
    local = _local(path.name, text)
    out: list[LiqFD] = []
    matches = list(RE_PAGO.finditer(text))
    for i, m in enumerate(matches):
        start = matches[i - 1].end() if i else 0
        bloque = text[start : m.start()]
        liq = LiqFD(
            archivo=path.name,
            nro=m.group(4),
            fecha=_fecha(m.group(3)),
            neto=money(m.group(1), m.group(2)),
            fecha_pres=_fecha(m.group(5) or ""),
            local=local,
            medio=medio,
            presentado=presentado,
        )
        for mv in re.finditer(
            r"\+\s*VENTAS\s+C/?DESCUENTO\s+CONTADO\s*\$\s*([\d.]+,\d{2})(-)?",
            bloque,
            re.I,
        ):
            liq.ventas_contado = round(liq.ventas_contado + money(mv.group(1), mv.group(2)), 2)
        for mv in re.finditer(
            r"\+\s*VENTAS\s+C/DTO\s+ANTIC POR FINANC ADQUIREN\s*\$\s*([\d.]+,\d{2})(-)?",
            bloque,
            re.I,
        ):
            liq.ventas_cuotas = round(liq.ventas_cuotas + money(mv.group(1), mv.group(2)), 2)
        for mv in re.finditer(
            r"\+\s*VENTAS\s+C/DTO\s+PLAN CUOTAS\s*\$\s*([\d.]+,\d{2})(-)?",
            bloque,
            re.I,
        ):
            liq.ventas_cuotas = round(liq.ventas_cuotas + money(mv.group(1), mv.group(2)), 2)
        qrpct = False
        for md in re.finditer(r"-\s*((?:(?!\$).){3,90}?)\$\s*([\d.]+,\d{2})(-)?", bloque):
            conc = re.sub(r"\s+", " ", md.group(1)).strip()
            if re.search(r"0800|www\.|Rep[uú]blica|Estimado|Centro de Atenci", conc, re.I):
                continue
            if "QRPCT" in conc.upper():
                qrpct = True
            _aplicar_deduccion(liq, conc, money(md.group(2), md.group(3)))
        for mc in RE_VENTA_CTDO.finditer(bloque):
            liq.cupones.append(
                VentaCupon(mc.group(1), money(mc.group(2)), money(mc.group(3)), money(mc.group(4)))
            )
        for mc in RE_VTA_CUO.finditer(bloque):
            liq.cupones.append(
                VentaCupon(mc.group(1), money(mc.group(2)), money(mc.group(3)), money(mc.group(4)))
            )
        if liq.ventas < 0:
            liq.tipo = "Nota de crédito"
        elif qrpct or abs(liq.ventas) < 0.01:
            liq.tipo = "Ajuste retenc. (QRPCT)"
        else:
            liq.tipo = "Liquidación de ventas"
        out.append(liq)
    return out


def parsear_naranja(path: Path, texto: str | None = None) -> dict[str, float | str]:
    t = texto if texto is not None else texto_pdf(path)

    def grab(pat: str) -> float:
        m = re.search(pat, t, re.I | re.S)
        return money(m.group(1)) if m else 0.0

    nro = ""
    m_nro = re.search(r"Tipo y N[º°o.]+\s*:\s*([A-Z0-9\-]+)", t, re.I)
    if m_nro:
        nro = m_nro.group(1).strip()
    m_pago = re.search(r"Fecha de Pago\s+(\d{2}/\d{2}/\d{2,4})", t, re.I)
    f_pago = _fecha(m_pago.group(1) if m_pago else "")
    m_comp = re.search(r"(\d{2}/\d{2}/\d{2,4})\s+\d{5,}", t)
    f_pres = _fecha(m_comp.group(1) if m_comp else "")
    local = _local(path.name, t)
    if local == "GUE":
        etiqueta = "GUE"
    elif local == "JBJ":
        etiqueta = "JBJ"
    else:
        etiqueta = path.stem
    arancel = grab(r"Arancel\s+-\s*\$\s*([\d.]+,\d{2})")
    interes = grab(r"Inter[eé]s\s+Plan[^\n]*\$\s*([\d.]+,\d{2})")
    return {
        "archivo": path.name,
        "nro": nro,
        "local": etiqueta if etiqueta in ("GUE", "JBJ") else local,
        "f_pres": f_pres,
        "f_pago": f_pago,
        "total": grab(r"Totales\s+1\s+\$\s*([\d.]+,\d{2})"),
        "arancel": arancel,
        "interes": interes,
        "iva": grab(r"IVA\s+21\.0\s*%\s+\$\s*([\d.]+,\d{2})"),
        "perc_iva": grab(r"Percepci[oó]n\s+de\s+IVA[^\n]*\$\s*([\d.]+,\d{2})"),
        "perc_iibb": grab(r"Percepci[oó]n\s+Ingresos?\s+Brutos[^\n]*\$\s*([\d.]+,\d{2})"),
        "sirtac": grab(r"Sirtac[^\n]*\$\s*([\d.]+,\d{2})"),
        "neto": grab(r"Neto Liquidado.*?Importe\s+\$\s*([\d.]+,\d{2})"),
        "bi": round(arancel + interes, 2),
    }


def parsear_fava(path: Path, texto: str | None = None) -> dict[str, float | str]:
    t = texto if texto is not None else texto_pdf(path)

    def grab(pat: str) -> float:
        m = re.search(pat, t, re.I)
        return money(m.group(1)) if m else 0.0

    nro = ""
    m_nro = re.search(r"N[°º]\s*A\s*(\d{4}-\d+)", t, re.I)
    if m_nro:
        nro = "A" + m_nro.group(1)
    m_pago = re.search(r"COMPROBANTELIQUIDACION PAGO.*?Fecha:\s*(\d{2}/\d{2}/\d{2,4})", t, re.I | re.S)
    f_pago = _fecha(m_pago.group(1) if m_pago else "")
    m_pres = re.search(
        r"DETALLE ULTIMA PRESENTACION.*?(\d{2}/\d{2}/\d{2,4})",
        t,
        re.I | re.S,
    )
    f_pres = _fecha(m_pres.group(1) if m_pres else "")
    local = _local(path.name, t)
    com = grab(r"Cargos y Bonific\.\s*Comisi[oó]n\s*:\s*-?\s*([\d.]+,\d{2})")
    dto = grab(r"Descuento Plan C\s*:\s*-?\s*([\d.]+,\d{2})")
    bon = grab(r"Bonificacion Plan C\s*:\s*([\d.]+,\d{2})")
    return {
        "archivo": path.name,
        "nro": nro,
        "local": local,
        "f_pres": f_pres,
        "f_pago": f_pago,
        "total": grab(r"Total de la presentaci[oó]n\s*:\s*([\d.]+,\d{2})"),
        "comision": com,
        "dto": dto,
        "bonif": bon,
        "iva": grab(r"Impuestos Iva 21,00[^\n]*:\s*-?\s*([\d.]+,\d{2})"),
        "perc_iva": grab(r"Percepcion de IVA\s*:\s*-?\s*([\d.]+,\d{2})"),
        "perc_iibb": grab(r"Percepcion I\.Brutos PBA[^\n]*:\s*-?\s*([\d.]+,\d{2})"),
        "sirtac": grab(r"Ret\.\s*IIBB SIRTAC[^\n]*:\s*-?\s*([\d.]+,\d{2})"),
        "neto": grab(r"NETO A PAGAR PESOS\s*\$\s*([\d.]+,\d{2})"),
        "bi": round(com + dto - bon, 2),
    }


def parsear_cabal(path: Path, texto: str | None = None) -> list[dict[str, float | str]]:
    t = texto if texto is not None else texto_pdf(path)
    local = _local(path.name, t) or "JBJ"
    bloques = re.split(r"FECHA DE PAGO\s*:", t)
    out: list[dict[str, float | str]] = []
    for b in bloques[1:]:
        m_nro = re.search(r"LIQUIDACION NRO\.\s*:\s*(\d+)", b, re.I)
        if not m_nro:
            continue

        def grab(pat: str) -> float:
            m = re.search(pat, b, re.I)
            return money(m.group(1)) if m else 0.0

        m_pago = re.search(r"(\d{2}/\d{2}/\d{4})", b)
        f_pago = _fecha(m_pago.group(1) if m_pago else "")
        m_ar = re.search(r"ARANCEL DE DESCUENTO\s+[\d.,]+%\s+([\d.]+,\d{2})(-)?", b, re.I)
        arancel = money(m_ar.group(1), m_ar.group(2)) if m_ar else 0.0
        cft_m = re.search(r"COSTO FINANCIERO TOTAL\s+([\d.]+,\d{2})(-)?", b, re.I)
        cft = money(cft_m.group(1), cft_m.group(2)) if cft_m else 0.0
        sirtac = 0.0
        m_sirtac = re.search(r"RETENCION IIBB SIRTAC[^\n]+", b, re.I)
        if m_sirtac:
            nums_s = re.findall(r"([\d.]+,\d{2})(-)?", m_sirtac.group(0))
            if nums_s:
                sirtac = money(nums_s[-1][0], nums_s[-1][1])
        iva = 0.0
        m_iva = re.search(r"IVA S/ARANCEL \+ COSTO FINANCIERO[^\n]+", b, re.I)
        if not m_iva:
            m_iva = re.search(r"IVA S/ARANCEL DE DESCUENTO[^\n]+", b, re.I)
        if m_iva:
            nums_i = re.findall(r"([\d.]+,\d{2})(-)?", m_iva.group(0))
            if nums_i:
                iva = money(nums_i[-1][0], nums_i[-1][1])
        m_pres = re.search(
            r"(\d{2}/\d{2}/\d{4})\s+\S+\s+\S+\s+\d+\s+\*TOTAL RESUMEN\*",
            b,
            re.I,
        )
        if m_pres:
            f_pres = _fecha(m_pres.group(1))
        else:
            fechas_venta = re.findall(r"(\d{2}/\d{2}/\d{4})\s+\d+", b)
            f_pres = _fecha(fechas_venta[-1] if fechas_venta else f_pago)
        out.append(
            {
                "archivo": path.name,
                "nro": m_nro.group(1).lstrip("0") or m_nro.group(1),
                "local": local,
                "f_pres": f_pres,
                "f_pago": f_pago,
                "total": grab(r"TOTAL DE VENTAS\.[^\n]*?([\d.]+,\d{2})"),
                "arancel": arancel,
                "cft": cft,
                "iva": iva,
                "sirtac": sirtac,
                "neto": grab(r"IMPORTE NETO FINAL A LIQUIDAR\s*\.*\s*([\d.]+,\d{2})"),
                "bi": round(arancel + cft, 2),
            }
        )
    return out


def _orden_fila(x: FilaLiq) -> tuple:
    return ({"JBJ": 0, "GUE": 1}.get(x.local, 9), x.f_pago, x.nro)


def _fmt_money(cell) -> None:
    cell.number_format = _NUM
    cell.font = _FONT_BODY
    cell.border = _THIN


def _filas_desde_datos(
    naranja: list[dict],
    fava: list[dict],
    cabal: list[dict],
    liqs: list[LiqFD],
) -> list[FilaLiq]:
    filas: list[FilaLiq] = []
    for x in liqs:
        filas.append(
            FilaLiq(
                nro=x.nro,
                local=x.local,
                tipo=x.tipo,
                f_pres=x.fecha_pres,
                f_pago=x.fecha,
                ventas=x.ventas,
                neto_gravado=x.neto_gravado,
                iva_cf=x.iva_cf,
                perc_iibb=x.perc_iibb,
                sirtac=x.sirtac,
                perc_iva=x.perc_iva,
                neto=x.neto,
                medio=x.medio,
                declarado=x.presentado,
                archivo=x.archivo,
            )
        )
    for item in cabal:
        filas.append(
            FilaLiq(
                nro=str(item.get("nro") or ""),
                local=str(item.get("local") or ""),
                tipo="Liquidación de ventas",
                f_pres=str(item.get("f_pres") or ""),
                f_pago=str(item.get("f_pago") or ""),
                ventas=float(item.get("total") or 0),
                neto_gravado=float(item.get("bi") or 0),
                iva_cf=float(item.get("iva") or 0),
                perc_iibb=0.0,
                sirtac=float(item.get("sirtac") or 0),
                perc_iva=0.0,
                neto=float(item.get("neto") or 0),
                medio="CABAL",
                declarado=float(item.get("total") or 0),
                archivo=str(item.get("archivo") or ""),
            )
        )
    for item in fava:
        filas.append(
            FilaLiq(
                nro=str(item.get("nro") or ""),
                local=str(item.get("local") or ""),
                tipo="Liquidación de ventas",
                f_pres=str(item.get("f_pres") or ""),
                f_pago=str(item.get("f_pago") or ""),
                ventas=float(item.get("total") or 0),
                neto_gravado=float(item.get("bi") or 0),
                iva_cf=float(item.get("iva") or 0),
                perc_iibb=float(item.get("perc_iibb") or 0),
                sirtac=float(item.get("sirtac") or 0),
                perc_iva=float(item.get("perc_iva") or 0),
                neto=float(item.get("neto") or 0),
                medio="FAVA",
                declarado=float(item.get("total") or 0),
                archivo=str(item.get("archivo") or ""),
            )
        )
    for item in naranja:
        filas.append(
            FilaLiq(
                nro=str(item.get("nro") or ""),
                local=str(item.get("local") or ""),
                tipo="Liquidación de ventas",
                f_pres=str(item.get("f_pres") or ""),
                f_pago=str(item.get("f_pago") or ""),
                ventas=float(item.get("total") or 0),
                neto_gravado=float(item.get("bi") or 0),
                iva_cf=float(item.get("iva") or 0),
                perc_iibb=float(item.get("perc_iibb") or 0),
                sirtac=float(item.get("sirtac") or 0),
                perc_iva=float(item.get("perc_iva") or 0),
                neto=float(item.get("neto") or 0),
                medio="Naranja",
                declarado=float(item.get("total") or 0),
                archivo=str(item.get("archivo") or ""),
            )
        )
    return filas


def _encabezado_hoja(ws: Worksheet, titulo: str, comercio: str, cuit: str, periodo: str) -> None:
    ws["A1"] = titulo
    ws["A1"].font = _FONT_TITLE
    ws.merge_cells("A1:N1")
    cuit_txt = f"CUIT {cuit}" if cuit else ""
    ws["A2"] = f"Comercio: {comercio or '—'} — {cuit_txt}".strip(" —")
    ws["A2"].font = _FONT_SUB
    ws.merge_cells("A2:N2")
    ws["A3"] = (
        "Entidad pagadora: Banco de la Nación Argentina (First Data/Fiserv), "
        "salvo CABAL y FAVA/Naranja (procesadoras propias)."
    )
    ws["A3"].font = _FONT_SUB
    ws.merge_cells("A3:N3")
    ws.row_dimensions[1].height = 22


def _escribir_hoja_medio(
    ws: Worksheet,
    medio: str,
    filas: list[FilaLiq],
    comercio: str,
    cuit: str,
    periodo: str,
) -> int:
    _encabezado_hoja(
        ws,
        f"Liquidación de Tarjeta — {medio} — {periodo}".strip(" —"),
        comercio,
        cuit,
        periodo,
    )
    header_row = 5
    for i, h in enumerate(HDR_DETALLE, 1):
        cell = ws.cell(header_row, i, h)
        cell.font = _FONT_HDR
        cell.fill = _FILL_HDR
        cell.alignment = _WRAP
        cell.border = _THIN
    ws.row_dimensions[header_row].height = 36
    start = header_row + 1
    for i, fila in enumerate(filas):
        r = start + i
        ws.cell(r, 1, fila.nro)
        ws.cell(r, 2, fila.local)
        ws.cell(r, 3, fila.tipo)
        ws.cell(r, 4, fila.f_pres)
        ws.cell(r, 5, fila.f_pago)
        ws.cell(r, 6, fila.ventas)
        ws.cell(r, 7, fila.neto_gravado)
        ws.cell(r, 8, fila.iva_cf)
        ws.cell(r, 9, fila.perc_iibb)
        ws.cell(r, 10, fila.sirtac)
        ws.cell(r, 11, fila.perc_iva)
        ws.cell(r, 12, fila.neto)
        ws.cell(r, 13, f"=F{r}-G{r}-H{r}-I{r}-J{r}-K{r}")
        ws.cell(r, 14, f"=L{r}-M{r}")
        for c in range(1, 15):
            cell = ws.cell(r, c)
            cell.font = _FONT_BODY
            cell.border = _THIN
            if c >= 6:
                cell.number_format = _NUM
            if i % 2:
                cell.fill = _FILL_ZEBRA
    last = start + len(filas) - 1
    tot = last + 1
    ws.cell(tot, 1, "TOTAL")
    ws.cell(tot, 2, f"{len(filas)} liq.")
    for col, letter in enumerate("FGHIJKLMN", 6):
        ws.cell(tot, col, f"=SUM({letter}{start}:{letter}{last})")
    for c in range(1, 15):
        cell = ws.cell(tot, c)
        cell.font = _FONT_BOLD
        cell.fill = _FILL_TOT
        cell.border = _THIN
        if c >= 6:
            cell.number_format = _NUM
    nota_row = tot + 2
    ws.merge_cells(start_row=nota_row, start_column=1, end_row=nota_row, end_column=14)
    ws.cell(nota_row, 1, NOTA_HOJA).font = _FONT_SUB
    ws.cell(nota_row, 1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[nota_row].height = 48
    ws.auto_filter.ref = f"A{header_row}:N{last}"
    ws.freeze_panes = f"A{start}"
    widths = [16, 8, 22, 13, 16, 13, 20, 15, 13, 15, 12, 16, 14, 11]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    return tot


def _escribir_resumen(
    ws: Worksheet,
    refs: list[tuple[str, int, float, int]],
    comercio: str,
    cuit: str,
    periodo: str,
    nros_comercio: str,
    sufijo: str = " (locales sumados)",
) -> None:
    ws["A1"] = f"Liquidaciones de Tarjeta — {periodo}{sufijo}".strip()
    ws["A1"].font = _FONT_TITLE
    ws.merge_cells("A1:K1")
    ws["A2"] = f"Comercio: {comercio or '—'}   |   CUIT: {cuit or '—'}"
    ws["A2"].font = _FONT_SUB
    ws.merge_cells("A2:K2")
    ws["A3"] = nros_comercio or "First Data/Fiserv, CABAL, FAVA y Naranja."
    ws["A3"].font = _FONT_SUB
    ws.merge_cells("A3:K3")
    ws.merge_cells("A4:K4")
    ws["A4"] = (
        "Notas: cada solapa suma las liquidaciones de los locales (columna Local). "
        'Percepción IIBB (col. F) incluye "PER B.A.I.BR.DN.01/04" y análogas. '
        "Retención IIBB SIRTAC (col. G) y Percepción IVA (col. H) van por separado. "
        '"Diferencia" (col. K) compara Ventas Brutas contra el total declarado; debe dar 0.'
    )
    ws["A4"].font = _FONT_SUB
    ws["A4"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[4].height = 48
    hdr = [
        "Medio de pago",
        "Cant. Liquidaciones",
        "Ventas Brutas",
        "Neto Gravado 21% (Arancel+Dto.Financ.)",
        "IVA Crédito Fiscal 21%",
        "Percepción IIBB",
        "Retención IIBB SIRTAC",
        "Percepción IVA",
        "Importe Neto Acreditado",
        "Total declarado (locales)",
        "Diferencia",
    ]
    for i, h in enumerate(hdr, 1):
        cell = ws.cell(6, i, h)
        cell.font = _FONT_HDR
        cell.fill = _FILL_HDR
        cell.alignment = _WRAP
        cell.border = _THIN
    ws.row_dimensions[6].height = 52
    start = 7
    for i, (medio, tot_row, declarado, n_liq) in enumerate(refs):
        r = start + i
        ws.cell(r, 1, medio)
        quoted = f"'{medio}'"
        ws.cell(r, 2, n_liq)
        ws.cell(r, 3, f"={quoted}!F{tot_row}")
        ws.cell(r, 4, f"={quoted}!G{tot_row}")
        ws.cell(r, 5, f"={quoted}!H{tot_row}")
        ws.cell(r, 6, f"={quoted}!I{tot_row}")
        ws.cell(r, 7, f"={quoted}!J{tot_row}")
        ws.cell(r, 8, f"={quoted}!K{tot_row}")
        ws.cell(r, 9, f"={quoted}!L{tot_row}")
        ws.cell(r, 10, declarado)
        ws.cell(r, 11, f"=C{r}-J{r}")
        for c in range(1, 12):
            cell = ws.cell(r, c)
            cell.font = _FONT_BODY
            cell.border = _THIN
            if c >= 2:
                cell.number_format = _NUM if c >= 3 else "0"
    last = start + len(refs) - 1
    tot = last + 1
    ws.cell(tot, 1, "TOTAL GENERAL")
    for col, letter in enumerate("BCDEFGHIJK", 2):
        ws.cell(tot, col, f"=SUM({letter}{start}:{letter}{last})")
    for c in range(1, 12):
        cell = ws.cell(tot, c)
        cell.font = _FONT_BOLD
        cell.fill = _FILL_TOT
        cell.border = _THIN
        if c >= 3:
            cell.number_format = _NUM
    ws.freeze_panes = "A7"
    widths = [18, 14, 17, 16, 14, 13, 15, 13, 16, 16, 12]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _escribir_otros(ws: Worksheet, otros: list[dict]) -> None:
    cols = [
        "Archivo",
        "Entidad",
        "Fecha",
        "Nro_Liquidacion",
        "Neto_Gravado",
        "IVA_21",
        "Percepcion_IVA",
        "Retencion_IVA",
        "Retencion_IIBB",
        "Percepcion_IIBB",
        "Total_Descontado",
    ]
    for c, h in enumerate(cols, 1):
        cell = ws.cell(1, c, h)
        cell.font = _FONT_HDR
        cell.fill = _FILL_HDR
        cell.border = _THIN
    for i, d in enumerate(otros, 2):
        for c, k in enumerate(cols, 1):
            ws.cell(i, c, d.get(k))
            if c >= 5:
                _fmt_money(ws.cell(i, c))
    ws.auto_filter.ref = f"A1:K{len(otros) + 1}"
    ws.freeze_panes = "A2"


def generar_excel_estudio(
    *,
    naranja: list[dict],
    fava: list[dict],
    cabal: list[dict],
    liqs: list[LiqFD],
    otros: list[dict],
    comercio: str,
    periodo: str,
    cuit: str = "",
    nros_comercio: str = "",
) -> bytes:
    filas = _filas_desde_datos(naranja, fava, cabal, liqs)
    por_medio: dict[str, list[FilaLiq]] = {m: [] for m in MEDIOS_ORDEN}
    for fila in filas:
        por_medio.setdefault(fila.medio, []).append(fila)

    wb = Workbook()
    ws_res = wb.active
    ws_res.title = "Resumen"
    refs: list[tuple[str, int, float, int]] = []
    for medio in MEDIOS_ORDEN:
        items = por_medio.get(medio) or []
        if not items:
            continue
        items = sorted(items, key=_orden_fila)
        ws = wb.create_sheet(medio)
        tot_row = _escribir_hoja_medio(ws, medio, items, comercio, cuit, periodo)
        vistos: dict[str, float] = {}
        for it in items:
            if it.archivo not in vistos:
                vistos[it.archivo] = it.declarado
        declarado = round(sum(vistos.values()), 2)
        if medio in ("CABAL", "FAVA", "Naranja"):
            declarado = round(sum(it.ventas for it in items), 2)
        refs.append((medio, tot_row, declarado, len(items)))
    extra = [m for m in por_medio if m not in MEDIOS_ORDEN and por_medio[m]]
    for medio in extra:
        items = por_medio[medio]
        items = sorted(items, key=_orden_fila)
        ws = wb.create_sheet(medio[:31])
        tot_row = _escribir_hoja_medio(ws, medio, items, comercio, cuit, periodo)
        declarado = round(sum(it.ventas for it in items), 2)
        refs.append((ws.title, tot_row, declarado, len(items)))
    locales = {f.local for f in filas if f.local}
    if locales >= {"GUE", "JBJ"}:
        sufijo = " (Güemes + Juan B. Justo sumados)"
    elif locales:
        sufijo = " (locales sumados)"
    else:
        sufijo = ""
    _escribir_resumen(ws_res, refs, comercio, cuit, periodo, nros_comercio, sufijo)
    if otros:
        _escribir_otros(wb.create_sheet("Otros"), otros)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _meta_pdf(texto: str, nombre: str) -> tuple[str, str, str, str]:
    comercio = ""
    if "GLOBAL RECIFE" in (texto or "").upper():
        comercio = "GLOBAL RECIFE SA"
    m_rs = re.search(r"RAZ[OÓ]N SOCIAL[:\s]+([^\n]+)", texto or "", re.I)
    if m_rs and not comercio:
        comercio = m_rs.group(1).strip()[:80]
    cuit = ""
    for m in RE_CUIT.finditer(texto or ""):
        if not m.group(1).startswith("30-52221156") and not m.group(1).startswith("30-50001091"):
            cuit = m.group(1)
            break
    periodo = ""
    m2 = RE_TITULO_FD.search(texto or "")
    if m2:
        periodo = f"{m2.group(2).title()} {m2.group(3)}"
    m3 = re.search(r"(20\d{2})(0[1-9]|1[0-2])", nombre)
    if m3 and not periodo:
        meses = "Enero Febrero Marzo Abril Mayo Junio Julio Agosto Septiembre Octubre Noviembre Diciembre".split()
        periodo = f"{meses[int(m3.group(2)) - 1]} {m3.group(1)}"
    nro = ""
    m_n = RE_NRO_COM.search(texto or "")
    if m_n:
        nro = re.sub(r"\s+", "", m_n.group(1))
    return comercio, cuit, periodo, nro


def procesar_pdfs_tarjetas_estudio(archivos: list[tuple[str, bytes]]) -> ResultadoTarjetasEstudio:
    if not archivos:
        return ResultadoTarjetasEstudio(ok=False, error="Subí al menos un PDF.")

    naranja: list[dict] = []
    fava: list[dict] = []
    cabal: list[dict] = []
    liqs: list[LiqFD] = []
    otros: list[dict] = []
    mensajes: list[str] = []
    advertencias: list[str] = []
    comercio = ""
    periodo = ""
    cuit = ""
    nros: list[str] = []

    with tempfile.TemporaryDirectory(prefix="liq_tarjetas_") as tmp:
        tmp_path = Path(tmp)
        for i, (nombre, data) in enumerate(archivos):
            safe = Path(nombre or f"liq_{i + 1}.pdf").name
            dest = tmp_path / safe
            dest.write_bytes(data)
            try:
                texto = texto_pdf(dest)
            except Exception as exc:
                advertencias.append(f"{safe}: no se pudo leer ({exc})")
                continue
            kind = clasificar_pdf(safe, texto)
            com, cuit_p, per, nro_c = _meta_pdf(texto, safe)
            comercio = comercio or com
            periodo = periodo or per
            cuit = cuit or cuit_p
            if nro_c and nro_c not in nros:
                nros.append(nro_c)
            try:
                if kind == "NARANJA":
                    naranja.append(parsear_naranja(dest, texto))
                    mensajes.append(f"{safe}: Naranja")
                elif kind == "FAVA":
                    fava.append(parsear_fava(dest, texto))
                    mensajes.append(f"{safe}: Favacard")
                elif kind == "CABAL":
                    parsed = parsear_cabal(dest, texto)
                    cabal.extend(parsed)
                    mensajes.append(f"{safe}: Cabal ({len(parsed)} liquidación/es)")
                elif kind == "FIRST_DATA":
                    parsed = parsear_fd(dest, texto)
                    liqs.extend(parsed)
                    mensajes.append(f"{safe}: First Data ({len(parsed)} liquidación/es)")
                else:
                    entidad = detectar_entidad_por_texto(f"{texto}\n{safe}")
                    fila = extraer_con_plantilla(texto, entidad)
                    fila["Archivo"] = safe
                    otros.append(fila)
                    mensajes.append(f"{safe}: {entidad} (hoja Otros)")
            except Exception as exc:
                advertencias.append(f"{safe}: {exc}")

    if not (naranja or fava or cabal or liqs or otros):
        return ResultadoTarjetasEstudio(
            ok=False,
            error="No se pudo extraer ninguna liquidación de los PDF.",
            mensajes=mensajes,
            advertencias=advertencias,
        )

    n_ok = sum(1 for x in liqs if abs(x.neto_calc - x.neto) <= TOL)
    n_diff = len(liqs) - n_ok
    n_ventas = sum(len(x.cupones) for x in liqs)
    for x in liqs:
        if x.otras_det:
            advertencias.append(f"Liq {x.nro}: deducción no mapeada: {x.otras_det}")

    nros_txt = ""
    if nros:
        nros_txt = "N° Comercio Fiserv/Nación: " + " y ".join(nros) + "."
    xlsx = generar_excel_estudio(
        naranja=naranja,
        fava=fava,
        cabal=cabal,
        liqs=liqs,
        otros=otros,
        comercio=comercio,
        periodo=periodo,
        cuit=cuit,
        nros_comercio=nros_txt,
    )
    medios_presentes = {f.medio for f in _filas_desde_datos(naranja, fava, cabal, liqs)}
    hojas = ["Resumen"] + [m for m in MEDIOS_ORDEN if m in medios_presentes]
    hojas += [m for m in sorted(medios_presentes) if m not in MEDIOS_ORDEN]
    if otros:
        hojas.append("Otros")
    slug = re.sub(r"[^A-Za-z0-9]+", "_", (comercio or "Tarjetas"))[:40].strip("_")
    per_slug = re.sub(r"\s+", "_", periodo or "")
    nombre = f"Liquidaciones_de_Tarjetas_{slug}_{per_slug}.xlsx".replace("__", "_")
    return ResultadoTarjetasEstudio(
        ok=True,
        excel_bytes=xlsx,
        nombre_archivo=nombre,
        mensajes=mensajes,
        advertencias=advertencias,
        n_pdf=len(archivos),
        n_fd=len(liqs),
        n_ventas=n_ventas,
        n_control_ok=n_ok,
        n_control_diff=n_diff,
        comercio=comercio,
        periodo=periodo,
        hojas=hojas,
    )
