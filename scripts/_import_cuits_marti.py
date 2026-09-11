# -*- coding: utf-8 -*-
"""Importa CUIT únicos desde Marti.xlsx al registry. NUNCA lee ni guarda claves."""
from __future__ import annotations

import os
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from afip_worker.jobs import normalizar_cuit
from afip_worker.registry import CuitEntry, list_cuits, merge_entries, _key

_RE_11 = re.compile(r"(?<!\d)(\d{11})(?!\d)")
_RE_DASH = re.compile(r"(?<!\d)(\d{2}-\d{8}-\d)(?!\d)")
_CLAVE_HEADERS = re.compile(r"clave|password|pass|pwd", re.I)


def _cuit_checksum(digits: str) -> bool:
    if len(digits) != 11 or not digits.isdigit():
        return False
    mult = (5, 4, 3, 2, 7, 6, 5, 4, 3, 2)
    total = sum(int(digits[i]) * mult[i] for i in range(10))
    dv = 11 - (total % 11)
    if dv == 11:
        dv = 0
    elif dv == 10:
        dv = 9
    return int(digits[10]) == dv


def _digits(raw: object) -> str:
    return re.sub(r"\D", "", str(raw or ""))


def _looks_name(raw: object) -> bool:
    s = str(raw or "").strip()
    if len(s) < 3 or s.isdigit():
        return False
    letters = sum(ch.isalpha() for ch in s)
    return letters >= 3


def extraer_de_claves(ws) -> dict[str, str]:
    """Hoja CLAVES: nombre+CUIT. Ignora columnas de clave."""
    found: dict[str, str] = {}
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None)
    # a veces fila 1 es fecha; fila 2 encabezados
    if header and not any(_CLAVE_HEADERS.search(str(c or "")) for c in header):
        header = next(rows, None)
    for row in rows:
        if not row:
            continue
        pairs = ((0, 1), (7, 8))  # nombre/cuit y bloque derecho
        for ni, ci in pairs:
            if ci >= len(row):
                continue
            digs = _digits(row[ci])
            if not _cuit_checksum(digs):
                continue
            name = str(row[ni] or "").strip() if ni < len(row) else ""
            if not _looks_name(name):
                name = ""
            found[digs] = name or found.get(digs, "")
        if len(row) > 2:
            soc = _digits(row[2])
            if _cuit_checksum(soc) and soc not in found:
                found[soc] = found.get(_digits(row[1]), "") or found.get(soc, "")
    return found


def extraer_todos_los_cuits(wb) -> dict[str, str]:
    found: dict[str, str] = {}
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        if sheet.strip().upper() == "CLAVES":
            found.update(extraer_de_claves(ws))
            continue
        clave_cols: set[int] = set()
        prev_name = ""
        for row_i, row in enumerate(ws.iter_rows(values_only=True)):
            if not row:
                continue
            if row_i < 3:
                for i, c in enumerate(row):
                    if _CLAVE_HEADERS.search(str(c or "")):
                        clave_cols.add(i)
            for i, cell in enumerate(row):
                if i in clave_cols or cell in (None, ""):
                    continue
                s = str(cell)
                if _looks_name(cell):
                    prev_name = str(cell).strip()
                for m in _RE_DASH.findall(s) + _RE_11.findall(s):
                    digs = _digits(m)
                    if not _cuit_checksum(digs):
                        continue
                    if digs not in found or not found[digs]:
                        found[digs] = prev_name
    return found


def main() -> int:
    src = Path(r"\\TANGOSRV\Compartido\CLIENTES\1-Guadi y Marti\Marti.xlsx")
    local = Path(os.environ.get("TEMP") or ".") / "Marti_cuits.xlsx"
    path = src if src.exists() else local
    if not path.exists():
        print("NO ENCONTRÉ Marti.xlsx")
        return 1
    print("leyendo", path)
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        found = extraer_todos_los_cuits(wb)
    finally:
        wb.close()

    before = {e.cuit for e in list_cuits()}
    updates: dict[str, CuitEntry] = {}
    named = 0
    for digs, nombre in found.items():
        cuit = normalizar_cuit(digs)
        if nombre:
            named += 1
        updates[_key(cuit)] = CuitEntry(
            cuit=cuit,
            razon_social=nombre,
            status="ready",
            note="CUIT de Marti.xlsx. Clave NO persistida: solo autofill de Chrome RECEPCION.",
        )
    merge_entries(updates)
    after = list_cuits()
    nuevos = sum(1 for e in after if e.cuit not in before)
    print(
        f"CUITs únicos en Excel: {len(found)} | "
        f"con nombre: {named} | nuevos en registry: {nuevos} | "
        f"registry total: {len(after)} (antes {len(before)}) | claves guardadas: 0"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
