"""Completa tenencia en FCI GROWTH_FIFO con VC del extracto 30/06/2026."""
from __future__ import annotations

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

OUT = r"C:\Users\recep\Desktop\FCI GROWTH_FIFO.xlsx"
VC = 81.856194
CUOTAS = 18440.04
COSTO = 1498206.63
VAL = round(CUOTAS * VC, 2)
TEN = round(VAL - COSTO, 2)

MONEY = '_-"$"\\ * #,##0.00_-;\\-"$"\\ * #,##0.00_-;_-"$"\\ * "-"??_-;_-@_-'
VC_FMT = "\\$\\ 0.000000"
CUOTAS_FMT = "#,##0.00"
YELLOW = PatternFill("solid", fgColor="FFF2CC")
BOLD = Font(name="Calibri", size=8, bold=True)
BODY = Font(name="Calibri", size=8)
ITAL = Font(name="Calibri", size=8, italic=True, color="666666")


def main() -> None:
    wb = load_workbook(OUT)
    ws = wb["FCI"]
    last = 10
    for r in range(10, 600):
        if ws.cell(r, 3).value in ("RESCATE", "SUSCRIPCION"):
            last = r

    for r in range(last + 1, last + 15):
        for c in range(2, 5):
            ws.cell(r, c).value = None

    r = last + 2
    ws.cell(r, 2, "Resultado por tenencia al 30/06/2026").font = BOLD
    r += 1
    ws.cell(r, 2, "Valor cuota (ultimo extracto del mes)").font = BODY
    c = ws.cell(r, 3, VC)
    c.number_format = VC_FMT
    c.fill = YELLOW
    c.font = BOLD
    vc_row = r
    r += 1
    ws.cell(r, 2, "Cuotas saldo final").font = BODY
    c = ws.cell(r, 3, CUOTAS)
    c.number_format = CUOTAS_FMT
    c.font = BODY
    cuotas_row = r
    r += 1
    ws.cell(r, 2, "Valuacion al cierre (cuotas x VC extracto)").font = BODY
    c = ws.cell(r, 3, f"=C{cuotas_row}*C{vc_row}")
    c.number_format = MONEY
    c.font = BODY
    val_row = r
    r += 1
    ws.cell(r, 2, "Costo de origen de las suscripciones del saldo").font = BODY
    c = ws.cell(r, 3, COSTO)
    c.number_format = MONEY
    c.font = BODY
    costo_row = r
    r += 1
    ws.cell(r, 2, "Resultado por tenencia (valuacion - costo origen)").font = BOLD
    c = ws.cell(r, 3, f"=C{val_row}-C{costo_row}")
    c.number_format = MONEY
    c.font = BOLD
    r += 2
    ws.cell(
        r,
        2,
        f"Fuente VC: extracto de fondos al 30/06/2026 — FIMA PREMIUM CLASE A — "
        f"{CUOTAS:,.2f} cuotas x {VC} = $ {VAL:,.2f}. "
        f"Resultado por tenencia = $ {TEN:,.2f}.",
    ).font = ITAL

    ws.cell(
        last + 1,
        2,
        "Vista FIFO: cada RESCATE se parte si consume varios lotes. "
        "G = (VC rescate - VC origen) x cuotas. H/I = formula a la suscripcion (o SALDO INICIAL). "
        "Fila 3 = intereses mensuales (SUMIFS). VC de cierre = ultimo extracto de fondos del mes.",
    ).font = ITAL

    wb.save(OUT)
    print("OK", OUT)
    print("valuacion", VAL, "tenencia", TEN)


if __name__ == "__main__":
    main()
