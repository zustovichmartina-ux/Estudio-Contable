"""Agrega hoja 'Insumos Alcon' con explicación + imágenes al Excel de compras Rele."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage

XLSX = Path(r"C:\Users\recep\Desktop\Compras_Oftalmologia_Rele_Julio2026.xlsx")
ASSETS = Path(
    r"C:\Users\recep\.cursor\projects\c-Users-recep-Desktop-Estudio-Contable\assets"
)
TMP = Path(r"C:\Users\recep\Desktop\Estudio Contable\_tmp_alcon_imgs")
OUT_IMG = TMP / "_embed"

COLOR_PRIMARIO = "1F4E79"
HDR_FILL = PatternFill("solid", fgColor=COLOR_PRIMARIO)
HDR_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(name="Calibri", bold=True, size=16, color=COLOR_PRIMARIO)
SUB_FONT = Font(name="Calibri", size=11, color="666666")
SECTION_FONT = Font(name="Calibri", bold=True, size=12, color=COLOR_PRIMARIO)
BODY = Font(name="Calibri", size=11)
BOLD = Font(name="Calibri", bold=True, size=11)
ZEBRA = PatternFill("solid", fgColor="F2F2F2")
THIN = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)


def shrink(src: Path, dst: Path, max_w: int) -> Path:
    im = PILImage.open(src)
    if im.mode in ("RGBA", "P"):
        bg = PILImage.new("RGB", im.size, (255, 255, 255))
        if im.mode == "P":
            im = im.convert("RGBA")
        bg.paste(im, mask=im.split()[-1] if im.mode == "RGBA" else None)
        im = bg
    elif im.mode != "RGB":
        im = im.convert("RGB")
    w, h = im.size
    if w > max_w:
        ratio = max_w / w
        im = im.resize((max_w, int(h * ratio)), PILImage.Resampling.LANCZOS)
    im.save(dst, format="JPEG", quality=85)
    return dst


def main() -> None:
    OUT_IMG.mkdir(parents=True, exist_ok=True)
    iol = shrink(ASSETS / "clareon_iol_ref.png", OUT_IMG / "iol.jpg", 420)
    cart = shrink(ASSETS / "monarch_cartridge_ref.png", OUT_IMG / "cart.jpg", 420)
    fc1 = shrink(
        TMP / "ALCON - FC 256442 -HOJA 1.jpeg", OUT_IMG / "fc256442.jpg", 520
    )

    wb = load_workbook(XLSX)
    if "Insumos Alcon" in wb.sheetnames:
        del wb["Insumos Alcon"]
    ws = wb.create_sheet("Insumos Alcon", 2)

    ws["A1"] = "Insumos Alcon — para qué se usan (cirugía de cataratas)"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:F1")
    ws["A2"] = (
        "Oftalmología Rele Mar del Plata SRL · Compras julio 2026 · "
        "Fuente: facturas Alcon + ficha técnica myalcon / IFU"
    )
    ws["A2"].font = SUB_FONT
    ws.merge_cells("A2:F2")

    ws["A4"] = "En simple"
    ws["A4"].font = SECTION_FONT
    ws["A5"] = (
        "Alcon les vendió el lente que se deja adentro del ojo cuando operan una "
        "catarata (Clareon SY60WF) y el cartucho descartable con el que el cirujano "
        "lo inyecta (Monarch III D). No es material de consultorio: son insumos de "
        "quirófano, un set por ojo."
    )
    ws["A5"].font = BODY
    ws["A5"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A5:F5")
    ws.row_dimensions[5].height = 48

    headers = [
        "Producto en factura",
        "Qué es",
        "Para qué se usa",
        "Cirugía",
        "Facturas julio 2026",
        "Cant. aprox.",
    ]
    for i, h in enumerate(headers, 1):
        c = ws.cell(8, i, h)
        c.fill = HDR_FILL
        c.font = HDR_FONT
        c.alignment = Alignment(wrap_text=True, vertical="center")

    rows = [
        (
            "Clareon SY60WF\n(ej. SY60WF.205 / 21.0 / 26.0…)",
            "Lente intraocular (LIO) monofocal, acrílico hidrófobo plegable",
            "Reemplaza el cristalino opaco del paciente. El número (20.5, 21.0, "
            "26.0…) es la potencia en dioptrías elegida para ese ojo.",
            "Cirugía de cataratas (facoemulsificación). Se coloca en el saco "
            "capsular.",
            "FC 256442, 256991, 277747",
            "10 lentes",
        ),
        (
            "Monarch III (D) Cartridges\nSGL USE",
            "Cartucho estéril de un solo uso del sistema Monarch III",
            "Sirve para plegar e inyectar la LIO a través de una incisión "
            "pequeña. Se usa junto con el mango reutilizable Monarch III "
            "(no viene en estas facturas).",
            "Misma cirugía de cataratas: es el “inyector” descartable del lente.",
            "FC 256442 (4), 256991 (4), 277747 (2)",
            "10 cartuchos",
        ),
    ]
    for r_i, row in enumerate(rows):
        excel_r = 9 + r_i
        for c_i, val in enumerate(row, 1):
            cell = ws.cell(excel_r, c_i, val)
            cell.font = BODY
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = THIN
            if r_i % 2 == 1:
                cell.fill = ZEBRA
        ws.row_dimensions[excel_r].height = 78

    ws["A12"] = "Clasificación contable sugerida"
    ws["A12"].font = SECTION_FONT
    ws["A13"] = (
        "Ambos son insumos quirúrgicos inventariables (activo / stock de "
        "materiales de cirugía), no gastos de oficina. Se consumen por cirugía: "
        "1 LIO + 1 cartucho por ojo operado."
    )
    ws["A13"].font = BODY
    ws["A13"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A13:F13")
    ws.row_dimensions[13].height = 36

    ws["A15"] = "Imágenes de referencia (ilustrativas)"
    ws["A15"].font = SECTION_FONT
    ws["A16"] = "Lente intraocular tipo Clareon (LIO)"
    ws["A16"].font = BOLD
    ws["D16"] = "Cartucho Monarch III (descartable)"
    ws["D16"].font = BOLD

    img1 = XLImage(str(iol))
    img1.anchor = "A17"
    ws.add_image(img1)

    img2 = XLImage(str(cart))
    img2.anchor = "D17"
    ws.add_image(img2)

    ws["A36"] = "Ejemplo de factura (FC 0020-00256442)"
    ws["A36"].font = SECTION_FONT
    ws["A37"] = (
        "Detalle real de compra: 4 LIO Clareon SY60WF de distintas potencias "
        "+ 4 cartuchos Monarch III D"
    )
    ws["A37"].font = SUB_FONT
    ws.merge_cells("A37:F37")

    img3 = XLImage(str(fc1))
    img3.anchor = "A38"
    ws.add_image(img3)

    ws["A58"] = "Referencias"
    ws["A58"].font = SECTION_FONT
    ws["A59"] = (
        "https://www.myalcon.com/es/professional/cataract-surgery/iols/clareon-monofocal/"
    )
    ws["A60"] = (
        "https://www.myalcon.com/international/professional/cataract-surgery/"
        "delivery-systems/monarch-iii/"
    )
    ws["A61"] = (
        "Indicación Clareon: corrección de afaquia en adultos tras cirugía de "
        "catarata; implantación en saco capsular."
    )
    for r in (59, 60, 61):
        ws.cell(r, 1).font = SUB_FONT
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)

    widths = [34, 28, 42, 32, 28, 14]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A8"
    wb.save(XLSX)
    print("OK", XLSX)


if __name__ == "__main__":
    main()
