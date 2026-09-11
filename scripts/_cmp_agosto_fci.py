# -*- coding: utf-8 -*-
"""Compara intereses de agosto: hoja manual vs PDF vs motor."""
from __future__ import annotations

import sys
from datetime import date, datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from extraccion_fci import parsear_pdf_fci
from inversiones import aplicar_fifo, identificar_movimientos

INV = Path(
    r"T:\CLIENTES\GROWTH MARKETING SA\Cierre 2025 01-07-2025 al 30-06-2026"
    r"\Papeles de Trabajo\Inversiones"
)
HOJA = Path(r"c:\Users\recep\Desktop\FCI GROWTH.xlsx")


def _fecha(v):
    if isinstance(v, datetime):
        return v.date()
    return v


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    wb = load_workbook(HOJA, data_only=False)
    ws = wb["FCI"]
    wb2 = load_workbook(HOJA, data_only=True)
    ws2 = wb2["FCI"]
    print("D3 cached (tu suma agosto) =", ws2["D3"].value)
    print()
    print("=== TU HOJA AGOSTO ===")
    for r in range(41, 59):
        tipo = ws.cell(r, 3).value
        if tipo not in {"RESCATE", "SUSCRIPCION"}:
            continue
        fd = _fecha(ws.cell(r, 2).value)
        print(
            f"R{r} {fd} {tipo:12} cuotas={ws.cell(r,4).value} "
            f"VC={ws.cell(r,5).value} tot={ws.cell(r,6).value} "
            f"G={ws.cell(r,7).value} => {ws2.cell(r,7).value}"
        )

    pdf = INV / "Extracto_Inversiones_Galicia_2025_08_29.pdf"
    parsed = parsear_pdf_fci(pdf.read_bytes(), pdf.name)
    print()
    print("=== PDF AGOSTO ===")
    for f in parsed["filas"]:
        print(f["Fecha"], f["Descripcion"], f["Cantidad de Cuotas"], f["Valor CC"], f["Total"])

    # Motor jul+ago desde apertura 30/06 (mismos PDFs)
    filas = []
    for nombre in (
        "Extracto_Inversiones_Galicia_2025_07_31.pdf",
        "Extracto_Inversiones_Galicia_2025_08_29.pdf",
    ):
        p = parsear_pdf_fci((INV / nombre).read_bytes(), nombre)
        filas.extend(p["filas"])
    df = pd.DataFrame(filas)
    df["_fecha"] = df["Fecha"]
    df["Especie"] = "FIMA PREMIUM CLASE A"
    df["Tipo_Operacion"] = df["Descripcion"].map(
        lambda x: "Rescate" if "RESCATE" in str(x).upper() else "Suscripcion"
    )
    df["Cantidad"] = df["Cantidad de Cuotas"]
    df["Precio"] = df["Valor CC"]
    df["Monto_Total"] = df["Total"]
    df["Moneda"] = "ARS"
    df["Nueva_Clasificacion"] = None
    df = identificar_movimientos(df)
    ini = pd.DataFrame([{
        "Especie": "FIMA PREMIUM CLASE A",
        "Cantidad": 310323.99,
        "Costo_Unitario": 64.0511,
        "Moneda": "ARS",
    }])
    res = aplicar_fifo(df, ini)
    print()
    print("=== MOTOR: cada tramo de rescate AGOSTO ===")
    tot = 0.0
    for a in res.aplicaciones:
        fe = datetime.strptime(a["Fecha_salida"], "%d/%m/%Y").date()
        if fe.month != 8:
            continue
        lote = a.get("Fecha_lote")
        qty = float(a["Cantidad"])
        vc_o = float(a["Valor_CC_salida"])
        vc_l = float(a["Valor_CC_lote"])
        inte = round((vc_o - vc_l) * qty, 2)
        tot += inte
        print(f"  {a['Fecha_salida']} lote={lote} qty={qty} (VC {vc_o}-{vc_l})*{qty} = {inte}")
    print("TOTAL MOTOR AGOSTO", round(tot, 2))

    print()
    print("=== MOTOR vs TU G por fecha de rescate ===")
    motor_por_fecha = {}
    for a in res.aplicaciones:
        fe = a["Fecha_salida"]
        d = datetime.strptime(fe, "%d/%m/%Y").date()
        if d.month != 8:
            continue
        motor_por_fecha.setdefault(fe, 0.0)
        motor_por_fecha[fe] += float(a["Interes"] or 0)
    for r in range(41, 59):
        if ws.cell(r, 3).value != "RESCATE":
            continue
        fd = _fecha(ws.cell(r, 2).value)
        key = fd.strftime("%d/%m/%Y")
        gval = ws2.cell(r, 7).value
        mval = motor_por_fecha.get(key)
        # puede haber 2 rescates el mismo día
        print(f"  hoja R{r} {key} G={gval}  (ver motor agrupado {mval})")


if __name__ == "__main__":
    main()
