"""Vista previa FCI en formato FCI GROWTH: origen por fórmula + intereses mensuales.

No modifica el motor ni FCI GROWTH.xlsx original. Salida:
  C:\\Users\\recep\\Desktop\\FCI GROWTH_FIFO.xlsx
"""
from __future__ import annotations

import sys
from collections import deque
from copy import copy
from datetime import date, datetime
from pathlib import Path
from shutil import copy2

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SRC = Path(r"C:\Users\recep\Desktop\FCI GROWTH.xlsx")
OUT = Path(r"C:\Users\recep\Desktop\FCI GROWTH_FIFO.xlsx")

MONEY_FMT = '_-"$"\\ * #,##0.00_-;\\-"$"\\ * #,##0.00_-;_-"$"\\ * "-"??_-;_-@_-'
CUOTAS_FMT = "#,##0.00"
VC_FMT = "\\$\\ 0.000000"
DATE_FMT = "dd/mm/yyyy"
MESES_COLS = list("CDEFGHIJKLMN")  # jul-25 … jun-26


def _as_date(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return None


def _copiar_estilo(src, dst) -> None:
    if src.has_style:
        dst.font = copy(src.font)
        dst.fill = copy(src.fill)
        dst.border = copy(src.border)
        dst.alignment = copy(src.alignment)
        dst.number_format = src.number_format


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    if not SRC.exists():
        raise SystemExit(f"No está {SRC}")

    copy2(SRC, OUT)
    wb = load_workbook(OUT)
    ws = wb["FCI"]

    # Estilos plantilla (antes de borrar datos)
    estilo_fecha = ws["B10"]
    estilo_tipo = ws["C10"]
    estilo_cuotas = ws["D10"]
    estilo_vc = ws["E10"]
    estilo_total = ws["F10"]
    estilo_hdr = ws["B9"]

    movs: list[dict] = []
    for r in range(10, (ws.max_row or 10) + 1):
        tipo = str(ws.cell(r, 3).value or "").strip().upper()
        if tipo not in {"RESCATE", "SUSCRIPCION"}:
            continue
        fecha = _as_date(ws.cell(r, 2).value)
        if fecha is None:
            continue
        cant = float(ws.cell(r, 4).value)
        vc = float(ws.cell(r, 5).value)
        tot = ws.cell(r, 6).value
        tot_f = float(tot) if tot not in (None, "") else round(cant * vc, 2)
        movs.append({"fecha": fecha, "tipo": tipo, "cant": cant, "vc": vc, "tot": tot_f})

    ini_cant = float(ws["C7"].value)

    # Noviembre en G2 venía como serial 45962 con formato moneda
    ws["G2"] = datetime(2025, 11, 1)
    ws["G2"].number_format = "mm-dd-yy"
    if ws["C2"].font:
        ws["G2"].font = copy(ws["C2"].font)
    if ws["C2"].fill:
        ws["G2"].fill = copy(ws["C2"].fill)

    # Limpiar bloque de movimientos (sobra espacio por si FIFO parte filas)
    for r in range(9, 600):
        for c in range(2, 10):
            cell = ws.cell(r, c)
            cell.value = None

    # Encabezados (mismo estilo de la fila 9 original)
    headers = {
        "B9": "Fecha",
        "C9": "Descripcion",
        "D9": "Cantidad de Cuotas",
        "E9": "Valor CC",
        "F9": "Total",
        "G9": "Interés",
        "H9": "Origen suscripción",
        "I9": "VC origen",
    }
    for addr, texto in headers.items():
        cell = ws[addr]
        cell.value = texto
        _copiar_estilo(estilo_hdr, cell)
        cell.font = Font(name="Calibri", size=8, bold=True, color="242021")

    # Cola FIFO: (cuotas restantes, fila Excel del lote, es_saldo_inicial)
    cola: deque[list] = deque()
    cola.append([ini_cant, 7, True])

    r = 10
    rescates_escritos = 0
    suscripciones = 0
    interes_check = 0.0

    for mov in movs:
        if mov["tipo"] == "SUSCRIPCION":
            ws.cell(r, 2, datetime(mov["fecha"].year, mov["fecha"].month, mov["fecha"].day))
            ws.cell(r, 3, "SUSCRIPCION")
            ws.cell(r, 4, mov["cant"])
            ws.cell(r, 5, mov["vc"])
            ws.cell(r, 6, mov["tot"])
            for c, est in (
                (2, estilo_fecha),
                (3, estilo_tipo),
                (4, estilo_cuotas),
                (5, estilo_vc),
                (6, estilo_total),
            ):
                _copiar_estilo(est, ws.cell(r, c))
                ws.cell(r, c).fill = PatternFill()  # suscripciones sin pintar
            ws.cell(r, 2).number_format = DATE_FMT
            ws.cell(r, 4).number_format = CUOTAS_FMT
            ws.cell(r, 5).number_format = VC_FMT
            ws.cell(r, 6).number_format = MONEY_FMT
            cola.append([mov["cant"], r, False])
            suscripciones += 1
            r += 1
            continue

        # RESCATE: una fila por lote FIFO consumido
        restante = mov["cant"]
        partes: list[tuple[float, int, bool]] = []
        while restante > 1e-12:
            if not cola:
                raise SystemExit(f"Sin stock en rescate {mov['fecha']} ({mov['cant']})")
            lote = cola[0]
            toma = min(float(lote[0]), restante)
            partes.append((toma, int(lote[1]), bool(lote[2])))
            lote[0] = float(lote[0]) - toma
            restante -= toma
            if lote[0] <= 1e-12:
                cola.popleft()

        totales_parciales: list[float] = []
        for i, (toma, fila_origen, es_ini) in enumerate(partes):
            ultimo = i == len(partes) - 1
            if ultimo:
                tot = round(mov["tot"] - sum(totales_parciales), 2)
            else:
                tot = round(toma * mov["vc"], 2)
                totales_parciales.append(tot)

            ws.cell(r, 2, datetime(mov["fecha"].year, mov["fecha"].month, mov["fecha"].day))
            ws.cell(r, 3, "RESCATE")
            ws.cell(r, 4, round(toma, 6))
            ws.cell(r, 5, mov["vc"])
            ws.cell(r, 6, tot)

            if es_ini:
                ws.cell(r, 7, f"=(E{r}-D$7)*D{r}")
                ws.cell(r, 8, "=$B$7")
                ws.cell(r, 9, "=D$7")
                interes_check += (mov["vc"] - float(ws["D7"].value)) * toma
            else:
                ws.cell(r, 7, f"=(E{r}-E${fila_origen})*D{r}")
                ws.cell(r, 8, f"=B${fila_origen}")
                ws.cell(r, 9, f"=E${fila_origen}")
                # VC origen lo leemos del movimiento original vía fila — no disponible aquí;
                # el check lo hace Excel. Acumulamos con el valor del lote desde cola no;
                # usamos E de la fila origen ya escrita.
                vc_origen = float(ws.cell(fila_origen, 5).value)
                interes_check += (mov["vc"] - vc_origen) * toma

            for c, est in (
                (2, estilo_fecha),
                (3, estilo_tipo),
                (4, estilo_cuotas),
                (5, estilo_vc),
                (6, estilo_total),
            ):
                _copiar_estilo(est, ws.cell(r, c))
            ws.cell(r, 2).number_format = DATE_FMT
            ws.cell(r, 4).number_format = CUOTAS_FMT
            ws.cell(r, 5).number_format = VC_FMT
            ws.cell(r, 6).number_format = MONEY_FMT
            ws.cell(r, 7).number_format = MONEY_FMT
            ws.cell(r, 7).font = Font(name="Calibri", size=8)
            ws.cell(r, 8).number_format = DATE_FMT
            ws.cell(r, 8).font = Font(name="Calibri", size=8)
            ws.cell(r, 8).alignment = Alignment(horizontal="center")
            ws.cell(r, 9).number_format = VC_FMT
            ws.cell(r, 9).font = Font(name="Calibri", size=8)

            rescates_escritos += 1
            r += 1

    last = r - 1

    # Cuadro mensual: SUMIFS sobre Interés (col G) por mes de fila 2
    for i, col in enumerate(MESES_COLS):
        if i + 1 < len(MESES_COLS):
            nxt = MESES_COLS[i + 1]
            formula = (
                f'=SUMIFS($G$10:$G${last},$B$10:$B${last},">="&{col}$2,'
                f'$B$10:$B${last},"<"&{nxt}$2)'
            )
        else:
            formula = (
                f'=SUMIFS($G$10:$G${last},$B$10:$B${last},">="&{col}$2,'
                f'$B$10:$B${last},"<"&EDATE({col}$2,1))'
            )
        cell = ws[f"{col}3"]
        cell.value = formula
        cell.number_format = MONEY_FMT

    ws["O3"] = "=SUM(C3:N3)"
    ws["O3"].number_format = MONEY_FMT

    # Fila 4 ya apunta a fila 3 (=+C3 …); asegurar formato
    for col in MESES_COLS + ["O"]:
        cell = ws[f"{col}4"]
        if not cell.value:
            cell.value = f"=+{col}3"
        cell.number_format = MONEY_FMT

    ws.column_dimensions["G"].width = 14
    ws.column_dimensions["H"].width = 18
    ws.column_dimensions["I"].width = 14

    # Nota breve debajo del bloque
    nota_row = last + 2
    ws.cell(
        nota_row,
        2,
        "Vista FIFO: cada RESCATE se parte si consume varios lotes. "
        "G = (VC rescate − VC origen) × cuotas. "
        "H/I = fórmula a la suscripción (o SALDO INICIAL). "
        "Fila 3 = intereses mensuales (SUMIFS). No modifica el archivo original.",
    ).font = Font(name="Calibri", size=8, italic=True, color="666666")

    wb.save(OUT)

    cuotas_fin = sum(float(x[0]) for x in cola)
    print("OK", OUT)
    print("Suscripciones", suscripciones, "filas rescate (tramos FIFO)", rescates_escritos)
    print("Última fila datos", last)
    print("Interés realizado (check)", round(interes_check, 2))
    print("Cuotas restantes cola", round(cuotas_fin, 4))
    print("Lotes abiertos", len(cola))
    for lote in cola:
        print(" ", "INI" if lote[2] else f"fila {lote[1]}", round(float(lote[0]), 4))


if __name__ == "__main__":
    main()
