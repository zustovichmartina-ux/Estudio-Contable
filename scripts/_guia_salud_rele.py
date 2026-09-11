"""Hojas guía (estilo Alcon) para insumos/equipos de salud — Rele jul/2026."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from PIL import Image as PILImage

XLSX = Path(r"C:\Users\recep\Desktop\Compras_Oftalmologia_Rele_Julio2026.xlsx")
ASSETS = Path(
    r"C:\Users\recep\.cursor\projects\c-Users-recep-Desktop-Estudio-Contable\assets"
)
TMP = Path(r"C:\Users\recep\Desktop\Estudio Contable\_tmp_rele_guia")
ALCON_TMP = Path(r"C:\Users\recep\Desktop\Estudio Contable\_tmp_alcon_imgs")
OUT_IMG = TMP / "_embed"

COLOR_PRIMARIO = "1F4E79"
HDR_FILL = PatternFill("solid", fgColor=COLOR_PRIMARIO)
HDR_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(name="Calibri", bold=True, size=16, color=COLOR_PRIMARIO)
SUB_FONT = Font(name="Calibri", size=11, color="666666")
SECTION_FONT = Font(name="Calibri", bold=True, size=12, color=COLOR_PRIMARIO)
BODY = Font(name="Calibri", size=11)
BOLD = Font(name="Calibri", bold=True, size=11)
ZEBRA = PatternFill("solid", fgColor="F2F2F2")
THIN = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)
HEADERS = [
    "Producto en factura",
    "Qué es",
    "Para qué se usa",
    "Cirugía / uso clínico",
    "Factura",
    "Cant. / monto",
]


def shrink(src: Path, dst: Path, max_w: int) -> Path:
    im = PILImage.open(src)
    if im.mode in ("RGBA", "P"):
        bg = PILImage.new("RGB", im.size, (255, 255, 255))
        if im.mode == "P":
            im = im.convert("RGBA")
        bg.paste(im, mask=im.split()[-1] if im.mode == "RGBA" else None)
        im = bg
    elif im.mode != "RGB":
        im = im.convert("RGB")
    w, h = im.size
    if w > max_w:
        ratio = max_w / w
        im = im.resize((max_w, int(h * ratio)), PILImage.Resampling.LANCZOS)
    im.save(dst, format="JPEG", quality=85)
    return dst


def style_title(ws: Worksheet, title: str, subtitle: str) -> None:
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:F1")
    ws["A2"] = subtitle
    ws["A2"].font = SUB_FONT
    ws.merge_cells("A2:F2")


def write_simple(ws: Worksheet, text: str, row: int = 4) -> int:
    ws.cell(row, 1, "En simple").font = SECTION_FONT
    cell = ws.cell(row + 1, 1, text)
    cell.font = BODY
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=row + 1, start_column=1, end_row=row + 1, end_column=6)
    ws.row_dimensions[row + 1].height = 52
    return row + 1


def write_table(ws: Worksheet, rows: list[tuple], start_row: int = 8) -> int:
    for i, h in enumerate(HEADERS, 1):
        c = ws.cell(start_row, i, h)
        c.fill = HDR_FILL
        c.font = HDR_FONT
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for r_i, row in enumerate(rows):
        excel_r = start_row + 1 + r_i
        for c_i, val in enumerate(row, 1):
            cell = ws.cell(excel_r, c_i, val)
            cell.font = BODY
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = THIN
            if r_i % 2 == 1:
                cell.fill = ZEBRA
        ws.row_dimensions[excel_r].height = 72
    return start_row + 1 + len(rows)


def write_clasif(ws: Worksheet, text: str, row: int) -> int:
    ws.cell(row, 1, "Clasificación contable sugerida").font = SECTION_FONT
    cell = ws.cell(row + 1, 1, text)
    cell.font = BODY
    cell.alignment = Alignment(wrap_text=True)
    ws.merge_cells(start_row=row + 1, start_column=1, end_row=row + 1, end_column=6)
    ws.row_dimensions[row + 1].height = 36
    return row + 1


def add_images(
    ws: Worksheet,
    labels_and_paths: list[tuple[str, Path]],
    start_row: int,
    invoice: tuple[str, str, Path] | None = None,
) -> None:
    ws.cell(start_row, 1, "Imágenes de referencia (ilustrativas)").font = SECTION_FONT
    # 2 columnas de imágenes
    for i, (label, path) in enumerate(labels_and_paths[:2]):
        col = 1 if i == 0 else 4
        ws.cell(start_row + 1, col, label).font = BOLD
        img = XLImage(str(path))
        img.anchor = f"{get_column_letter(col)}{start_row + 2}"
        ws.add_image(img)
    row_fc = start_row + 22
    if invoice:
        titulo, sub, path = invoice
        ws.cell(row_fc, 1, titulo).font = SECTION_FONT
        ws.cell(row_fc + 1, 1, sub).font = SUB_FONT
        ws.merge_cells(
            start_row=row_fc + 1, start_column=1, end_row=row_fc + 1, end_column=6
        )
        img = XLImage(str(path))
        img.anchor = f"A{row_fc + 2}"
        ws.add_image(img)


def finalize_sheet(ws: Worksheet) -> None:
    widths = [34, 28, 42, 32, 26, 16]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A8"


def replace_sheet(wb, name: str, index: int | None = None) -> Worksheet:
    if name in wb.sheetnames:
        del wb[name]
    if index is None:
        return wb.create_sheet(name)
    return wb.create_sheet(name, index)


def main() -> None:
    OUT_IMG.mkdir(parents=True, exist_ok=True)

    iol = shrink(ASSETS / "clareon_iol_ref.png", OUT_IMG / "iol.jpg", 400)
    cart = shrink(ASSETS / "monarch_cartridge_ref.png", OUT_IMG / "cart.jpg", 400)
    ovd = shrink(ASSETS / "viscotec_ovd_ref.png", OUT_IMG / "ovd.jpg", 400)
    lamp = shrink(ASSETS / "slit_lamp_ref.png", OUT_IMG / "lamp.jpg", 400)
    knife = shrink(ASSETS / "mani_knife_ref.png", OUT_IMG / "knife.jpg", 400)
    drops = shrink(ASSETS / "colirio_ampolla_ref.png", OUT_IMG / "drops.jpg", 400)
    tono = shrink(ASSETS / "tonometro_adap_ref.png", OUT_IMG / "tono.jpg", 400)

    fc_biomat = shrink(
        TMP / "FC A 00004-00071964 -BIOMAT.jpeg", OUT_IMG / "fc_biomat.jpg", 500
    )
    fc_gsj = shrink(TMP / "FC GSJ SA- 70515.jpeg", OUT_IMG / "fc_gsj.jpg", 500)
    fc_msz = shrink(TMP / "MSZ - FC 24086.jpeg", OUT_IMG / "fc_msz.jpg", 500)
    fc_farm = shrink(
        TMP / "FC FARMACIA MAGISTER - GAONA 3050SRL.jpeg",
        OUT_IMG / "fc_farm.jpg",
        420,
    )
    # PDF pages → use rendered? skip invoice photo if PDF only; use placeholder note
    # For Implantec/Med we don't have JPEG of full invoice easily — skip or convert
    # Convert first page of key PDFs with pymupdf if available
    try:
        import fitz

        for pdf_name, jpg_name in [
            ("FC IMPLANTEC -006_00012_000020240.pdf", "fc_implantec.jpg"),
            ("med lampara FA-A 00004-00046610.pdf", "fc_med_lamp.jpg"),
            ("FA-A 00004-00046797 - MED SRL.pdf", "fc_med_tono.jpg"),
        ]:
            doc = fitz.open(TMP / pdf_name)
            page = doc[0]
            pix = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5))
            raw = OUT_IMG / f"_raw_{jpg_name}"
            pix.save(str(raw))
            shrink(raw, OUT_IMG / jpg_name, 520)
            doc.close()
        fc_implantec = OUT_IMG / "fc_implantec.jpg"
        fc_med_lamp = OUT_IMG / "fc_med_lamp.jpg"
        fc_med_tono = OUT_IMG / "fc_med_tono.jpg"
    except Exception:
        fc_implantec = fc_biomat
        fc_med_lamp = fc_biomat
        fc_med_tono = fc_biomat

    wb = load_workbook(XLSX)

    # ── Índice ──────────────────────────────────────────────────────────────
    ws = replace_sheet(wb, "Índice guías salud", 2)
    style_title(
        ws,
        "Guías de productos de salud / quirófano",
        "Oftalmología Rele · julio 2026 · solo facturas con denominación técnica o insumos clínicos",
    )
    write_simple(
        ws,
        "Estas hojas explican, en criollo, qué compraron cuando la factura dice "
        "códigos raros (SY60WF, Viscotec, PureSee, MSL28SK, LS-4, etc.). "
        "No incluye honorarios médicos, limpieza, fletes, electricidad ni librería.",
    )
    idx_rows = [
        (
            "Insumos Alcon",
            "LIO Clareon SY60WF + cartuchos Monarch III",
            "Cirugía de cataratas",
            "Activo / stock quirófano",
            "3 FC",
            "10 LIO + 10 cartuchos",
        ),
        (
            "Biomat — TECNIS",
            "TECNIS PureSee Simplicity 24.0 (J&J)",
            "LIO premium EDOF + inyector precargado",
            "Activo / stock quirófano",
            "FC 00004-00071964",
            "1 LIO USD 713,90",
        ),
        (
            "GSJ — VisTor",
            "VISTOR 23.00 CYL −3.00 + inyector",
            "LIO tórica (corrige astigmatismo)",
            "Activo / stock quirófano",
            "FC 0060-00070515",
            "1 LIO USD 312,50",
        ),
        (
            "Implantec — Viscotec",
            "Viscotec HPMC 2% x 3 ml (dispersiva)",
            "Gel viscoelástico protector en cirugía",
            "Activo / stock quirófano",
            "FC 00012-00020240",
            "20 packs · $1.976.000",
        ),
        (
            "MSZ — Bisturíes Mani",
            "Bisturíes oftalmológicos de un solo uso",
            "Incisiones de córnea en cataratas",
            "Activo / stock quirófano",
            "FC 0003-00024086",
            "16 u. · $755.820",
        ),
        (
            "Med SRL — Equipos",
            "Lámpara de hendidura LS-4, mesas, adaptador tonómetro",
            "Equipamiento de consultorio / diagnóstico",
            "Activo fijo / equipamiento",
            "FC 46610 + 46797",
            "USD 3.943,70 + 58,57",
        ),
        (
            "Farmacia Magister",
            "Loción, colirio estéril, frascos ampolla",
            "Medicación / preparados oftálmicos",
            "Gasto / insumos farmacia",
            "FC 00010-00010645",
            "$300.000",
        ),
    ]
    # reuse product headers with slight rename
    for i, h in enumerate(
        [
            "Hoja",
            "Producto",
            "Para qué",
            "Clase sugerida",
            "Comprobante",
            "Resumen",
        ],
        1,
    ):
        c = ws.cell(8, i, h)
        c.fill = HDR_FILL
        c.font = HDR_FONT
    for r_i, row in enumerate(idx_rows):
        excel_r = 9 + r_i
        for c_i, val in enumerate(row, 1):
            cell = ws.cell(excel_r, c_i, val)
            cell.font = BODY
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = THIN
            if r_i % 2:
                cell.fill = ZEBRA
        ws.row_dimensions[excel_r].height = 40
    ws["A17"] = "Excluidas a propósito (no son insumos clínicos raros)"
    ws["A17"].font = SECTION_FONT
    ws["A18"] = (
        "OACI (electricidad/LED), Integral Pack y Rabbione (flete), Limpa (limpieza), "
        "Biomat DDCA (solo diferencia de cambio), honorarios médicos, Payway, soda, VETRA."
    )
    ws["A18"].font = SUB_FONT
    ws.merge_cells("A18:F18")
    finalize_sheet(ws)

    # ── Biomat ──────────────────────────────────────────────────────────────
    ws = replace_sheet(wb, "Biomat — TECNIS")
    style_title(
        ws,
        "Biomat — TECNIS PureSee Simplicity 24.0",
        "Factura A 00004-00071964 · 01/07/2026 · Johnson & Johnson Vision vía Biomat Instrumental",
    )
    write_simple(
        ws,
        "Compraron un lente intraocular premium de J&J (no Alcon): TECNIS PureSee. "
        "“Simplicity” es el inyector ya precargado; “24.0” es la potencia en dioptrías. "
        "Es LIO de profundidad de foco extendida (EDOF): mejora visión intermedia vs monofocal.",
    )
    write_table(
        ws,
        [
            (
                "TECNIS PURESEE SIMPLICITY 24.0\n(código DENVVI240)",
                "Lente intraocular EDOF hidrófoba + sistema de inyección precargado Simplicity",
                "Reemplaza el cristalino en cirugía de catarata. Da visión de lejos comparable a monofocal y mejor intermedia (pantalla, tablero).",
                "Cirugía de cataratas / afaquia. Colocación en saco capsular.",
                "FC A00004-00071964",
                "1 u. · USD 713,90",
            )
        ],
    )
    write_clasif(
        ws,
        "Activo / stock de insumos quirúrgicos inventariables (igual lógica que Clareon). "
        "1 lente = 1 ojo. No es gasto de consultorio.",
        11,
    )
    add_images(
        ws,
        [
            ("Lente intraocular (referencia)", iol),
            ("Inyector/cartucho tipo (referencia)", cart),
        ],
        14,
        (
            "Factura Biomat",
            "Detalle: TECNIS PureSee Simplicity 24.0 — despacho Ezeiza / Países Bajos",
            fc_biomat,
        ),
    )
    ws["A58"] = "Referencias"
    ws["A58"].font = SECTION_FONT
    ws["A59"] = "https://www.jnjvisionpro.com/en-us/products/tecnis-puresee/"
    ws["A59"].font = SUB_FONT
    finalize_sheet(ws)

    # ── GSJ VisTor ──────────────────────────────────────────────────────────
    ws = replace_sheet(wb, "GSJ — VisTor")
    style_title(
        ws,
        "GSJ / Rosinov — VisTor 23.00 CYL −3.00",
        "Factura A 0060-00070515 · 20/07/2026 · LIO tórica Hanita + inyector monouso",
    )
    write_simple(
        ws,
        "Es otra lente para cataratas, pero tórica: además de reemplazar el cristalino "
        "corrige astigmatismo (CYL −3.00). “23.00” es la potencia esférica. Viene con "
        "inyector de un solo uso. Marca VisTor (Hanita Lenses), distribuida por GSJ/Rosinov.",
    )
    write_table(
        ws,
        [
            (
                "VISTOR23.00 CYL-3.00\nIntraocular Lenses With Single Use Injector",
                "LIO monofocal asférica tórica + inyector descartable",
                "Cirugía de catarata en pacientes con astigmatismo corneal previo. El cilindro (−3.00) se alinea al eje calculado.",
                "Cirugía de cataratas con corrección de astigmatismo.",
                "FC 0060-00070515",
                "1 u. · USD 312,50",
            )
        ],
    )
    write_clasif(
        ws,
        "Activo / stock quirúrgico inventariable. Misma lógica que Alcon/Biomat.",
        11,
    )
    add_images(
        ws,
        [
            ("Lente intraocular tórica (ref.)", iol),
            ("Inyector monouso (ref.)", cart),
        ],
        14,
        ("Factura GSJ / Rosinov", "VISTOR 23.00 CYL-3.00 + single use injector", fc_gsj),
    )
    ws["A58"] = "Referencias"
    ws["A58"].font = SECTION_FONT
    ws["A59"] = "https://www.hanitalenses.com/intraocular-implants/monofocal/vistor/"
    ws["A59"].font = SUB_FONT
    finalize_sheet(ws)

    # ── Implantec ───────────────────────────────────────────────────────────
    ws = replace_sheet(wb, "Implantec — Viscotec")
    style_title(
        ws,
        "Implantec — Viscotec HPMC 2% (viscoelástica)",
        "Factura 00012-00020240 · 29/07/2026 · 20 packs × 5 jeringas de 3 ml",
    )
    write_simple(
        ws,
        "No es un lente: es un gel estéril (HPMC 2%) que el cirujano inyecta en el ojo "
        "durante la cirugía para mantener la cámara anterior abierta y proteger la córnea "
        "mientras saca la catarata e introduce la LIO. “Dispersiva” = se adhiere bien a los tejidos.",
    )
    write_table(
        ws,
        [
            (
                "SUSTANCIA VISCOELÁSTICA DISPERSIVA\nVISCOTEC HPMC 2% X 3 ML X 5U",
                "Dispositivo viscoquirúrgico oftálmico (OVD) en jeringa",
                "Protege el endotelio corneal, mantiene espacio intraoperatorio y facilita manipular tejidos / implantar la LIO.",
                "Cirugía de segmento anterior (cataratas, IOL, etc.). Se aspira al final.",
                "FC 00012-00020240",
                "20 packs · $1.976.000",
            )
        ],
    )
    write_clasif(
        ws,
        "Activo / stock de insumos quirúrgicos. Se consume por cirugía (no es honorario ni limpieza).",
        11,
    )
    add_images(
        ws,
        [("Jeringa viscoelástica (ref.)", ovd), ("Contexto LIO (ref.)", iol)],
        14,
        (
            "Factura Implantec",
            "20 × Viscotec HPMC 2% dispersiva (origen Argentina)",
            fc_implantec,
        ),
    )
    ws["A58"] = "Referencias"
    ws["A58"].font = SECTION_FONT
    ws["A59"] = "IFU Viscotec / Implantec — OVD segmento anterior"
    ws["A59"].font = SUB_FONT
    finalize_sheet(ws)

    # ── MSZ ─────────────────────────────────────────────────────────────────
    ws = replace_sheet(wb, "MSZ — Bisturíes Mani")
    style_title(
        ws,
        "MSZ — Bisturíes oftalmológicos Mani",
        "Factura 0003-00024086 · 08/07/2026 · 16 unidades descartables",
    )
    write_simple(
        ws,
        "Son cuchillitos estériles de un solo uso para abrir la córnea en cirugía de "
        "cataratas (no bisturí de quirófano general). Cada código es un tipo/ángulo distinto "
        "de incisión. Marca Mani (Japón/Vietnam), proveedor MSZ Medical Supplies.",
    )
    write_table(
        ws,
        [
            (
                "MSL28SK — Bisturí Mani 2.8 mm Safety",
                "Keratome / slit knife 2.8 mm con tapa de seguridad",
                "Incisión principal de córnea (~2.8 mm) para entrar con faco / inyector de LIO.",
                "Cirugía de cataratas",
                "FC 0003-00024086",
                "6 u.",
            ),
            (
                "MVR21ASK — Bisturí 21G angulado Safety",
                "Cuchilla MVR 21G angulada con seguridad",
                "Incisiones auxiliares / paracentesis (puertos laterales).",
                "Cirugía de cataratas / vitrectomía anterior",
                "FC 0003-00024086",
                "6 u.",
            ),
            (
                "MCU26 — Bisturí Crescent Mani",
                "Cuchilla crescent (media luna)",
                "Disección / tunelización de tejidos en segmento anterior.",
                "Cirugía de segmento anterior",
                "FC 0003-00024086",
                "2 u.",
            ),
            (
                "MST15 — Bisturí 15° Mani",
                "Cuchilla puntual 15 grados",
                "Incisiones pequeñas precisas / sideport.",
                "Cirugía oftalmológica",
                "FC 0003-00024086",
                "2 u.",
            ),
        ],
    )
    write_clasif(
        ws,
        "Activo / stock de instrumental descartable de quirófano. Total factura $755.820 (IVA 10,5%).",
        14,
    )
    add_images(
        ws,
        [("Bisturí oftalmológico (ref.)", knife), ("Uso: cirugía de catarata", iol)],
        17,
        ("Factura MSZ", "4 modelos Mani · 16 unidades", fc_msz),
    )
    finalize_sheet(ws)

    # ── Med SRL ─────────────────────────────────────────────────────────────
    ws = replace_sheet(wb, "Med SRL — Equipos")
    style_title(
        ws,
        "Med SRL — Equipamiento de consultorio",
        "FC 00004-00046610 (01/07) + FC 00004-00046797 (29/07) · USD",
    )
    write_simple(
        ws,
        "Acá no hay lentes: son equipos de diagnóstico. La lámpara de hendidura sirve para "
        "examinar el ojo con microscopio; las mesas eléctricas la sostienen; el adaptador "
        "permite medir presión ocular (tonometría de aplanación) sobre esa lámpara.",
    )
    write_table(
        ws,
        [
            (
                "[LS-4] Lámpara de hendidura LS-4 (tipo HS)",
                "Biomicroscopio de consultorio (slit lamp)",
                "Examen de córnea, cristalino, cámara anterior, etc. Base del diagnóstico oftalmológico.",
                "Consultorio / pre y post cirugía",
                "FC 00004-00046610",
                "1 u. (dto. voucher)",
            ),
            (
                "[M120/T01] Mesa eléctrica con tabla chica",
                "Mesa motorizada para equipos",
                "Soporta la lámpara / equipos a altura del paciente.",
                "Consultorio",
                "FC 00004-00046610",
                "2 u.",
            ),
            (
                "[ADAP.TON] Adaptador tonómetro aplanático",
                "Accesorio para tonómetro de Goldmann",
                "Mide presión intraocular (glaucoma / control prequirúrgico) montado en la lámpara.",
                "Consultorio / diagnóstico",
                "FC 00004-00046797",
                "1 u. · USD 58,57",
            ),
        ],
    )
    write_clasif(
        ws,
        "Activo fijo / equipamiento médico (no se consume por cirugía). Totales: USD 3.943,70 + USD 58,57. "
        "Rabbione 0021-00029406 es el flete del pallet desde Med SRL (gasto de logística, no el equipo).",
        13,
    )
    add_images(
        ws,
        [("Lámpara de hendidura (ref.)", lamp), ("Adaptador tonómetro (ref.)", tono)],
        16,
        (
            "Factura Med SRL — lámpara y mesas",
            "LS-4 + 2 mesas eléctricas · Total USD 3.943,70",
            fc_med_lamp,
        ),
    )
    # segunda factura un poco más abajo
    ws["A52"] = "Factura Med SRL — adaptador tonómetro"
    ws["A52"].font = SECTION_FONT
    ws["A53"] = "ADAP.TON · Total USD 58,57"
    ws["A53"].font = SUB_FONT
    img = XLImage(str(fc_med_tono))
    img.anchor = "A54"
    ws.add_image(img)
    finalize_sheet(ws)

    # ── Farmacia ────────────────────────────────────────────────────────────
    ws = replace_sheet(wb, "Farmacia Magister")
    style_title(
        ws,
        "Farmacia Magister — preparados oftálmicos",
        "Tique Factura A 00010-00010645 · 03/07/2026 · $300.000",
    )
    write_simple(
        ws,
        "Compra en farmacia (no quirófano de lentes): loción, colirio estéril y frascos "
        "ampolla. En oftalmología suelen ser preparados magistrales o especialidades para "
        "tratamiento/postoperatorio. Los nombres en el ticket son genéricos (sin principio "
        "activo legible en la foto).",
    )
    write_table(
        ws,
        [
            (
                "LOCION",
                "Loción / preparación tópica",
                "Uso dermatológico u oftalmológico periocular según fórmula.",
                "Tratamiento ambulatorio",
                "FC 00010-00010645",
                "1 u.",
            ),
            (
                "COLIRIO - ESTER",
                "Colirio estéril (gotas oculares)",
                "Medicación tópica ocular (antibiótico, antiinflamatorio u otra según fórmula).",
                "Consultorio / postoperatorio",
                "FC 00010-00010645",
                "25 u.",
            ),
            (
                "FCO. / FRASCO AMPOLLA ES",
                "Frasco ampolla inyectable o reconstituible",
                "Presentación unitaria estéril; en oftalmo a veces para inyecciones o preparación.",
                "Uso clínico / quirúrgico según fórmula",
                "FC 00010-00010645",
                "6 u. total",
            ),
        ],
    )
    write_clasif(
        ws,
        "Gasto / insumos de farmacia (medicamentos). No inventariable como LIO; se consume en atención.",
        13,
    )
    add_images(
        ws,
        [("Colirio y ampollas (ref.)", drops), ("Contexto clínico ocular", iol)],
        16,
        (
            "Ticket Farmacia Magister",
            "Loción + 25 colirios + frascos ampolla · Total $300.000",
            fc_farm,
        ),
    )
    finalize_sheet(ws)

    # Mejorar detalle compras para filas clave
    if "Detalle compras" in wb.sheetnames:
        det = wb["Detalle compras"]
        headers = [c.value for c in det[1]]
        try:
            col_arch = headers.index("Archivo") + 1
            col_det = headers.index("Detalle / conceptos") + 1
            col_obs = headers.index("Observaciones") + 1
            col_clase = headers.index("Clase (Gasto/Activo)") + 1
            col_tipo = headers.index("Tipo") + 1
        except ValueError:
            col_arch = col_det = col_obs = col_clase = col_tipo = None

        updates = {
            "BIOMAT.jpeg": (
                "Activo",
                "Insumos quirúrgicos inventariables",
                "TECNIS PureSee Simplicity 24.0 (LIO EDOF J&J)",
                "Ver hoja Biomat — TECNIS",
            ),
            "GSJ SA": (
                "Activo",
                "Insumos quirúrgicos inventariables",
                "VisTor 23.00 CYL-3.00 LIO tórica + inyector",
                "Ver hoja GSJ — VisTor",
            ),
            "MSZ": (
                "Activo",
                "Insumos quirúrgicos inventariables",
                "Bisturíes Mani MSL28SK / MVR21ASK / MCU26 / MST15",
                "Ver hoja MSZ — Bisturíes Mani",
            ),
            "IMPLANTEC": (
                "Activo",
                "Insumos quirúrgicos inventariables",
                "Viscotec HPMC 2% viscoelástica dispersiva × 20 packs",
                "Ver hoja Implantec — Viscotec",
            ),
            "med lampara": (
                "Activo",
                "Equipamiento médico",
                "Lámpara hendidura LS-4 + 2 mesas eléctricas",
                "Ver hoja Med SRL — Equipos",
            ),
            "00046797": (
                "Activo",
                "Equipamiento médico",
                "Adaptador tonómetro aplanático",
                "Ver hoja Med SRL — Equipos",
            ),
            "FARMACIA MAGISTER": (
                "Gasto",
                "Farmacia / medicamentos",
                "Loción + colirio estéril ×25 + frascos ampolla",
                "Ver hoja Farmacia Magister",
            ),
            "ALCON - FC 256442": (
                None,
                None,
                "CLAREON SY60WF ×4 potencias + Monarch III D ×4",
                "Ver hoja Insumos Alcon",
            ),
            "ALCON - FC 256991": (
                None,
                None,
                "CLAREON SY60WF ×4 potencias + Monarch III D ×4",
                "Ver hoja Insumos Alcon",
            ),
            "ALCON - FC 277747": (
                None,
                None,
                "CLAREON SY60WF ×2 + Monarch III D ×2",
                "Ver hoja Insumos Alcon",
            ),
        }
        if col_arch and col_det:
            for row in det.iter_rows(min_row=2):
                arch = str(row[col_arch - 1].value or "")
                for key, (clase, tipo, detalle, obs) in updates.items():
                    if key.lower() in arch.lower():
                        if detalle:
                            row[col_det - 1].value = detalle
                        if obs and col_obs:
                            row[col_obs - 1].value = obs
                        if clase and col_clase:
                            row[col_clase - 1].value = clase
                        if tipo and col_tipo:
                            row[col_tipo - 1].value = tipo
                        break

    wb.save(XLSX)
    print("OK", XLSX)
    print("Hojas:", wb.sheetnames)


if __name__ == "__main__":
    main()
