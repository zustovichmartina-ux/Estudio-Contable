# -*- coding: utf-8 -*-
"""Extrae carátulas, facturas FCE/A y reportes de Centro Médico (Gastro) a Excel."""
from __future__ import annotations

import re
import sys
from pathlib import Path

import pandas as pd
import pdfplumber
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from caratulas_centro_medico import (  # noqa: E402
    BODY,
    BOLD,
    HDR_FILL,
    HDR_FONT,
    MONEY_FMT,
    ZEBRA,
    consolidar_carpeta,
    parse_us_money,
)

CARPETA = Path(
    r"\\TANGOSRV\Compartido\CLIENTES"
    r"\GASTROENTEROLOGIA Y ENDOSCOPIA DIGESTIVA MAR DEL PLATA S.A"
    r"\Liquidaciones Centro Medico y Cli. Colon\Liquidaciones Centro Medico"
)
OUT = CARPETA / "Liquidaciones_Centro_Medico_consolidado.xlsx"
OUT_DESKTOP = Path(r"C:\Users\recep\Desktop") / "Liquidaciones_Centro_Medico_consolidado.xlsx"


def _texto(path: Path, max_pages: int = 6) -> str:
    with pdfplumber.open(path) as pdf:
        return "\n".join((pg.extract_text() or "") for pg in pdf.pages[:max_pages])


def parse_ar_money(s: str) -> float | None:
    s = re.sub(r"[^\d,.\-]", "", str(s or "").strip())
    if not s or s in ("-", ".", ","):
        return None
    if "," in s and "." in s:
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif "," in s:
        s = s.replace(".", "").replace(",", ".")
    try:
        return round(float(s), 2)
    except ValueError:
        return None


def _grab(text: str, pat: str, ar: bool = True):
    m = re.search(pat, text, re.I | re.S)
    if not m:
        return None
    return parse_ar_money(m.group(1)) if ar else parse_us_money(m.group(1))


def parse_factura_afip(path: Path) -> dict | None:
    try:
        text = _texto(path)
    except Exception:
        return None
    if "Punto de Venta" not in text and "Punto de Venta:" not in text:
        return None
    if "FACTURA" not in text.upper() and "CR" not in text.upper():
        return None
    tipo = "FCE A (201)" if re.search(r"C[OÓ]D\.?\s*201", text, re.I) else (
        "Factura A (01)" if re.search(r"COD\.?\s*01", text, re.I) else "Factura"
    )
    pto = re.search(r"Punto de Venta:\s*(\d+)", text, re.I)
    nro = re.search(r"Comp\.?\s*Nro:\s*(\d+)", text, re.I)
    fem = re.search(r"Fecha de Emisi[oó]n:\s*(\d{1,2}/\d{1,2}/\d{4})", text, re.I)
    vto = re.search(r"Fecha de Vto\. para el pago:\s*(\d{1,2}/\d{1,2}/\d{4})", text, re.I)
    dsd = re.search(r"Desde:\s*(\d{1,2}/\d{1,2}/\d{4})", text, re.I)
    hst = re.search(r"Hasta:\s*(\d{1,2}/\d{1,2}/\d{4})", text, re.I)
    cli = re.search(r"Raz[oó]n Social:([^\n]+)", text[text.find("CUIT: 30") :] if "CUIT: 30" in text else text)
    # receptor after second CUIT
    rec = re.search(
        r"CUIT:\s*(30541190518)[^\n]*Raz[oó]n Social:\s*([^\n]+)",
        text,
        re.I,
    )
    cae = re.search(r"CAE N[°º]?:\s*(\d+)", text, re.I)
    return {
        "Tipo": tipo,
        "Punto venta": pto.group(1) if pto else "",
        "Nro": nro.group(1) if nro else "",
        "Fecha emision": fem.group(1) if fem else "",
        "Vto pago": vto.group(1) if vto else "",
        "Periodo desde": dsd.group(1) if dsd else "",
        "Periodo hasta": hst.group(1) if hst else "",
        "Cliente": (rec.group(2).strip() if rec else "CENTRO MEDICO DE MAR DEL PLATA"),
        "CUIT cliente": rec.group(1) if rec else "30541190518",
        "Honorarios exentos": _grab(text, r"Honorarios Exentos\s+[\d.,]+\s+\S+\s+([\d.,]+)"),
        "Honorarios gravados": _grab(text, r"Honorarios Gravados\s+[\d.,]+\s+\S+\s+([\d.,]+)"),
        "Gastos exentos": _grab(text, r"Gastos Exentos\s+[\d.,]+\s+\S+\s+([\d.,]+)"),
        "Gastos gravados": _grab(text, r"Gastos Gravados\s+[\d.,]+\s+\S+\s+([\d.,]+)"),
        "Importe exento": _grab(text, r"Importe Exento:\s*\$?\s*([\d.,]+)"),
        "Neto gravado": _grab(text, r"Importe Neto Gravado:\s*\$?\s*([\d.,]+)"),
        "IVA 10,5%": _grab(text, r"IVA\s*10[.,]5%\s*:\s*\$?\s*([\d.,]+)"),
        "Importe total": _grab(text, r"Importe Total:\s*\$?\s*([\d.,]+)"),
        "CAE": cae.group(1) if cae else "",
        "Archivo": str(path.relative_to(CARPETA)),
    }


def parse_pendientes(path: Path) -> list[dict]:
    try:
        text = _texto(path)
    except Exception:
        return []
    if "Pendiente" not in text and "PENDIENTE" not in text.upper():
        return []
    tipo = "IVA" if "IVA" in text.upper() else "Honorarios"
    rows = []
    # |02/2026| 523,208.00| ... US money
    pat = re.compile(
        r"\|(\d{2}/\d{4})\|([^|\n]+)\|([^|\n]+)\|([^|\n]+)\|([^|\n]+)\|([^|\n]+)\|([^|\n]+)\|([^|\n]+)\|([^|\n]+)\|"
    )
    for m in pat.finditer(text.replace("\n", "")):
        vals = [parse_us_money(x) for x in m.groups()[1:]]
        rows.append(
            {
                "Tipo reporte": f"Facturas pendientes de {tipo}",
                "Mes/Año": m.group(1),
                "Gravado honorarios": vals[0],
                "Exento honorarios": vals[1],
                "IVA honorarios": vals[2],
                "Gravado gastos": vals[3],
                "Exento gastos": vals[4],
                "IVA gastos": vals[5],
                "Total IVA": vals[6],
                "Total general": vals[7],
                "Archivo": str(path.relative_to(CARPETA)),
            }
        )
    return rows


def parse_resumen_percibido(path: Path) -> dict | None:
    try:
        text = _texto(path)
    except Exception:
        return None
    if "Resumen Percibido" not in text and "LIQUIDACION A PROFESIONALES" not in text.upper():
        return None
    per = re.search(r"entre el\s*(\d{2}/\d{2}/\d{4})\s*y el\s*(\d{2}/\d{2}/\d{4})", text, re.I)
    def linea(nombre: str) -> float | None:
        m = re.search(rf"{nombre}\s+([\d.,]+)", text, re.I)
        return parse_us_money(m.group(1)) if m else None
    return {
        "Desde": per.group(1) if per else "",
        "Hasta": per.group(2) if per else "",
        "Facturado bruto honorarios": linea("FACTURADO BRUTO HONORARIOS"),
        "Facturado bruto gastos": linea("FACTURADO BRUTO GASTOS"),
        "Retencion IIBB": linea(r"RET\.?IIBB[^\n]*"),
        "Jubilacion": linea("JUBILACION"),
        "Imp. ganancias": linea(r"IMP\.\s*A LAS GANANCIAS"),
        "Derecho administrativo": linea("DERECHO ADMINISTRATIVO"),
        "Aporte fondo compensador": linea(r"APORTE FONDO COMPENSADOR[^\n]*"),
        "Recup. bancarias": linea(r"RECUP\.?TRANSACCIONES BANCARIAS"),
        "Archivo": str(path.relative_to(CARPETA)),
    }


def _escribir_hoja(wb, nombre: str, df: pd.DataFrame, money_cols: list[str]) -> None:
    if df is None or df.empty:
        return
    ws = wb.create_sheet(nombre[:31])
    headers = list(df.columns)
    for i, h in enumerate(headers, 1):
        cell = ws.cell(1, i, h)
        cell.font = HDR_FONT
        cell.fill = HDR_FILL
        cell.alignment = Alignment(horizontal="right" if h in money_cols else "left", wrap_text=True)
    for r_i, row in enumerate(df.itertuples(index=False), 2):
        for c_i, val in enumerate(row, 1):
            h = headers[c_i - 1]
            cell = ws.cell(r_i, c_i, val if val is not None and not (isinstance(val, float) and pd.isna(val)) else None)
            cell.font = BODY
            if h in money_cols and cell.value is not None:
                try:
                    cell.value = float(cell.value)
                except (TypeError, ValueError):
                    pass
                cell.number_format = MONEY_FMT
            if (r_i % 2) == 1:
                cell.fill = ZEBRA
    last = 1 + len(df)
    tot_row = last + 1
    ws.cell(tot_row, 1, "TOTAL").font = BOLD
    for i, h in enumerate(headers, 1):
        if h in money_cols:
            c = ws.cell(tot_row, i, float(pd.to_numeric(df[h], errors="coerce").fillna(0).sum()))
            c.number_format = MONEY_FMT
            c.font = BOLD
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last}"
    ws.freeze_panes = "A2"
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, len(str(h)) + 2))


def extraer_otros(carpeta: Path) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    facturas: list[dict] = []
    vistos_fct: set[tuple] = set()
    pendientes: list[dict] = []
    resúmenes: list[dict] = []

    for pdf in sorted(carpeta.rglob("*")):
        if not pdf.is_file() or pdf.suffix.lower() != ".pdf":
            continue
        name = pdf.name.lower()
        if name.startswith("caratula_") or name.startswith("car_"):
            continue
        if "pendiente" in name or "pendientes" in name or name.startswith("facturas_iva"):
            pendientes.extend(parse_pendientes(pdf))
            continue
        if name.startswith("res_") or "resumen" in name:
            rec = parse_resumen_percibido(pdf)
            if rec:
                resúmenes.append(rec)
                continue
        # FCE / factura A (incluye PDFs mal llamados "Liquidacion Iva")
        rec = parse_factura_afip(pdf)
        if rec and rec.get("Nro"):
            key = (rec["Tipo"], rec["Punto venta"], rec["Nro"])
            if key in vistos_fct:
                continue
            vistos_fct.add(key)
            facturas.append(rec)

    return (
        pd.DataFrame(facturas),
        pd.DataFrame(pendientes),
        pd.DataFrame(resúmenes),
    )


def main() -> int:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    print(f"Carpeta: {CARPETA}")
    out = consolidar_carpeta(CARPETA, OUT)
    print(f"Carátulas -> {out}")

    from openpyxl import load_workbook

    fct, pend, resu = extraer_otros(CARPETA)
    print(f"Facturas: {len(fct)}  Pendientes: {len(pend)}  Resúmenes percibido: {len(resu)}")
    wb = load_workbook(out)
    money_fct = [
        "Honorarios exentos",
        "Honorarios gravados",
        "Gastos exentos",
        "Gastos gravados",
        "Importe exento",
        "Neto gravado",
        "IVA 10,5%",
        "Importe total",
    ]
    money_pend = [
        "Gravado honorarios",
        "Exento honorarios",
        "IVA honorarios",
        "Gravado gastos",
        "Exento gastos",
        "IVA gastos",
        "Total IVA",
        "Total general",
    ]
    money_res = [
        "Facturado bruto honorarios",
        "Facturado bruto gastos",
        "Retencion IIBB",
        "Jubilacion",
        "Imp. ganancias",
        "Derecho administrativo",
        "Aporte fondo compensador",
        "Recup. bancarias",
    ]
    _escribir_hoja(wb, "Facturas Centro Medico", fct, money_fct)
    _escribir_hoja(wb, "Pendientes", pend, money_pend)
    _escribir_hoja(wb, "Resumen percibido", resu, money_res)
    wb.save(out)
    try:
        import shutil

        shutil.copy2(out, OUT_DESKTOP)
    except OSError as exc:
        print("No se pudo copiar al escritorio:", exc)
    print(f"Excel: {out}")
    print(f"Copia: {OUT_DESKTOP}")
    if not fct.empty:
        print("Total facturas:", f"{fct['Importe total'].fillna(0).sum():,.2f}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
